from __future__ import annotations

from pathlib import Path
import csv
import math

import matplotlib.pyplot as plt
import numpy as np
import openpyxl

# This script intentionally reuses the correction functions and constants from
# the latest drag-corrected evaluation script. Keep both files in the same wtt
# folder when running this script.
from plot_cp_ct_by_yaw_2000_blockage_drag_corrected import (
    PROJECT_ROOT,
    DATA_DIR,
    TEST_MATRIX_PATH,
    ROTOR_RADIUS_M,
    ROTOR_AREA_M2,
    TORQUE_SCALE_TO_NM,
    HUB_HEIGHT_M,
    THRUST_LEVER_ARM_M,
    TUNNEL_AREA_M2,
    BLOCKAGE_RATIO_ALPHA,
    TARGET_LENGTH,
    load_model1,
    get_array,
    resample_to_2000,
    drag_corrected_rotor_thrust_and_velocity,
    safe_stat,
)


# ============================================================
# Output
# ============================================================

OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct_grouped_by_yaw_drag_corrected"


# ============================================================
# Excel mapping and yaw grouping
# ============================================================

def load_test_matrix(path: Path) -> list[dict]:
    """Read the test matrix and keep the Excel mapping exactly.

    The columns 'Yaw [deg]', 'Top' and 'File number' define which MAT file
    belongs to which test and which yaw case.
    """

    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["Task1 Performance G2"]

    headers = [
        sheet.cell(row=1, column=col).value
        for col in range(1, sheet.max_column + 1)
    ]

    records = []

    for row_index in range(2, sheet.max_row + 1):
        test_number = sheet.cell(row=row_index, column=1).value

        if test_number is None:
            continue

        record = {}

        for col_index, header in enumerate(headers, start=1):
            record[header] = sheet.cell(row=row_index, column=col_index).value

        records.append(record)

    return records


def yaw_suffix(yaw_value: float) -> str:
    """Create suffixes such as 0, plus_30 and minus_30."""

    yaw_int = int(round(float(yaw_value)))

    if yaw_int > 0:
        return f"plus_{yaw_int}"

    if yaw_int < 0:
        return f"minus_{abs(yaw_int)}"

    return "0"


def group_records_by_yaw(records: list[dict]) -> dict[str, dict]:
    """Group records by the Excel column 'Yaw [deg]' before MAT processing."""

    groups: dict[str, dict] = {}

    for record in records:
        yaw_target = record.get("Yaw [deg]")

        if yaw_target is None:
            continue

        suffix = yaw_suffix(yaw_target)
        yaw_key = f"yaw_{suffix}"

        if yaw_key not in groups:
            groups[yaw_key] = {
                "yaw_target": yaw_target,
                "suffix": suffix,
                "records": [],
            }

        groups[yaw_key]["records"].append(record)

    # Keep Excel/test order inside each yaw group.
    for group in groups.values():
        group["records"].sort(
            key=lambda r: (
                r.get("Test number", math.inf),
                r.get("Top", math.inf),
            )
        )

    return groups


def mat_path_from_record(record: dict) -> Path:
    """Build the MAT file path from Excel columns File number and Top."""

    file_number = record["File number"]
    top = record["Top"]
    mat_name = f"Results_File_{int(file_number)}_TOP_{int(top)}.mat"

    return DATA_DIR / mat_name


# ============================================================
# MAT access
# ============================================================

def get_nested_array(root, *names: str) -> np.ndarray:
    current = root

    for name in names:
        if not hasattr(current, name):
            return np.array([np.nan], dtype=float)
        current = getattr(current, name)

    return get_array(current)


def load_resampled_time_series(mat_path: Path) -> dict[str, np.ndarray]:
    """Load all relevant measured time series and resample each to 2000 values.

    This function is called inside one yaw group. Therefore, all interpolation
    and block averaging is performed independently for each yaw case.
    """

    model1 = load_model1(mat_path)

    raw = {
        # Channels directly required for cP/cT.
        "rotor_speed_rpm": get_nested_array(model1, "RotorSpeed"),
        "pitot_velocity": get_nested_array(model1, "PitotVelocity"),
        "air_density": get_nested_array(model1, "AirDensity"),
        "hub_torque_raw": get_nested_array(model1, "Hub", "Torque"),
        "tower_fa_moment": get_nested_array(model1, "Tower", "FA"),
        "yaw_actual": get_nested_array(model1, "Yaw"),

        # Additional measured channels kept separated by yaw as requested.
        "azimuth_deg": get_nested_array(model1, "Azimuth"),
        "tower_ss_moment": get_nested_array(model1, "Tower", "SS"),
        "hub_yawing_moment_raw": get_nested_array(model1, "Hub", "Yawing"),
        "hub_nodding_moment_raw": get_nested_array(model1, "Hub", "Nodding"),
        "pitch_blade1": get_nested_array(model1, "Pitch", "Blade1"),
        "pitch_blade1_dem": get_nested_array(model1, "Pitch", "Blade1Dem"),
        "pitch_blade2": get_nested_array(model1, "Pitch", "Blade2"),
        "pitch_blade2_dem": get_nested_array(model1, "Pitch", "Blade2Dem"),
        "pitch_blade3": get_nested_array(model1, "Pitch", "Blade3"),
        "pitch_blade3_dem": get_nested_array(model1, "Pitch", "Blade3Dem"),
        "yaw_dem": get_nested_array(model1, "YawDem"),
    }

    measured = {
        name: resample_to_2000(values, name)
        for name, values in raw.items()
    }

    measured["omega_rad_s"] = measured["rotor_speed_rpm"] * 2.0 * math.pi / 60.0
    measured["hub_torque_nm"] = measured["hub_torque_raw"] * TORQUE_SCALE_TO_NM
    measured["shaft_bending_moment_nm"] = np.sqrt(
        (measured["hub_yawing_moment_raw"] * TORQUE_SCALE_TO_NM) ** 2
        + (measured["hub_nodding_moment_raw"] * TORQUE_SCALE_TO_NM) ** 2
    )

    return measured


# ============================================================
# cP/cT calculation with unchanged drag and Glauert logic
# ============================================================

def compute_derived_time_series(measured: dict[str, np.ndarray]) -> dict[str, np.ndarray]:
    omega = measured["omega_rad_s"]
    torque_nm = measured["hub_torque_nm"]
    power_w = torque_nm * omega

    correction = drag_corrected_rotor_thrust_and_velocity(
        tower_fa_moment=measured["tower_fa_moment"],
        pitot_velocity=measured["pitot_velocity"],
        air_density=measured["air_density"],
    )

    thrust_raw_from_tower_fa = -measured["tower_fa_moment"] / THRUST_LEVER_ARM_M
    thrust_n = correction["rotor_thrust_n"]
    u_blockage_corrected = correction["u_blockage_corrected"]

    tsr = np.divide(
        omega * ROTOR_RADIUS_M,
        u_blockage_corrected,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=np.isfinite(u_blockage_corrected) & (u_blockage_corrected != 0.0),
    )

    cp_denominator = (
        0.5
        * measured["air_density"]
        * ROTOR_AREA_M2
        * u_blockage_corrected**3
    )
    cp = np.divide(
        power_w,
        cp_denominator,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=np.isfinite(cp_denominator) & (cp_denominator != 0.0),
    )

    ct_denominator = (
        0.5
        * measured["air_density"]
        * ROTOR_AREA_M2
        * u_blockage_corrected**2
    )
    ct = np.divide(
        thrust_n,
        ct_denominator,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=np.isfinite(ct_denominator) & (ct_denominator != 0.0),
    )

    return {
        "power_w": power_w,
        "thrust_raw_from_tower_fa": thrust_raw_from_tower_fa,
        "rotor_thrust_drag_corrected": thrust_n,
        "u_blockage_corrected": u_blockage_corrected,
        "ct_for_blockage": correction["ct_for_blockage"],
        "blockage_factor": correction["blockage_factor"],
        "nacelle_drag": correction["nacelle_drag_n"],
        "nacelle_drag_moment": correction["nacelle_drag_moment_nm"],
        "tower_drag": correction["tower_drag_n"],
        "tower_drag_moment": correction["tower_drag_moment_nm"],
        "tower_velocity": correction["tower_velocity"],
        "axial_induction": correction["axial_induction"],
        "tsr": tsr,
        "cp": cp,
        "ct": ct,
    }


# ============================================================
# Group-level processing
# ============================================================

def stack_time_series(items: list[dict[str, np.ndarray]]) -> dict[str, np.ndarray]:
    """Stack per-test time series into arrays with shape n_tests x 2000."""

    if not items:
        return {}

    keys = sorted(items[0].keys())
    stacked = {}

    for key in keys:
        stacked[key] = np.vstack([item[key] for item in items])

    return stacked


def build_yaw_case_arrays(groups: dict[str, dict]) -> dict[str, dict]:
    """Process each yaw case independently.

    Example resulting arrays:
        yaw_cases["yaw_0"]["measured"]["rotor_speed_rpm"]
        yaw_cases["yaw_plus_30"]["measured"]["rotor_speed_rpm"]
        yaw_cases["yaw_minus_30"]["measured"]["rotor_speed_rpm"]
    """

    yaw_cases: dict[str, dict] = {}

    for yaw_key, group in groups.items():
        measured_list = []
        derived_list = []
        processed_records = []

        for record in group["records"]:
            mat_path = mat_path_from_record(record)

            if not mat_path.exists():
                print(f"MAT file missing, skipped: {mat_path.name}")
                continue

            measured = load_resampled_time_series(mat_path)
            derived = compute_derived_time_series(measured)

            measured_list.append(measured)
            derived_list.append(derived)

            record_copy = dict(record)
            record_copy["mat_file"] = mat_path.name
            processed_records.append(record_copy)

        yaw_cases[yaw_key] = {
            "yaw_target": group["yaw_target"],
            "suffix": group["suffix"],
            "records": processed_records,
            "measured": stack_time_series(measured_list),
            "derived": stack_time_series(derived_list),
        }

    return yaw_cases


def compute_summary_rows(yaw_cases: dict[str, dict]) -> list[dict]:
    rows = []

    for yaw_key, yaw_case in yaw_cases.items():
        records = yaw_case["records"]
        measured = yaw_case["measured"]
        derived = yaw_case["derived"]

        if not records:
            continue

        for i, record in enumerate(records):
            cp = derived["cp"][i]
            ct = derived["ct"][i]
            tsr = derived["tsr"][i]

            cp_median = safe_stat(cp, np.median)
            cp_std = safe_stat(cp, np.std)
            ct_median = safe_stat(ct, np.median)
            ct_std = safe_stat(ct, np.std)

            rows.append(
                {
                    "yaw_group": yaw_key,
                    "test": record["Test number"],
                    "top": record["Top"],
                    "file_number": record["File number"],
                    "mat_file": record["mat_file"],
                    "yaw_target": record["Yaw [deg]"],
                    "yaw_actual_median": safe_stat(measured["yaw_actual"][i], np.median),
                    "tsr_expected": record["TSR [-]"],
                    "tsr_median": safe_stat(tsr, np.median),
                    "tsr_std": safe_stat(tsr, np.std),
                    "pitot_velocity_median": safe_stat(measured["pitot_velocity"][i], np.median),
                    "u_blockage_corrected_median": safe_stat(derived["u_blockage_corrected"][i], np.median),
                    "ct_for_blockage_median": safe_stat(derived["ct_for_blockage"][i], np.median),
                    "blockage_factor_median": safe_stat(derived["blockage_factor"][i], np.median),
                    "rotor_thrust_drag_corrected_median": safe_stat(derived["rotor_thrust_drag_corrected"][i], np.median),
                    "nacelle_drag_median": safe_stat(derived["nacelle_drag"][i], np.median),
                    "tower_drag_median": safe_stat(derived["tower_drag"][i], np.median),
                    "cp_median": cp_median,
                    "cp_plus_std": cp_median + cp_std if np.isfinite(cp_median) and np.isfinite(cp_std) else math.nan,
                    "cp_minus_std": cp_median - cp_std if np.isfinite(cp_median) and np.isfinite(cp_std) else math.nan,
                    "cp_min": safe_stat(cp, np.min),
                    "cp_max": safe_stat(cp, np.max),
                    "ct_median": ct_median,
                    "ct_plus_std": ct_median + ct_std if np.isfinite(ct_median) and np.isfinite(ct_std) else math.nan,
                    "ct_minus_std": ct_median - ct_std if np.isfinite(ct_median) and np.isfinite(ct_std) else math.nan,
                    "ct_min": safe_stat(ct, np.min),
                    "ct_max": safe_stat(ct, np.max),
                }
            )

    return rows


# ============================================================
# Output files and plots
# ============================================================

def save_summary_csv(rows: list[dict]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / "cp_ct_summary_grouped_by_yaw.csv"

    fieldnames = [
        "yaw_group",
        "test",
        "top",
        "file_number",
        "mat_file",
        "yaw_target",
        "yaw_actual_median",
        "tsr_expected",
        "tsr_median",
        "tsr_std",
        "pitot_velocity_median",
        "u_blockage_corrected_median",
        "ct_for_blockage_median",
        "blockage_factor_median",
        "rotor_thrust_drag_corrected_median",
        "nacelle_drag_median",
        "tower_drag_median",
        "cp_median",
        "cp_plus_std",
        "cp_minus_std",
        "cp_min",
        "cp_max",
        "ct_median",
        "ct_plus_std",
        "ct_minus_std",
        "ct_min",
        "ct_max",
    ]

    with output_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)

    return output_path


def save_grouped_arrays_npz(yaw_cases: dict[str, dict]) -> Path:
    """Save arrays with names like rotor_speed_rpm_0 and cp_plus_30."""

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / "yaw_grouped_time_series_arrays.npz"

    arrays = {}

    for yaw_case in yaw_cases.values():
        suffix = yaw_case["suffix"]

        for source in ["measured", "derived"]:
            for variable_name, values in yaw_case[source].items():
                arrays[f"{variable_name}_{suffix}"] = values

    if arrays:
        np.savez_compressed(output_path, **arrays)

    return output_path


def save_group_shapes_csv(yaw_cases: dict[str, dict]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / "yaw_group_array_shapes.csv"

    with output_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=["yaw_group", "source", "variable", "shape"],
            delimiter=";",
        )
        writer.writeheader()

        for yaw_key, yaw_case in yaw_cases.items():
            for source in ["measured", "derived"]:
                for variable_name, values in yaw_case[source].items():
                    writer.writerow(
                        {
                            "yaw_group": yaw_key,
                            "source": source,
                            "variable": variable_name,
                            "shape": str(values.shape),
                        }
                    )

    return output_path


def plot_for_yaw_group(rows: list[dict], yaw_key: str, yaw_target: float, coefficient: str) -> None:
    selected = [row for row in rows if row["yaw_group"] == yaw_key]
    selected.sort(key=lambda row: row["tsr_median"] if row["tsr_median"] is not None else math.inf)

    x = [row["tsr_median"] for row in selected]
    y_median = [row[f"{coefficient}_median"] for row in selected]
    y_plus_std = [row[f"{coefficient}_plus_std"] for row in selected]
    y_minus_std = [row[f"{coefficient}_minus_std"] for row in selected]
    y_min = [row[f"{coefficient}_min"] for row in selected]
    y_max = [row[f"{coefficient}_max"] for row in selected]

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(x, y_median, marker="o", label="Median")
    plt.plot(x, y_plus_std, marker="o", linestyle="--", label="Median + standard deviation")
    plt.plot(x, y_minus_std, marker="o", linestyle="--", label="Median - standard deviation")
    plt.plot(x, y_min, marker="o", linestyle=":", label="Minimum")
    plt.plot(x, y_max, marker="o", linestyle=":", label="Maximum")

    plt.xlabel("TSR with Glauert-corrected velocity")
    plt.ylabel(f"{label} [-]")
    plt.title(f"{label} over TSR at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    suffix = yaw_suffix(yaw_target)
    plt.savefig(OUTPUT_DIR / f"{label}_yaw_{suffix}.png", dpi=200)
    plt.close()


def main() -> None:
    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(f"Excel file not found: {TEST_MATRIX_PATH}")

    if not DATA_DIR.exists():
        raise FileNotFoundError(f"data folder not found: {DATA_DIR}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    records = load_test_matrix(TEST_MATRIX_PATH)

    # Step 1: group by Excel yaw case BEFORE loading/processing MAT data.
    grouped_records = group_records_by_yaw(records)

    # Step 2: process each yaw case independently.
    yaw_cases = build_yaw_case_arrays(grouped_records)

    # Explicit example aliases requested in the prompt.
    rotor_speed_rpm_0 = yaw_cases.get("yaw_0", {}).get("measured", {}).get("rotor_speed_rpm")
    rotor_speed_rpm_plus_30 = yaw_cases.get("yaw_plus_30", {}).get("measured", {}).get("rotor_speed_rpm")
    rotor_speed_rpm_minus_30 = yaw_cases.get("yaw_minus_30", {}).get("measured", {}).get("rotor_speed_rpm")
    _ = (rotor_speed_rpm_0, rotor_speed_rpm_plus_30, rotor_speed_rpm_minus_30)

    rows = compute_summary_rows(yaw_cases)

    if not rows:
        raise RuntimeError("No valid MAT files could be processed.")

    summary_csv = save_summary_csv(rows)
    arrays_npz = save_grouped_arrays_npz(yaw_cases)
    shapes_csv = save_group_shapes_csv(yaw_cases)

    for yaw_key, yaw_case in yaw_cases.items():
        if not yaw_case["records"]:
            continue
        plot_for_yaw_group(rows, yaw_key, yaw_case["yaw_target"], "cp")
        plot_for_yaw_group(rows, yaw_key, yaw_case["yaw_target"], "ct")

    print("Finished.")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")
    print(f"Summary CSV: {summary_csv.resolve()}")
    print(f"Grouped arrays NPZ: {arrays_npz.resolve()}")
    print(f"Array-shape CSV: {shapes_csv.resolve()}")

    print("\nYaw groups processed independently:")
    for yaw_key, yaw_case in yaw_cases.items():
        print(f" - {yaw_key}: {len(yaw_case['records'])} tests, yaw target = {yaw_case['yaw_target']}°")

    print("\nImportant:")
    print("The Excel columns 'Yaw [deg]', 'Top' and 'File number' define the grouping and MAT-file mapping.")
    print("Each yaw group is loaded, resampled, drag-corrected and Glauert-corrected independently.")
    print("All drag and Glauert correction logic is imported from the previous drag-corrected script.")
    print(f"Blockage ratio alpha: {BLOCKAGE_RATIO_ALPHA:.5f}")
    print(f"Wind tunnel area: {TUNNEL_AREA_M2:.3f} m²")
    print(f"Hub height / thrust arm: {HUB_HEIGHT_M:.3f} m")


if __name__ == "__main__":
    main()
