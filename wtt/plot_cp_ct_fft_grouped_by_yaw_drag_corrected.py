from __future__ import annotations

import csv
import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat


# ============================================================
# Paths and constants
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
TEST_MATRIX_PATH = DATA_DIR / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct_fft_grouped_by_yaw_drag_corrected"

MEASUREMENT_DURATION_S = 8.0
TARGET_LENGTH = 2000

ROTOR_RADIUS_M = 0.5436
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2

TORQUE_SCALE_TO_NM = 1e-3

# Geometry for the thrust / moment correction.
# The user-provided hub height is used as the thrust/nacelle moment arm.
HUB_HEIGHT_M = 0.80
THRUST_LEVER_ARM_M = HUB_HEIGHT_M
NACELLE_HUB_MOMENT_ARM_M = HUB_HEIGHT_M

# Nacelle-hub drag from the slides:
# D_nacelle = 0.5 * rho * U^2 * SCD
NACELLE_HUB_SCD_M2 = 0.0175

# Tower drag from the slides:
# D_T = 0.5 * rho * D * integral_0^H V(h)^2 * C_D(h,Re) dh
TOWER_HEIGHT_M = HUB_HEIGHT_M
TOWER_DIAMETER_M = 0.047

# Required modelling assumption for C_D(h,Re).
# A typical first estimate for a circular cylinder is used.
TOWER_DRAG_COEFFICIENT = 1.2

# Wind-tunnel dimensions for blockage ratio alpha = A_rotor / A_tunnel.
TUNNEL_WIDTH_M = 2.7
TUNNEL_HEIGHT_M = 1.8
TUNNEL_CROSS_SECTION_M2 = TUNNEL_WIDTH_M * TUNNEL_HEIGHT_M
BLOCKAGE_RATIO_ALPHA = ROTOR_AREA_M2 / TUNNEL_CROSS_SECTION_M2

CORRECTION_MAX_ITERATIONS = 50
CORRECTION_TOLERANCE = 1e-8


# ============================================================
# Data helpers
# ============================================================

def unwrap_singleton(value):
    current = value

    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]

    return current


def get_array(value) -> np.ndarray:
    array = np.asarray(
        unwrap_singleton(value),
        dtype=float,
    ).ravel(order="C")

    if array.size == 0:
        return np.array([np.nan], dtype=float)

    return array


def get_nested_array(root, *names: str) -> np.ndarray:
    current = root

    for name in names:
        current = unwrap_singleton(current)

        if not hasattr(current, name):
            return np.array([np.nan], dtype=float)

        current = getattr(current, name)

    return get_array(current)


def load_model1(mat_path: Path):
    mat_data = loadmat(
        mat_path,
        squeeze_me=True,
        struct_as_record=False,
    )
    data = unwrap_singleton(mat_data["Data"])
    models_data = unwrap_singleton(data.ModelsData)
    model1 = unwrap_singleton(models_data.Model1)

    return model1


def block_average_to_target_length(
    values: np.ndarray,
    target_length: int,
) -> np.ndarray:
    """Downsamples by block averaging.

    Example:
    Hub.Torque has 20000 samples. For target_length=2000, this function
    forms 2000 blocks with 10 samples each and returns one mean per block.
    """

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size % target_length != 0:
        raise ValueError(
            "Block averaging requires the signal length to be an integer "
            "multiple of target_length."
        )

    block_size = values.size // target_length
    blocks = values.reshape(target_length, block_size)

    with np.errstate(invalid="ignore"):
        return np.nanmean(blocks, axis=1)


def interpolate_to_target_length(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    """Linearly interpolates a signal over the common measurement period."""

    values = np.asarray(values, dtype=float).ravel(order="C")

    if target_length <= 0:
        raise ValueError("target_length must be greater than zero.")

    if values.size == 0:
        return np.full(target_length, np.nan)

    if values.size == 1:
        return np.full(target_length, values[0])

    if values.size == target_length:
        return values.astype(float, copy=False)

    finite = np.isfinite(values)

    if finite.sum() == 0:
        return np.full(target_length, np.nan)

    if finite.sum() == 1:
        return np.full(target_length, values[finite][0])

    source_time = np.linspace(
        0.0,
        duration_s,
        values.size,
        endpoint=False,
    )
    target_time = np.linspace(
        0.0,
        duration_s,
        target_length,
        endpoint=False,
    )

    return np.interp(
        target_time,
        source_time[finite],
        values[finite],
    )


def resample_to_2000(values: np.ndarray, signal_name: str) -> np.ndarray:
    """Resamples one signal to 2000 samples.

    This function is called within a yaw group. Therefore, the yaw cases are
    processed independently and no interpolation or clustering operation mixes
    data from different yaw cases.

    - 2000-sample signals remain unchanged.
    - Longer integer-ratio signals are block-averaged.
      This is intended for Hub.Torque: 20000 -> 2000.
    - Shorter signals are interpolated.
      This is intended for AirDensity: 1000 -> 2000.
    """

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size == TARGET_LENGTH:
        return values.astype(float, copy=False)

    if values.size > TARGET_LENGTH and values.size % TARGET_LENGTH == 0:
        return block_average_to_target_length(values, TARGET_LENGTH)

    return interpolate_to_target_length(values, TARGET_LENGTH)


def fill_nonfinite(values: np.ndarray) -> np.ndarray:
    """Fills invalid values linearly before the Fourier transform."""

    values = np.asarray(values, dtype=float).copy()
    finite = np.isfinite(values)

    if finite.sum() == 0:
        return np.full_like(values, np.nan)

    if finite.sum() == 1:
        return np.full_like(values, values[finite][0])

    if not finite.all():
        indices = np.arange(values.size)
        values[~finite] = np.interp(
            indices[~finite],
            indices[finite],
            values[finite],
        )

    return values


def safe_stat(values: np.ndarray, fn) -> float:
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]

    if values.size == 0:
        return math.nan

    return float(fn(values))


# ============================================================
# Fourier analysis
# ============================================================

def fourier_analysis(
    values: np.ndarray,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> dict:
    """Returns the one-sided FFT spectrum and the DC component.

    The normalized FFT component at 0 Hz equals the arithmetic mean of the
    original time series.
    """

    values = fill_nonfinite(values)

    if values.size == 0 or not np.isfinite(values).any():
        return {
            "dc": math.nan,
            "frequency_hz": np.array([]),
            "amplitude": np.array([]),
            "dominant_frequency_hz": math.nan,
            "dominant_amplitude": math.nan,
        }

    n = values.size
    dt = duration_s / n

    fft_values = np.fft.rfft(values)
    frequency_hz = np.fft.rfftfreq(n, d=dt)
    amplitude = np.abs(fft_values) / n

    if amplitude.size > 1:
        if n % 2 == 0:
            amplitude[1:-1] *= 2.0
        else:
            amplitude[1:] *= 2.0

    dc = float(np.real(fft_values[0]) / n)

    if amplitude.size > 1:
        dominant_index = int(np.argmax(amplitude[1:])) + 1
        dominant_frequency_hz = float(frequency_hz[dominant_index])
        dominant_amplitude = float(amplitude[dominant_index])
    else:
        dominant_frequency_hz = math.nan
        dominant_amplitude = math.nan

    return {
        "dc": dc,
        "frequency_hz": frequency_hz,
        "amplitude": amplitude,
        "dominant_frequency_hz": dominant_frequency_hz,
        "dominant_amplitude": dominant_amplitude,
    }


# ============================================================
# Blockage, nacelle-hub drag and tower drag
# ============================================================

def axial_induction_from_ct(ct: np.ndarray) -> np.ndarray:
    """Momentum-theory estimate for axial induction.

    CT = 4a(1-a) gives a = 0.5 * (1 - sqrt(1 - CT)).
    Values are clipped to the classical range.
    """

    ct = np.asarray(ct, dtype=float)
    ct_limited = np.clip(ct, 0.0, 0.999999)

    return 0.5 * (1.0 - np.sqrt(1.0 - ct_limited))


def glauert_velocity_from_thrust(
    u_inf: np.ndarray,
    air_density: np.ndarray,
    rotor_thrust_n: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Applies the Glauert blockage correction to the time series.

    Formula:
        U'_inf / U_inf = 1 + (alpha / 4) * C_T / (1 - C_T)

    with:
        C_T = T / (0.5 * rho * A_R * U_inf^2)

    C_T is the rotor thrust coefficient, not the torque coefficient.
    """

    u_inf = np.asarray(u_inf, dtype=float)
    air_density = np.asarray(air_density, dtype=float)
    rotor_thrust_n = np.asarray(rotor_thrust_n, dtype=float)

    ct_denominator_uncorrected = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * u_inf**2
    )

    ct_for_blockage = np.divide(
        rotor_thrust_n,
        ct_denominator_uncorrected,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(rotor_thrust_n)
            & np.isfinite(ct_denominator_uncorrected)
            & (ct_denominator_uncorrected > 0.0)
        ),
    )

    blockage_factor = np.full(TARGET_LENGTH, np.nan)

    valid = (
        np.isfinite(u_inf)
        & np.isfinite(ct_for_blockage)
        & (u_inf > 0.0)
        & (ct_for_blockage < 1.0)
        & np.isfinite(1.0 - ct_for_blockage)
        & ((1.0 - ct_for_blockage) != 0.0)
    )

    blockage_factor[valid] = (
        1.0
        + (BLOCKAGE_RATIO_ALPHA / 4.0)
        * ct_for_blockage[valid]
        / (1.0 - ct_for_blockage[valid])
    )

    u_blockage_corrected = u_inf * blockage_factor

    return u_blockage_corrected, ct_for_blockage, blockage_factor


def drag_corrected_rotor_thrust_and_velocity(
    tower_fa_moment: np.ndarray,
    pitot_velocity: np.ndarray,
    air_density: np.ndarray,
) -> dict[str, np.ndarray]:
    """Removes nacelle-hub and tower-drag moments from Tower.FA.

    Moment balance around the tower-base load measurement location:

        -M_FA = T_rotor * L
                + D_nacelle * z_hub
                + M_tower_drag

    Therefore:

        T_rotor = (-M_FA - D_nacelle*z_hub - M_tower_drag) / L

    The correction logic is kept consistent with the uploaded FFT script.
    """

    tower_fa_moment = np.asarray(tower_fa_moment, dtype=float)
    pitot_velocity = np.asarray(pitot_velocity, dtype=float)
    air_density = np.asarray(air_density, dtype=float)

    measured_positive_moment = -tower_fa_moment

    # Initial estimate: all measured moment is due to rotor thrust.
    rotor_thrust_n = measured_positive_moment / THRUST_LEVER_ARM_M

    u_blockage_corrected = pitot_velocity.copy()
    ct_for_blockage = np.full(TARGET_LENGTH, np.nan)
    blockage_factor = np.full(TARGET_LENGTH, np.nan)

    nacelle_drag_n = np.zeros(TARGET_LENGTH)
    nacelle_drag_moment_nm = np.zeros(TARGET_LENGTH)

    tower_drag_n = np.zeros(TARGET_LENGTH)
    tower_drag_moment_nm = np.zeros(TARGET_LENGTH)
    tower_velocity = np.zeros(TARGET_LENGTH)
    axial_induction = np.zeros(TARGET_LENGTH)

    for _ in range(CORRECTION_MAX_ITERATIONS):
        previous_rotor_thrust_n = rotor_thrust_n.copy()

        (
            u_blockage_corrected,
            ct_for_blockage,
            blockage_factor,
        ) = glauert_velocity_from_thrust(
            u_inf=pitot_velocity,
            air_density=air_density,
            rotor_thrust_n=rotor_thrust_n,
        )

        ct_corrected_denominator = (
            0.5
            * air_density
            * ROTOR_AREA_M2
            * u_blockage_corrected**2
        )

        ct_corrected = np.divide(
            rotor_thrust_n,
            ct_corrected_denominator,
            out=np.full(TARGET_LENGTH, np.nan),
            where=(
                np.isfinite(rotor_thrust_n)
                & np.isfinite(ct_corrected_denominator)
                & (ct_corrected_denominator > 0.0)
            ),
        )

        axial_induction = axial_induction_from_ct(ct_corrected)

        nacelle_drag_n = (
            0.5
            * air_density
            * u_blockage_corrected**2
            * NACELLE_HUB_SCD_M2
        )

        nacelle_drag_moment_nm = (
            nacelle_drag_n
            * NACELLE_HUB_MOMENT_ARM_M
        )

        tower_velocity = u_blockage_corrected * (1.0 - 2.0 * axial_induction)
        tower_velocity = np.maximum(tower_velocity, 0.0)

        tower_drag_n = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * TOWER_HEIGHT_M
        )

        tower_drag_moment_nm = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * (TOWER_HEIGHT_M**2 / 2.0)
        )

        rotor_moment_nm = (
            measured_positive_moment
            - nacelle_drag_moment_nm
            - tower_drag_moment_nm
        )

        rotor_thrust_n = np.divide(
            rotor_moment_nm,
            THRUST_LEVER_ARM_M,
            out=np.full(TARGET_LENGTH, np.nan),
            where=np.isfinite(rotor_moment_nm),
        )

        rotor_thrust_n = np.where(
            np.isfinite(rotor_thrust_n),
            np.maximum(rotor_thrust_n, 0.0),
            np.nan,
        )

        difference = np.abs(rotor_thrust_n - previous_rotor_thrust_n)
        finite_difference = difference[np.isfinite(difference)]

        if finite_difference.size == 0:
            break

        if np.max(finite_difference) < CORRECTION_TOLERANCE:
            break

    return {
        "measured_positive_moment": measured_positive_moment,
        "rotor_thrust_n": rotor_thrust_n,
        "u_blockage_corrected": u_blockage_corrected,
        "ct_for_blockage": ct_for_blockage,
        "blockage_factor": blockage_factor,

        "nacelle_drag_n": nacelle_drag_n,
        "nacelle_drag_moment_nm": nacelle_drag_moment_nm,

        "tower_drag_n": tower_drag_n,
        "tower_drag_moment_nm": tower_drag_moment_nm,
        "tower_velocity": tower_velocity,
        "axial_induction": axial_induction,
    }


# ============================================================
# Excel mapping and yaw grouping
# ============================================================

def load_test_matrix() -> list[dict]:
    workbook = openpyxl.load_workbook(
        TEST_MATRIX_PATH,
        data_only=True,
    )
    sheet = workbook["Task1 Performance G2"]

    headers = [
        sheet.cell(row=1, column=column_index).value
        for column_index in range(1, sheet.max_column + 1)
    ]

    rows = []

    for row_index in range(2, sheet.max_row + 1):
        test_number = sheet.cell(row_index, 1).value

        if test_number is None:
            continue

        record = {}

        for column_index, header in enumerate(headers, start=1):
            record[header] = sheet.cell(row_index, column_index).value

        rows.append(record)

    return rows


def yaw_suffix(yaw_value: float) -> str:
    yaw_int = int(round(float(yaw_value)))

    if yaw_int > 0:
        return f"plus_{yaw_int}"

    if yaw_int < 0:
        return f"minus_{abs(yaw_int)}"

    return "0"


def group_records_by_yaw(records: list[dict]) -> dict[str, dict]:
    """Groups Excel records by the Yaw [deg] column before MAT processing."""

    grouped = {}

    for record in records:
        yaw_target = record["Yaw [deg]"]

        if yaw_target is None:
            continue

        suffix = yaw_suffix(yaw_target)
        key = f"yaw_{suffix}"

        if key not in grouped:
            grouped[key] = {
                "yaw_target": yaw_target,
                "suffix": suffix,
                "records": [],
            }

        grouped[key]["records"].append(record)

    for group in grouped.values():
        group["records"].sort(
            key=lambda record: (
                record["Test number"],
                record["Top"],
            )
        )

    return grouped


def mat_file_from_record(record: dict) -> Path:
    top = record["Top"]
    file_number = record["File number"]

    mat_name = f"Results_File_{int(file_number)}_TOP_{int(top)}.mat"

    return DATA_DIR / mat_name


# ============================================================
# MAT loading and time-series calculations per yaw group
# ============================================================

def load_resampled_time_series_for_record(mat_path: Path) -> dict[str, np.ndarray]:
    """Loads relevant measured time series and maps them to 2000 samples."""

    model1 = load_model1(mat_path)

    raw = {
        "azimuth_deg": get_nested_array(model1, "Azimuth"),
        "rotor_speed_rpm": get_nested_array(model1, "RotorSpeed"),
        "pitot_velocity": get_nested_array(model1, "PitotVelocity"),
        "air_density": get_nested_array(model1, "AirDensity"),

        "hub_torque_raw": get_nested_array(model1, "Hub", "Torque"),
        "hub_yawing_moment_raw": get_nested_array(model1, "Hub", "Yawing"),
        "hub_nodding_moment_raw": get_nested_array(model1, "Hub", "Nodding"),

        "tower_fa_moment": get_nested_array(model1, "Tower", "FA"),
        "tower_ss_moment": get_nested_array(model1, "Tower", "SS"),

        "pitch_blade1": get_nested_array(model1, "Pitch", "Blade1"),
        "pitch_blade1_dem": get_nested_array(model1, "Pitch", "Blade1Dem"),
        "pitch_blade2": get_nested_array(model1, "Pitch", "Blade2"),
        "pitch_blade2_dem": get_nested_array(model1, "Pitch", "Blade2Dem"),
        "pitch_blade3": get_nested_array(model1, "Pitch", "Blade3"),
        "pitch_blade3_dem": get_nested_array(model1, "Pitch", "Blade3Dem"),

        "yaw_dem": get_nested_array(model1, "YawDem"),
        "yaw_actual": get_nested_array(model1, "Yaw"),
    }

    resampled = {
        name: resample_to_2000(values, name)
        for name, values in raw.items()
    }

    resampled["omega_rad_s"] = (
        resampled["rotor_speed_rpm"]
        * 2.0
        * math.pi
        / 60.0
    )

    resampled["tip_speed"] = (
        resampled["omega_rad_s"]
        * ROTOR_RADIUS_M
    )

    resampled["hub_torque_nm"] = (
        resampled["hub_torque_raw"]
        * TORQUE_SCALE_TO_NM
    )

    resampled["shaft_bending_moment_nm"] = np.sqrt(
        (resampled["hub_yawing_moment_raw"] * TORQUE_SCALE_TO_NM) ** 2
        + (resampled["hub_nodding_moment_raw"] * TORQUE_SCALE_TO_NM) ** 2
    )

    return resampled


def compute_derived_time_series(
    measured: dict[str, np.ndarray],
) -> dict[str, np.ndarray]:
    """Computes drag-corrected thrust, Glauert velocity, TSR, cP and cT."""

    omega = measured["omega_rad_s"]
    torque_nm = measured["hub_torque_nm"]
    power_w = torque_nm * omega

    correction = drag_corrected_rotor_thrust_and_velocity(
        tower_fa_moment=measured["tower_fa_moment"],
        pitot_velocity=measured["pitot_velocity"],
        air_density=measured["air_density"],
    )

    thrust_raw_from_tower_fa = (
        -measured["tower_fa_moment"]
        / THRUST_LEVER_ARM_M
    )

    thrust_n = correction["rotor_thrust_n"]
    u_blockage_corrected = correction["u_blockage_corrected"]

    tsr = np.divide(
        measured["tip_speed"],
        u_blockage_corrected,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(u_blockage_corrected)
            & (u_blockage_corrected != 0.0)
        ),
    )

    cp_denominator = (
        0.5
        * measured["air_density"]
        * ROTOR_AREA_M2
        * u_blockage_corrected**3
    )

    cp = np.divide(
        power_w,
        cp_denominator,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(cp_denominator)
            & (cp_denominator != 0.0)
        ),
    )

    ct_denominator = (
        0.5
        * measured["air_density"]
        * ROTOR_AREA_M2
        * u_blockage_corrected**2
    )

    ct = np.divide(
        thrust_n,
        ct_denominator,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(ct_denominator)
            & (ct_denominator != 0.0)
        ),
    )

    return {
        "power_w": power_w,
        "thrust_raw_from_tower_fa": thrust_raw_from_tower_fa,
        "rotor_thrust_drag_corrected": thrust_n,
        "u_blockage_corrected": u_blockage_corrected,
        "ct_for_blockage": correction["ct_for_blockage"],
        "blockage_factor": correction["blockage_factor"],

        "nacelle_drag": correction["nacelle_drag_n"],
        "nacelle_drag_moment": correction["nacelle_drag_moment_nm"],
        "tower_drag": correction["tower_drag_n"],
        "tower_drag_moment": correction["tower_drag_moment_nm"],
        "tower_velocity": correction["tower_velocity"],
        "axial_induction": correction["axial_induction"],

        "tsr": tsr,
        "cp": cp,
        "ct": ct,
    }


def stack_series(series_list: list[dict[str, np.ndarray]]) -> dict[str, np.ndarray]:
    """Turns per-test dictionaries into one 2D array per variable.

    Shape of each array:
        number_of_tests_in_this_yaw_case x 2000
    """

    if not series_list:
        return {}

    keys = sorted(series_list[0].keys())
    stacked = {}

    for key in keys:
        stacked[key] = np.vstack([item[key] for item in series_list])

    return stacked


def build_yaw_case_arrays(
    grouped_records: dict[str, dict],
) -> dict[str, dict]:
    """Loads and processes each yaw case independently.

    This creates arrays equivalent to names such as:
        rotor_speed_rpm_0
        rotor_speed_rpm_plus_30
        rotor_speed_rpm_minus_30

    Stored as:
        yaw_cases["yaw_0"]["measured"]["rotor_speed_rpm"]
        yaw_cases["yaw_plus_30"]["measured"]["rotor_speed_rpm"]
        yaw_cases["yaw_minus_30"]["measured"]["rotor_speed_rpm"]
    """

    yaw_cases = {}

    for yaw_key, group in grouped_records.items():
        measured_series_list = []
        derived_series_list = []
        processed_records = []

        for record in group["records"]:
            mat_path = mat_file_from_record(record)

            if not mat_path.exists():
                print(f"MAT file missing, skipped: {mat_path.name}")
                continue

            measured = load_resampled_time_series_for_record(mat_path)
            derived = compute_derived_time_series(measured)

            measured_series_list.append(measured)
            derived_series_list.append(derived)

            record_with_file = dict(record)
            record_with_file["mat_file"] = mat_path.name
            processed_records.append(record_with_file)

        yaw_cases[yaw_key] = {
            "yaw_target": group["yaw_target"],
            "suffix": group["suffix"],
            "records": processed_records,
            "measured": stack_series(measured_series_list),
            "derived": stack_series(derived_series_list),
        }

    return yaw_cases


# ============================================================
# FFT per yaw group
# ============================================================

def fft_for_single_test(
    measured: dict[str, np.ndarray],
    derived: dict[str, np.ndarray],
    test_index: int,
) -> dict:
    """Applies Fourier analysis to one already grouped and processed test."""

    cp_fft = fourier_analysis(derived["cp"][test_index])
    ct_fft = fourier_analysis(derived["ct"][test_index])

    return {
        "yaw_fft": fourier_analysis(measured["yaw_actual"][test_index]),
        "tsr_fft": fourier_analysis(derived["tsr"][test_index]),
        "tip_speed_fft": fourier_analysis(measured["tip_speed"][test_index]),
        "cp_fft": cp_fft,
        "ct_fft": ct_fft,

        "pitot_fft": fourier_analysis(measured["pitot_velocity"][test_index]),
        "u_blockage_corrected_fft": fourier_analysis(
            derived["u_blockage_corrected"][test_index]
        ),
        "blockage_factor_fft": fourier_analysis(
            derived["blockage_factor"][test_index]
        ),
        "ct_for_blockage_fft": fourier_analysis(
            derived["ct_for_blockage"][test_index]
        ),

        "thrust_raw_fft": fourier_analysis(
            derived["thrust_raw_from_tower_fa"][test_index]
        ),
        "rotor_thrust_fft": fourier_analysis(
            derived["rotor_thrust_drag_corrected"][test_index]
        ),

        "nacelle_drag_fft": fourier_analysis(
            derived["nacelle_drag"][test_index]
        ),
        "nacelle_drag_moment_fft": fourier_analysis(
            derived["nacelle_drag_moment"][test_index]
        ),

        "tower_drag_fft": fourier_analysis(
            derived["tower_drag"][test_index]
        ),
        "tower_drag_moment_fft": fourier_analysis(
            derived["tower_drag_moment"][test_index]
        ),
        "tower_velocity_fft": fourier_analysis(
            derived["tower_velocity"][test_index]
        ),
        "axial_induction_fft": fourier_analysis(
            derived["axial_induction"][test_index]
        ),
    }


def compute_fft_rows(yaw_cases: dict[str, dict]) -> list[dict]:
    rows = []

    for yaw_key, yaw_case in yaw_cases.items():
        records = yaw_case["records"]
        measured = yaw_case["measured"]
        derived = yaw_case["derived"]

        if not records:
            continue

        for i, record in enumerate(records):
            fft_result = fft_for_single_test(
                measured=measured,
                derived=derived,
                test_index=i,
            )

            rows.append(
                {
                    "yaw_group": yaw_key,
                    "test": record["Test number"],
                    "yaw_target": record["Yaw [deg]"],
                    "top": record["Top"],
                    "file_number": record["File number"],
                    "mat_file": record["mat_file"],

                    "yaw_fft_dc": fft_result["yaw_fft"]["dc"],
                    "tsr_expected": record["TSR [-]"],
                    "tsr_fft_dc": fft_result["tsr_fft"]["dc"],
                    "tip_speed_fft_dc": fft_result["tip_speed_fft"]["dc"],

                    "cp_fft_dc": fft_result["cp_fft"]["dc"],
                    "ct_fft_dc": fft_result["ct_fft"]["dc"],

                    "pitot_fft_dc": fft_result["pitot_fft"]["dc"],
                    "u_blockage_corrected_fft_dc":
                        fft_result["u_blockage_corrected_fft"]["dc"],
                    "blockage_factor_fft_dc":
                        fft_result["blockage_factor_fft"]["dc"],
                    "ct_for_blockage_fft_dc":
                        fft_result["ct_for_blockage_fft"]["dc"],

                    "thrust_raw_fft_dc": fft_result["thrust_raw_fft"]["dc"],
                    "rotor_thrust_fft_dc":
                        fft_result["rotor_thrust_fft"]["dc"],

                    "nacelle_drag_fft_dc":
                        fft_result["nacelle_drag_fft"]["dc"],
                    "nacelle_drag_moment_fft_dc":
                        fft_result["nacelle_drag_moment_fft"]["dc"],

                    "tower_drag_fft_dc": fft_result["tower_drag_fft"]["dc"],
                    "tower_drag_moment_fft_dc":
                        fft_result["tower_drag_moment_fft"]["dc"],
                    "tower_velocity_fft_dc":
                        fft_result["tower_velocity_fft"]["dc"],
                    "axial_induction_fft_dc":
                        fft_result["axial_induction_fft"]["dc"],

                    "cp_dominant_frequency_hz":
                        fft_result["cp_fft"]["dominant_frequency_hz"],
                    "cp_dominant_amplitude":
                        fft_result["cp_fft"]["dominant_amplitude"],
                    "ct_dominant_frequency_hz":
                        fft_result["ct_fft"]["dominant_frequency_hz"],
                    "ct_dominant_amplitude":
                        fft_result["ct_fft"]["dominant_amplitude"],

                    "cp_frequency_hz": fft_result["cp_fft"]["frequency_hz"],
                    "cp_amplitude": fft_result["cp_fft"]["amplitude"],
                    "ct_frequency_hz": fft_result["ct_fft"]["frequency_hz"],
                    "ct_amplitude": fft_result["ct_fft"]["amplitude"],
                }
            )

    return rows


# ============================================================
# Output
# ============================================================

def save_summary_csv(results: list[dict]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / "cp_ct_fft_summary_grouped_by_yaw.csv"

    columns = [
        "yaw_group",
        "test",
        "yaw_target",
        "top",
        "file_number",
        "mat_file",

        "yaw_fft_dc",
        "tsr_expected",
        "tsr_fft_dc",
        "tip_speed_fft_dc",

        "cp_fft_dc",
        "ct_fft_dc",

        "pitot_fft_dc",
        "u_blockage_corrected_fft_dc",
        "blockage_factor_fft_dc",
        "ct_for_blockage_fft_dc",

        "thrust_raw_fft_dc",
        "rotor_thrust_fft_dc",

        "nacelle_drag_fft_dc",
        "nacelle_drag_moment_fft_dc",

        "tower_drag_fft_dc",
        "tower_drag_moment_fft_dc",
        "tower_velocity_fft_dc",
        "axial_induction_fft_dc",

        "cp_dominant_frequency_hz",
        "cp_dominant_amplitude",
        "ct_dominant_frequency_hz",
        "ct_dominant_amplitude",
    ]

    with output_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=columns,
            delimiter=";",
        )
        writer.writeheader()

        for result in results:
            writer.writerow(
                {column: result[column] for column in columns}
            )

    return output_path


def save_grouped_arrays_npz(yaw_cases: dict[str, dict]) -> Path:
    """Saves grouped time-series arrays for checking/reuse.

    Example keys:
        rotor_speed_rpm_0
        rotor_speed_rpm_plus_30
        rotor_speed_rpm_minus_30
        cp_0
        cp_plus_30
        cp_minus_30
        ct_0
        ct_plus_30
        ct_minus_30
    """

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_npz = OUTPUT_DIR / "yaw_grouped_time_series_arrays_fft_input.npz"

    arrays = {}

    for yaw_key, yaw_case in yaw_cases.items():
        suffix = yaw_case["suffix"]

        for source_name in ["measured", "derived"]:
            for variable_name, values in yaw_case[source_name].items():
                arrays[f"{variable_name}_{suffix}"] = values

    if arrays:
        np.savez_compressed(out_npz, **arrays)

    return out_npz


def save_group_shapes_csv(yaw_cases: dict[str, dict]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_csv = OUTPUT_DIR / "yaw_group_array_shapes_fft_input.csv"

    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["yaw_group", "source", "variable", "shape"],
            delimiter=";",
        )
        writer.writeheader()

        for yaw_key, yaw_case in yaw_cases.items():
            for source_name in ["measured", "derived"]:
                for variable_name, values in yaw_case[source_name].items():
                    writer.writerow(
                        {
                            "yaw_group": yaw_key,
                            "source": source_name,
                            "variable": variable_name,
                            "shape": str(values.shape),
                        }
                    )

    return out_csv


def plot_performance(
    results: list[dict],
    yaw_group: str,
    yaw_target: float,
    coefficient: str,
) -> None:
    selected = [
        result
        for result in results
        if result["yaw_group"] == yaw_group
    ]
    selected.sort(key=lambda result: result["tsr_fft_dc"])

    x = [result["tsr_fft_dc"] for result in selected]
    y = [result[f"{coefficient}_fft_dc"] for result in selected]

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(
        x,
        y,
        marker="o",
        label="FFT DC component",
    )
    plt.xlabel("TSR from FFT DC component")
    plt.ylabel(f"{label} [-]")
    plt.title(f"{label} over TSR at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    suffix = yaw_suffix(yaw_target)
    plt.savefig(
        OUTPUT_DIR / f"{label}_fft_yaw_{suffix}.png",
        dpi=200,
    )
    plt.close()


def plot_mean_spectrum(
    results: list[dict],
    yaw_group: str,
    yaw_target: float,
    coefficient: str,
) -> None:
    selected = [
        result
        for result in results
        if result["yaw_group"] == yaw_group
    ]

    if not selected:
        return

    frequency_key = f"{coefficient}_frequency_hz"
    amplitude_key = f"{coefficient}_amplitude"

    lengths = [
        len(result[frequency_key])
        for result in selected
        if len(result[frequency_key]) > 0
    ]

    if not lengths:
        return

    common_length = min(lengths)
    frequency = selected[0][frequency_key][:common_length]
    amplitudes = np.vstack(
        [
            result[amplitude_key][:common_length]
            for result in selected
        ]
    )
    mean_amplitude = np.nanmean(amplitudes, axis=0)

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(
        frequency,
        mean_amplitude,
        label="Mean one-sided FFT spectrum",
    )
    plt.xlabel("Frequency [Hz]")
    plt.ylabel(f"{label} amplitude")
    plt.title(f"{label} Fourier spectrum at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    suffix = yaw_suffix(yaw_target)
    plt.savefig(
        OUTPUT_DIR / f"{label}_spectrum_yaw_{suffix}.png",
        dpi=200,
    )
    plt.close()


def main() -> None:
    if not DATA_DIR.exists():
        raise FileNotFoundError(f"Data folder not found: {DATA_DIR}")

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(
            f"Excel file not found: {TEST_MATRIX_PATH}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1) Excel is the reference for MAT-file mapping.
    # 2) Group by Yaw [deg] before any MAT-file processing.
    records = load_test_matrix()
    grouped_records = group_records_by_yaw(records)

    # 3) Load/resample/drag-correct/Glauert-correct each yaw group independently.
    yaw_cases = build_yaw_case_arrays(grouped_records)

    # Optional aliases showing the requested separated structure.
    rotor_speed_rpm_0 = yaw_cases.get("yaw_0", {}).get("measured", {}).get("rotor_speed_rpm")
    rotor_speed_rpm_plus_30 = yaw_cases.get("yaw_plus_30", {}).get("measured", {}).get("rotor_speed_rpm")
    rotor_speed_rpm_minus_30 = yaw_cases.get("yaw_minus_30", {}).get("measured", {}).get("rotor_speed_rpm")

    _ = (
        rotor_speed_rpm_0,
        rotor_speed_rpm_plus_30,
        rotor_speed_rpm_minus_30,
    )

    # 4) Fourier transformation is then applied within each separated yaw group.
    results = compute_fft_rows(yaw_cases)

    if not results:
        raise RuntimeError("No tests could be processed.")

    summary_csv = save_summary_csv(results)
    arrays_npz = save_grouped_arrays_npz(yaw_cases)
    shapes_csv = save_group_shapes_csv(yaw_cases)

    for yaw_group, yaw_case in yaw_cases.items():
        if not yaw_case["records"]:
            continue

        yaw_target = yaw_case["yaw_target"]
        plot_performance(results, yaw_group, yaw_target, "cp")
        plot_performance(results, yaw_group, yaw_target, "ct")
        plot_mean_spectrum(results, yaw_group, yaw_target, "cp")
        plot_mean_spectrum(results, yaw_group, yaw_target, "ct")

    print("Finished.")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")
    print(f"Summary CSV: {summary_csv.resolve()}")
    print(f"Grouped arrays NPZ: {arrays_npz.resolve()}")
    print(f"Array-shape CSV: {shapes_csv.resolve()}")

    print("\nYaw groups processed independently before FFT:")
    for yaw_key, yaw_case in yaw_cases.items():
        print(
            f" - {yaw_key}: "
            f"{len(yaw_case['records'])} tests, "
            f"yaw target = {yaw_case['yaw_target']}°"
        )

    print("\nImportant:")
    print("The Excel columns 'Yaw [deg]', 'Top', and 'File number' define grouping and MAT-file mapping.")
    print("Each yaw group is loaded, resampled, drag-corrected, Glauert-corrected and Fourier-transformed independently.")
    print("Hub.Torque is block-averaged from 20000 to 2000 samples when applicable.")
    print("AirDensity is interpolated from 1000 to 2000 samples when applicable.")
    print("All drag and Glauert correction logic is kept consistent with the uploaded FFT script.")
    print("The normalized FFT value at 0 Hz equals the arithmetic mean, not the median.")
    print(f"Hub height / thrust arm: {HUB_HEIGHT_M:.3f} m")
    print(f"Nacelle-hub SCD: {NACELLE_HUB_SCD_M2:.5f} m²")
    print(f"Tower diameter: {TOWER_DIAMETER_M:.3f} m")
    print(f"Tower drag coefficient assumption: {TOWER_DRAG_COEFFICIENT:.2f}")
    print(f"Blockage ratio alpha: {BLOCKAGE_RATIO_ALPHA:.5f}")


if __name__ == "__main__":
    main()
