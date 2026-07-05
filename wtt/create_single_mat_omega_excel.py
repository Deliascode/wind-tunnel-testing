from __future__ import annotations

import math
import sys
from pathlib import Path

import numpy as np
from scipy.io import loadmat
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# Settings
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
OUTPUT_DIR = PROJECT_ROOT / "omega_single_mat"

# Default single case. You can override this from the command line:
# py .\wtt\create_single_mat_omega_excel.py Results_File_5_TOP_1.mat
DEFAULT_MAT_FILE_NAME = "Results_File_5_TOP_1.mat"

MEASUREMENT_DURATION_S = 8.0
TARGET_LENGTH = 2000

RPM_TO_RAD_PER_S = 2.0 * math.pi / 60.0


# ============================================================
# MAT helpers
# ============================================================

def unwrap_singleton(value):
    current = value

    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]

    return current


def load_model1(mat_path: Path):
    mat_data = loadmat(
        mat_path,
        squeeze_me=True,
        struct_as_record=False,
    )

    data = unwrap_singleton(mat_data["Data"])
    models_data = unwrap_singleton(data.ModelsData)
    model1 = unwrap_singleton(models_data.Model1)

    return model1


def get_array(value) -> np.ndarray:
    array = np.asarray(
        unwrap_singleton(value),
        dtype=float,
    ).ravel(order="C")

    if array.size == 0:
        return np.array([np.nan], dtype=float)

    return array


def block_average_to_target_length(
    values: np.ndarray,
    target_length: int,
) -> np.ndarray:
    """Downsamples by block averaging if the ratio is an exact integer."""

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size % target_length != 0:
        raise ValueError(
            "Block averaging requires an integer ratio between signal length "
            "and target length."
        )

    block_size = values.size // target_length
    blocks = values.reshape(target_length, block_size)

    with np.errstate(invalid="ignore"):
        return np.nanmean(blocks, axis=1)


def interpolate_to_target_length(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    """Interpolates a signal onto a target-length time axis."""

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size == target_length:
        return values.astype(float, copy=False)

    if values.size == 0:
        return np.full(target_length, np.nan, dtype=float)

    if values.size == 1:
        return np.full(target_length, values[0], dtype=float)

    finite_mask = np.isfinite(values)

    if finite_mask.sum() == 0:
        return np.full(target_length, np.nan, dtype=float)

    if finite_mask.sum() == 1:
        return np.full(
            target_length,
            values[finite_mask][0],
            dtype=float,
        )

    source_time = np.linspace(
        0.0,
        duration_s,
        values.size,
        endpoint=False,
    )
    target_time = np.linspace(
        0.0,
        duration_s,
        target_length,
        endpoint=False,
    )

    return np.interp(
        target_time,
        source_time[finite_mask],
        values[finite_mask],
    )


def resample_to_2000(values: np.ndarray) -> np.ndarray:
    """Returns exactly 2000 values.

    For RotorSpeed this normally keeps the original 2000 values unchanged.
    This function is included so the rule is explicit and consistent.
    """

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size == TARGET_LENGTH:
        return values.astype(float, copy=False)

    if values.size > TARGET_LENGTH and values.size % TARGET_LENGTH == 0:
        return block_average_to_target_length(values, TARGET_LENGTH)

    return interpolate_to_target_length(values, TARGET_LENGTH)


def finite_stat(values: np.ndarray, function) -> float:
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]

    if values.size == 0:
        return math.nan

    return float(function(values))


def fourier_dc(values: np.ndarray) -> float:
    """The normalized FFT 0 Hz component equals the arithmetic mean."""

    values = np.asarray(values, dtype=float)

    finite = np.isfinite(values)

    if finite.sum() == 0:
        return math.nan

    # For the DC component, the mean is exactly the required value.
    return float(np.mean(values[finite]))


# ============================================================
# Omega calculation
# ============================================================

def compute_omega_time_series(mat_path: Path) -> dict:
    """Loads one MAT file and computes only omega from RotorSpeed."""

    model1 = load_model1(mat_path)

    rotor_speed_raw = get_array(model1.RotorSpeed)
    rotor_speed_2000 = resample_to_2000(rotor_speed_raw)

    omega = rotor_speed_2000 * RPM_TO_RAD_PER_S

    time_s = np.linspace(
        0.0,
        MEASUREMENT_DURATION_S,
        TARGET_LENGTH,
        endpoint=False,
    )

    return {
        "time_s": time_s,
        "rotor_speed_raw_count": int(rotor_speed_raw.size),
        "rotor_speed_rpm": rotor_speed_2000,
        "omega_rad_s": omega,
    }


# ============================================================
# Excel output
# ============================================================

def write_excel(mat_file_name: str, data: dict) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    output_path = (
        OUTPUT_DIR
        / f"omega_time_series_{Path(mat_file_name).stem}.xlsx"
    )

    time_s = data["time_s"]
    rotor_speed_rpm = data["rotor_speed_rpm"]
    omega = data["omega_rad_s"]

    workbook = Workbook()
    time_sheet = workbook.active
    time_sheet.title = "Omega_TimeSeries"

    summary_sheet = workbook.create_sheet("Omega_Summary")

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    # --------------------------------------------------------
    # Time series sheet
    # --------------------------------------------------------
    headers = [
        "sample_index",
        "time_s",
        "rotor_speed_rpm",
        "omega_rad_s",
    ]

    time_sheet.append(headers)

    for cell in time_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i in range(TARGET_LENGTH):
        time_sheet.append(
            [
                i + 1,
                float(time_s[i]),
                float(rotor_speed_rpm[i]),
                float(omega[i]),
            ]
        )

    time_sheet.freeze_panes = "A2"

    column_widths = {
        "A": 14,
        "B": 14,
        "C": 20,
        "D": 18,
    }

    for column_letter, width in column_widths.items():
        time_sheet.column_dimensions[column_letter].width = width

    for row in time_sheet.iter_rows(
        min_row=2,
        max_row=TARGET_LENGTH + 1,
        min_col=2,
        max_col=4,
    ):
        for cell in row:
            cell.number_format = "0.000000"

    # Line chart of omega over time.
    chart = LineChart()
    chart.title = "Omega over time"
    chart.y_axis.title = "Omega [rad/s]"
    chart.x_axis.title = "Time [s]"
    chart.height = 12
    chart.width = 24

    y_values = Reference(
        time_sheet,
        min_col=4,
        min_row=1,
        max_row=TARGET_LENGTH + 1,
    )
    x_values = Reference(
        time_sheet,
        min_col=2,
        min_row=2,
        max_row=TARGET_LENGTH + 1,
    )

    chart.add_data(y_values, titles_from_data=True)
    chart.set_categories(x_values)

    time_sheet.add_chart(chart, "F2")

    # --------------------------------------------------------
    # Summary sheet
    # --------------------------------------------------------
    omega_min = finite_stat(omega, np.min)
    omega_max = finite_stat(omega, np.max)
    omega_median = finite_stat(omega, np.median)
    omega_mean = finite_stat(omega, np.mean)
    omega_fourier_dc = fourier_dc(omega)

    summary_rows = [
        ["mat_file", mat_file_name],
        ["source", "Data.ModelsData.Model1.RotorSpeed"],
        ["raw_rotor_speed_sample_count", data["rotor_speed_raw_count"]],
        ["processed_sample_count", TARGET_LENGTH],
        ["measurement_duration_s", MEASUREMENT_DURATION_S],
        ["formula", "omega_rad_s = rotor_speed_rpm * 2*pi/60"],
        ["omega_min_rad_s", omega_min],
        ["omega_max_rad_s", omega_max],
        ["omega_median_rad_s", omega_median],
        ["omega_mean_rad_s", omega_mean],
        ["omega_fourier_dc_rad_s", omega_fourier_dc],
    ]

    summary_sheet.append(["quantity", "value"])

    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in summary_rows:
        summary_sheet.append(row)

    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 60

    for row in summary_sheet.iter_rows(
        min_row=2,
        max_row=summary_sheet.max_row,
        min_col=2,
        max_col=2,
    ):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"

    workbook.save(output_path)

    return output_path


def main() -> None:
    if len(sys.argv) >= 2:
        mat_file_name = sys.argv[1]
    else:
        mat_file_name = DEFAULT_MAT_FILE_NAME

    mat_path = DATA_DIR / mat_file_name

    if not mat_path.exists():
        raise FileNotFoundError(
            f"MAT file not found: {mat_path}\n"
            "You can pass a file name explicitly, for example:\n"
            "py .\\wtt\\create_single_mat_omega_excel.py Results_File_5_TOP_1.mat"
        )

    data = compute_omega_time_series(mat_path)
    output_path = write_excel(mat_file_name, data)

    omega = data["omega_rad_s"]

    print("Finished.")
    print(f"MAT file: {mat_file_name}")
    print(f"Excel file: {output_path.resolve()}")
    print(f"Omega formula: omega = RotorSpeed * 2*pi/60")
    print(f"Number of omega values: {omega.size}")
    print(f"Omega min [rad/s]: {finite_stat(omega, np.min):.6f}")
    print(f"Omega max [rad/s]: {finite_stat(omega, np.max):.6f}")
    print(f"Omega median [rad/s]: {finite_stat(omega, np.median):.6f}")


if __name__ == "__main__":
    main()
