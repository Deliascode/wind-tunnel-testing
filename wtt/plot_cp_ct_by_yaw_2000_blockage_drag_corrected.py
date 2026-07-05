from __future__ import annotations

from pathlib import Path
import math
import csv

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat


# =========================
# User settings / constants
# =========================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct_2000_blockage_drag_corrected"
TEST_MATRIX_PATH = DATA_DIR / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"

# If your report uses the exact lookup-table value, change this to 0.5436 m.
ROTOR_RADIUS_M = 0.55
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2

# Hub.Torque in the MAT files appears to be in mNm.
TORQUE_SCALE_TO_NM = 1e-3

# Geometry for the moment balance.
# The user-provided hub height is used as the thrust/nacelle moment arm.
HUB_HEIGHT_M = 0.80
THRUST_LEVER_ARM_M = HUB_HEIGHT_M
NACELLE_HUB_MOMENT_ARM_M = HUB_HEIGHT_M

# Tower geometry.
TOWER_HEIGHT_M = HUB_HEIGHT_M
TOWER_DIAMETER_M = 0.047

# Nacelle-hub drag parameter from the slides:
# D_nacelle-hub = 0.5 * rho * V^2 * SCD
NACELLE_HUB_SCD_M2 = 0.0175

# Assumption required by the slides for the tower drag.
# A value around 1.2 is a standard first estimate for a circular cylinder.
TOWER_DRAG_COEFFICIENT = 1.2

# Wind-tunnel dimensions for blockage ratio alpha = A_rotor / A_tunnel.
TUNNEL_WIDTH_M = 2.7
TUNNEL_HEIGHT_M = 1.8
TUNNEL_AREA_M2 = TUNNEL_WIDTH_M * TUNNEL_HEIGHT_M
BLOCKAGE_RATIO_ALPHA = ROTOR_AREA_M2 / TUNNEL_AREA_M2

# Most measurement channels have 2000 samples. Hub.Torque has 20000 samples
# and is block-averaged down to 2000 samples. AirDensity has 1000 samples and
# is interpolated up to 2000 samples.
TARGET_LENGTH = 2000
MEASUREMENT_DURATION_S = 8.0

CORRECTION_MAX_ITERATIONS = 50
CORRECTION_TOLERANCE = 1e-8


# =========================
# Helper functions
# =========================

def unwrap_singleton(value):
    current = value
    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]
    return current


def load_model1(mat_path: Path):
    mat_data = loadmat(mat_path, squeeze_me=True, struct_as_record=False)
    data = unwrap_singleton(mat_data["Data"])
    models_data = unwrap_singleton(data.ModelsData)
    model1 = unwrap_singleton(models_data.Model1)
    return model1


def get_array(value) -> np.ndarray:
    array = np.asarray(unwrap_singleton(value), dtype=float).ravel(order="C")

    if array.size == 0:
        return np.array([np.nan], dtype=float)

    return array


def block_average_to_target_length(
    values: np.ndarray,
    target_length: int,
) -> np.ndarray:
    """Downsamples by forming block means.

    Example:
    Hub.Torque has 20000 samples. For target_length=2000 this forms
    2000 blocks with 10 samples each and returns one mean value per block.
    """

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size % target_length != 0:
        raise ValueError(
            "Block averaging requires the signal length to be an integer "
            "multiple of target_length."
        )

    block_size = values.size // target_length
    blocks = values.reshape(target_length, block_size)

    with np.errstate(invalid="ignore"):
        return np.nanmean(blocks, axis=1)


def interpolate_to_target_length(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    """Interpolates a signal over the shared measurement duration."""

    values = np.asarray(values, dtype=float).ravel(order="C")

    if target_length <= 0:
        raise ValueError("target_length must be greater than zero.")

    if values.size == 0:
        return np.full(target_length, np.nan, dtype=float)

    if values.size == 1:
        return np.full(target_length, values[0], dtype=float)

    if values.size == target_length:
        return values.astype(float, copy=False)

    finite_mask = np.isfinite(values)

    if finite_mask.sum() == 0:
        return np.full(target_length, np.nan, dtype=float)

    if finite_mask.sum() == 1:
        return np.full(target_length, values[finite_mask][0], dtype=float)

    source_time = np.linspace(
        0.0,
        duration_s,
        values.size,
        endpoint=False,
    )
    target_time = np.linspace(
        0.0,
        duration_s,
        target_length,
        endpoint=False,
    )

    return np.interp(
        target_time,
        source_time[finite_mask],
        values[finite_mask],
    )


def resample_to_2000(values: np.ndarray, name: str) -> np.ndarray:
    """Resamples all signals to 2000 samples.

    - Signals with exactly 2000 samples are kept unchanged.
    - Signals with more than 2000 samples and an integer ratio are block-averaged.
      This is the intended path for Hub.Torque: 20000 -> 2000.
    - Signals with fewer than 2000 samples are interpolated.
      This is the intended path for AirDensity: 1000 -> 2000.
    - Any other non-integer case falls back to interpolation.
    """

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size == TARGET_LENGTH:
        return values.astype(float, copy=False)

    if values.size > TARGET_LENGTH and values.size % TARGET_LENGTH == 0:
        return block_average_to_target_length(values, TARGET_LENGTH)

    return interpolate_to_target_length(values, TARGET_LENGTH)


def safe_stat(values: np.ndarray, fn):
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]

    if values.size == 0:
        return math.nan

    return float(fn(values))


def axial_induction_from_ct(ct: np.ndarray) -> np.ndarray:
    """Momentum-theory estimate for axial induction.

    CT = 4a(1-a) gives a = 0.5 * (1 - sqrt(1 - CT)).
    Values are clipped to the classical range to avoid complex numbers.
    """

    ct = np.asarray(ct, dtype=float)
    ct_limited = np.clip(ct, 0.0, 0.999999)

    return 0.5 * (1.0 - np.sqrt(1.0 - ct_limited))


def glauert_velocity_from_thrust(
    u_inf: np.ndarray,
    air_density: np.ndarray,
    rotor_thrust_n: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Applies Glauert blockage correction to the time series.

    Formula:
        U'_inf / U_inf = 1 + (alpha / 4) * C_T / (1 - C_T)

    with:
        C_T = T / (0.5 * rho * A_R * U_inf^2)

    C_T is the rotor thrust coefficient, not the torque coefficient.
    """

    u_inf = np.asarray(u_inf, dtype=float)
    air_density = np.asarray(air_density, dtype=float)
    rotor_thrust_n = np.asarray(rotor_thrust_n, dtype=float)

    ct_denominator_uncorrected = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * u_inf**2
    )

    ct_for_blockage = np.divide(
        rotor_thrust_n,
        ct_denominator_uncorrected,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(rotor_thrust_n)
            & np.isfinite(ct_denominator_uncorrected)
            & (ct_denominator_uncorrected > 0.0)
        ),
    )

    blockage_factor = np.full(TARGET_LENGTH, np.nan, dtype=float)

    valid = (
        np.isfinite(u_inf)
        & np.isfinite(ct_for_blockage)
        & (u_inf > 0.0)
        & (ct_for_blockage < 1.0)
        & np.isfinite(1.0 - ct_for_blockage)
        & ((1.0 - ct_for_blockage) != 0.0)
    )

    blockage_factor[valid] = (
        1.0
        + (BLOCKAGE_RATIO_ALPHA / 4.0)
        * ct_for_blockage[valid]
        / (1.0 - ct_for_blockage[valid])
    )

    u_blockage_corrected = u_inf * blockage_factor

    return u_blockage_corrected, ct_for_blockage, blockage_factor


def drag_corrected_rotor_thrust_and_velocity(
    tower_fa_moment: np.ndarray,
    pitot_velocity: np.ndarray,
    air_density: np.ndarray,
) -> dict[str, np.ndarray]:
    """Computes rotor thrust by removing nacelle-hub and tower drag moments.

    Measured moment balance about the tower-base load measurement location:

        -M_FA = T_rotor * L
                + D_nacelle * z_hub
                + M_tower_drag

    Therefore:

        T_rotor = (-M_FA - D_nacelle*z_hub - M_tower_drag) / L

    The nacelle-hub drag is:
        D_nacelle = 0.5 * rho * U^2 * SCD

    The tower drag force is approximated from the slide integral:
        D_T = 0.5 * rho * D * integral_0^H V(h)^2 * C_D(h,Re) dh

    For the tower-base bending moment, the corresponding moment integral is:
        M_T = 0.5 * rho * D * integral_0^H h * V(h)^2 * C_D(h,Re) dh

    Assumption:
    C_D is constant and V(h) is approximated by a uniform downstream velocity
    V_tower = U'_inf * (1 - 2a), where a is estimated from CT.
    """

    tower_fa_moment = np.asarray(tower_fa_moment, dtype=float)
    pitot_velocity = np.asarray(pitot_velocity, dtype=float)
    air_density = np.asarray(air_density, dtype=float)

    measured_positive_moment = -tower_fa_moment

    # Initial estimate: all measured moment is due to rotor thrust.
    rotor_thrust_n = measured_positive_moment / THRUST_LEVER_ARM_M

    u_blockage_corrected = pitot_velocity.copy()
    ct_for_blockage = np.full(TARGET_LENGTH, np.nan, dtype=float)
    blockage_factor = np.full(TARGET_LENGTH, np.nan, dtype=float)

    nacelle_drag_n = np.zeros(TARGET_LENGTH, dtype=float)
    nacelle_drag_moment_nm = np.zeros(TARGET_LENGTH, dtype=float)

    tower_drag_n = np.zeros(TARGET_LENGTH, dtype=float)
    tower_drag_moment_nm = np.zeros(TARGET_LENGTH, dtype=float)
    tower_velocity = np.zeros(TARGET_LENGTH, dtype=float)
    axial_induction = np.zeros(TARGET_LENGTH, dtype=float)

    for _ in range(CORRECTION_MAX_ITERATIONS):
        previous_rotor_thrust_n = rotor_thrust_n.copy()

        (
            u_blockage_corrected,
            ct_for_blockage,
            blockage_factor,
        ) = glauert_velocity_from_thrust(
            u_inf=pitot_velocity,
            air_density=air_density,
            rotor_thrust_n=rotor_thrust_n,
        )

        ct_corrected_denominator = (
            0.5
            * air_density
            * ROTOR_AREA_M2
            * u_blockage_corrected**2
        )

        ct_corrected = np.divide(
            rotor_thrust_n,
            ct_corrected_denominator,
            out=np.full(TARGET_LENGTH, np.nan, dtype=float),
            where=(
                np.isfinite(rotor_thrust_n)
                & np.isfinite(ct_corrected_denominator)
                & (ct_corrected_denominator > 0.0)
            ),
        )

        axial_induction = axial_induction_from_ct(ct_corrected)

        # Nacelle-hub drag at hub height.
        nacelle_drag_n = (
            0.5
            * air_density
            * u_blockage_corrected**2
            * NACELLE_HUB_SCD_M2
        )

        nacelle_drag_moment_nm = (
            nacelle_drag_n
            * NACELLE_HUB_MOMENT_ARM_M
        )

        # Tower drag in the rotor wake.
        tower_velocity = u_blockage_corrected * (1.0 - 2.0 * axial_induction)
        tower_velocity = np.maximum(tower_velocity, 0.0)

        # Integral approximation with constant C_D and constant V(h):
        # D_T = 0.5*rho*D*V^2*C_D*H
        tower_drag_n = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * TOWER_HEIGHT_M
        )

        # Moment integral:
        # M_T = 0.5*rho*D*V^2*C_D*H^2/2
        tower_drag_moment_nm = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * (TOWER_HEIGHT_M**2 / 2.0)
        )

        rotor_moment_nm = (
            measured_positive_moment
            - nacelle_drag_moment_nm
            - tower_drag_moment_nm
        )

        rotor_thrust_n = np.divide(
            rotor_moment_nm,
            THRUST_LEVER_ARM_M,
            out=np.full(TARGET_LENGTH, np.nan, dtype=float),
            where=np.isfinite(rotor_moment_nm),
        )

        # Negative corrected rotor thrust is non-physical here.
        rotor_thrust_n = np.where(
            np.isfinite(rotor_thrust_n),
            np.maximum(rotor_thrust_n, 0.0),
            np.nan,
        )

        difference = np.abs(rotor_thrust_n - previous_rotor_thrust_n)
        finite_difference = difference[np.isfinite(difference)]

        if finite_difference.size == 0:
            break

        if np.max(finite_difference) < CORRECTION_TOLERANCE:
            break

    return {
        "measured_positive_moment": measured_positive_moment,
        "rotor_thrust_n": rotor_thrust_n,
        "u_blockage_corrected": u_blockage_corrected,
        "ct_for_blockage": ct_for_blockage,
        "blockage_factor": blockage_factor,

        "nacelle_drag_n": nacelle_drag_n,
        "nacelle_drag_moment_nm": nacelle_drag_moment_nm,

        "tower_drag_n": tower_drag_n,
        "tower_drag_moment_nm": tower_drag_moment_nm,
        "tower_velocity": tower_velocity,
        "axial_induction": axial_induction,
    }


# =========================
# Main calculation
# =========================

def compute_test_statistics(mat_path: Path) -> dict:
    """Computes TSR, cP and cT statistics on a 2000-sample time axis."""

    model1 = load_model1(mat_path)

    rotor_speed_rpm = get_array(model1.RotorSpeed)
    pitot_velocity = get_array(model1.PitotVelocity)
    air_density = get_array(model1.AirDensity)
    hub_torque_raw = get_array(model1.Hub.Torque)
    tower_fa_moment = get_array(model1.Tower.FA)
    yaw_actual = get_array(model1.Yaw)

    # Fixed common target length: 2000 samples.
    # Hub.Torque 20000 -> 2000 by block averaging.
    # AirDensity 1000 -> 2000 by interpolation.
    rotor_speed_common = resample_to_2000(rotor_speed_rpm, "RotorSpeed")
    pitot_velocity_common = resample_to_2000(pitot_velocity, "PitotVelocity")
    air_density_common = resample_to_2000(air_density, "AirDensity")
    torque_common_raw = resample_to_2000(hub_torque_raw, "Hub.Torque")
    tower_fa_common = resample_to_2000(tower_fa_moment, "Tower.FA")
    yaw_common = resample_to_2000(yaw_actual, "Yaw")

    omega = (
        rotor_speed_common
        * 2.0
        * math.pi
        / 60.0
    )

    torque_nm = (
        torque_common_raw
        * TORQUE_SCALE_TO_NM
    )

    power_w = torque_nm * omega

    correction = drag_corrected_rotor_thrust_and_velocity(
        tower_fa_moment=tower_fa_common,
        pitot_velocity=pitot_velocity_common,
        air_density=air_density_common,
    )

    thrust_n_raw_from_tower_fa = (
        -tower_fa_common
        / THRUST_LEVER_ARM_M
    )

    thrust_n = correction["rotor_thrust_n"]
    u_blockage_corrected = correction["u_blockage_corrected"]

    # From this point onward, only u_blockage_corrected is used for TSR, cP and cT.
    # The thrust used in cT and in the blockage correction is the rotor thrust
    # after nacelle-hub and tower-drag moment removal.
    cp_denominator = (
        0.5
        * air_density_common
        * ROTOR_AREA_M2
        * u_blockage_corrected**3
    )

    cp_series = np.divide(
        power_w,
        cp_denominator,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(cp_denominator)
            & (cp_denominator != 0.0)
        ),
    )

    tsr_series = np.divide(
        omega * ROTOR_RADIUS_M,
        u_blockage_corrected,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(u_blockage_corrected)
            & (u_blockage_corrected != 0.0)
        ),
    )

    ct_denominator_corrected = (
        0.5
        * air_density_common
        * ROTOR_AREA_M2
        * u_blockage_corrected**2
    )

    ct_series = np.divide(
        thrust_n,
        ct_denominator_corrected,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(ct_denominator_corrected)
            & (ct_denominator_corrected != 0.0)
        ),
    )

    return {
        "yaw_actual_median": safe_stat(
            yaw_common,
            np.median,
        ),

        "pitot_velocity_median": safe_stat(
            pitot_velocity_common,
            np.median,
        ),
        "u_blockage_corrected_median": safe_stat(
            u_blockage_corrected,
            np.median,
        ),
        "ct_for_blockage_median": safe_stat(
            correction["ct_for_blockage"],
            np.median,
        ),
        "blockage_factor_median": safe_stat(
            correction["blockage_factor"],
            np.median,
        ),

        "thrust_raw_from_tower_fa_median": safe_stat(
            thrust_n_raw_from_tower_fa,
            np.median,
        ),
        "rotor_thrust_drag_corrected_median": safe_stat(
            thrust_n,
            np.median,
        ),

        "nacelle_drag_median": safe_stat(
            correction["nacelle_drag_n"],
            np.median,
        ),
        "nacelle_drag_moment_median": safe_stat(
            correction["nacelle_drag_moment_nm"],
            np.median,
        ),

        "tower_drag_median": safe_stat(
            correction["tower_drag_n"],
            np.median,
        ),
        "tower_drag_moment_median": safe_stat(
            correction["tower_drag_moment_nm"],
            np.median,
        ),

        "tower_velocity_median": safe_stat(
            correction["tower_velocity"],
            np.median,
        ),
        "axial_induction_median": safe_stat(
            correction["axial_induction"],
            np.median,
        ),

        "tsr_median": safe_stat(
            tsr_series,
            np.median,
        ),
        "tsr_std": safe_stat(
            tsr_series,
            np.std,
        ),

        "cp_median": safe_stat(
            cp_series,
            np.median,
        ),
        "cp_std": safe_stat(
            cp_series,
            np.std,
        ),
        "cp_min": safe_stat(
            cp_series,
            np.min,
        ),
        "cp_max": safe_stat(
            cp_series,
            np.max,
        ),

        "ct_median": safe_stat(
            ct_series,
            np.median,
        ),
        "ct_std": safe_stat(
            ct_series,
            np.std,
        ),
        "ct_min": safe_stat(
            ct_series,
            np.min,
        ),
        "ct_max": safe_stat(
            ct_series,
            np.max,
        ),
    }


def load_test_matrix(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Task1 Performance G2"]

    headers = [ws.cell(row=1, column=i).value for i in range(1, 31)]

    records = []

    for row_idx in range(2, ws.max_row + 1):
        test_number = ws.cell(row=row_idx, column=1).value

        if test_number is None:
            continue

        record = {}

        for col_idx, header in enumerate(headers, start=1):
            record[header] = ws.cell(row=row_idx, column=col_idx).value

        records.append(record)

    return records


def build_dataframe_like_records() -> list[dict]:
    records = load_test_matrix(TEST_MATRIX_PATH)
    rows = []

    for record in records:
        top = record["Top"]
        file_number = record["File number"]

        if top is None or file_number is None:
            continue

        mat_name = f"Results_File_{int(file_number)}_TOP_{int(top)}.mat"
        mat_path = DATA_DIR / mat_name

        if not mat_path.exists():
            print(f"Datei fehlt, übersprungen: {mat_name}")
            continue

        stats = compute_test_statistics(mat_path)

        rows.append(
            {
                "test": record["Test number"],
                "yaw_target": record["Yaw [deg]"],
                "yaw_actual_median": stats["yaw_actual_median"],

                "tsr_expected": record["TSR [-]"],
                "tsr_median": stats["tsr_median"],
                "tsr_std": stats["tsr_std"],

                "pitot_velocity_median": stats["pitot_velocity_median"],
                "u_blockage_corrected_median": stats[
                    "u_blockage_corrected_median"
                ],
                "ct_for_blockage_median": stats["ct_for_blockage_median"],
                "blockage_factor_median": stats["blockage_factor_median"],

                "thrust_raw_from_tower_fa_median": stats[
                    "thrust_raw_from_tower_fa_median"
                ],
                "rotor_thrust_drag_corrected_median": stats[
                    "rotor_thrust_drag_corrected_median"
                ],

                "nacelle_drag_median": stats["nacelle_drag_median"],
                "nacelle_drag_moment_median": stats[
                    "nacelle_drag_moment_median"
                ],

                "tower_drag_median": stats["tower_drag_median"],
                "tower_drag_moment_median": stats[
                    "tower_drag_moment_median"
                ],

                "tower_velocity_median": stats["tower_velocity_median"],
                "axial_induction_median": stats["axial_induction_median"],

                "cp_median": stats["cp_median"],
                "cp_plus_std": (
                    stats["cp_median"] + stats["cp_std"]
                    if np.isfinite(stats["cp_median"])
                    and np.isfinite(stats["cp_std"])
                    else math.nan
                ),
                "cp_minus_std": (
                    stats["cp_median"] - stats["cp_std"]
                    if np.isfinite(stats["cp_median"])
                    and np.isfinite(stats["cp_std"])
                    else math.nan
                ),
                "cp_min": stats["cp_min"],
                "cp_max": stats["cp_max"],

                "ct_median": stats["ct_median"],
                "ct_plus_std": (
                    stats["ct_median"] + stats["ct_std"]
                    if np.isfinite(stats["ct_median"])
                    and np.isfinite(stats["ct_std"])
                    else math.nan
                ),
                "ct_minus_std": (
                    stats["ct_median"] - stats["ct_std"]
                    if np.isfinite(stats["ct_median"])
                    and np.isfinite(stats["ct_std"])
                    else math.nan
                ),
                "ct_min": stats["ct_min"],
                "ct_max": stats["ct_max"],

                "mat_file": mat_name,
            }
        )

    return rows


def plot_for_yaw(
    rows: list[dict],
    yaw_value: float,
    coefficient_prefix: str,
):
    """
    coefficient_prefix: 'cp' or 'ct'

    Produces one plot with:
    median, median+std, median-std, min, max
    """

    yaw_rows = [r for r in rows if r["yaw_target"] == yaw_value]

    yaw_rows.sort(
        key=lambda r: (
            r["tsr_median"]
            if r["tsr_median"] is not None
            else math.inf
        )
    )

    x = [r["tsr_median"] for r in yaw_rows]
    y_median = [r[f"{coefficient_prefix}_median"] for r in yaw_rows]
    y_plus_std = [r[f"{coefficient_prefix}_plus_std"] for r in yaw_rows]
    y_minus_std = [r[f"{coefficient_prefix}_minus_std"] for r in yaw_rows]
    y_min = [r[f"{coefficient_prefix}_min"] for r in yaw_rows]
    y_max = [r[f"{coefficient_prefix}_max"] for r in yaw_rows]

    coeff_label = "cP" if coefficient_prefix == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(x, y_median, marker="o", label="Median")
    plt.plot(
        x,
        y_plus_std,
        marker="o",
        linestyle="--",
        label="Median + Standardabweichung",
    )
    plt.plot(
        x,
        y_minus_std,
        marker="o",
        linestyle="--",
        label="Median - Standardabweichung",
    )
    plt.plot(
        x,
        y_min,
        marker="o",
        linestyle=":",
        label="Minimum",
    )
    plt.plot(
        x,
        y_max,
        marker="o",
        linestyle=":",
        label="Maximum",
    )

    plt.xlabel("TSR mit Glauert-korrigierter Geschwindigkeit")
    plt.ylabel(coeff_label)
    plt.title(f"{coeff_label} über TSR bei Yaw = {yaw_value}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    out_name = (
        f"{coeff_label}_yaw_{int(yaw_value):+d}.png"
        .replace("+", "plus")
        .replace("-", "minus")
    )

    plt.savefig(OUTPUT_DIR / out_name, dpi=200)
    plt.close()


def save_summary_csv(rows: list[dict]):
    out_csv = OUTPUT_DIR / "cp_ct_summary_by_test.csv"

    fieldnames = [
        "test",
        "yaw_target",
        "yaw_actual_median",

        "tsr_expected",
        "tsr_median",
        "tsr_std",

        "pitot_velocity_median",
        "u_blockage_corrected_median",
        "ct_for_blockage_median",
        "blockage_factor_median",

        "thrust_raw_from_tower_fa_median",
        "rotor_thrust_drag_corrected_median",
        "nacelle_drag_median",
        "nacelle_drag_moment_median",
        "tower_drag_median",
        "tower_drag_moment_median",
        "tower_velocity_median",
        "axial_induction_median",

        "cp_median",
        "cp_plus_std",
        "cp_minus_std",
        "cp_min",
        "cp_max",

        "ct_median",
        "ct_plus_std",
        "ct_minus_std",
        "ct_min",
        "ct_max",

        "mat_file",
    ]

    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(
            f"Excel-Datei nicht gefunden: {TEST_MATRIX_PATH}"
        )

    if not DATA_DIR.exists():
        raise FileNotFoundError(
            f"data-Ordner nicht gefunden: {DATA_DIR}"
        )

    rows = build_dataframe_like_records()

    if not rows:
        raise RuntimeError("Es konnten keine Testdaten verarbeitet werden.")

    save_summary_csv(rows)

    yaw_values = sorted(
        set(
            r["yaw_target"]
            for r in rows
            if r["yaw_target"] is not None
        )
    )

    for yaw in yaw_values:
        plot_for_yaw(rows, yaw, "cp")
        plot_for_yaw(rows, yaw, "ct")

    print("Fertig.")
    print(f"Ausgabeordner: {OUTPUT_DIR.resolve()}")
    print("Erzeugte Dateien:")

    for path in sorted(OUTPUT_DIR.iterdir()):
        print(" -", path.name)

    print("\nWICHTIG:")
    print("Alle Signale werden auf 2000 Werte gebracht.")
    print("Hub.Torque wird von 20000 auf 2000 Werte blockweise gemittelt.")
    print("AirDensity wird von 1000 auf 2000 Werte interpoliert.")
    print("PitotVelocity wird nur für die Glauert-Korrektur verwendet.")
    print("Nacelle-hub drag and tower drag moments are removed from Tower.FA.")
    print("Danach werden TSR, cP und cT mit u_blockage_corrected berechnet.")
    print(f"Hub height / thrust arm: {HUB_HEIGHT_M:.3f} m")
    print(f"Nacelle-hub SCD: {NACELLE_HUB_SCD_M2:.5f} m²")
    print(f"Tower diameter: {TOWER_DIAMETER_M:.3f} m")
    print(f"Tower drag coefficient assumption: {TOWER_DRAG_COEFFICIENT:.2f}")
    print(f"Blockage ratio alpha: {BLOCKAGE_RATIO_ALPHA:.5f}")
    print(f"Windkanalfläche: {TUNNEL_AREA_M2:.3f} m²")


if __name__ == "__main__":
    main()
