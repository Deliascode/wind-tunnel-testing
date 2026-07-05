from __future__ import annotations

from pathlib import Path
import math
import csv

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat
from scipy.signal import welch


# ============================================================
# Settings
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
TEST_MATRIX_PATH = DATA_DIR / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct_welch_grouped_by_yaw_zero_glauert"

MEASUREMENT_DURATION_S = 8.0
TARGET_LENGTH = 2000
FS_HZ = TARGET_LENGTH / MEASUREMENT_DURATION_S

ROTOR_RADIUS_M = 0.5436
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2
TORQUE_SCALE_TO_NM = 1e-3

HUB_HEIGHT_M = 0.80
THRUST_LEVER_ARM_M = HUB_HEIGHT_M
NACELLE_HUB_MOMENT_ARM_M = HUB_HEIGHT_M
TOWER_HEIGHT_M = HUB_HEIGHT_M
TOWER_DIAMETER_M = 0.047
NACELLE_HUB_SCD_M2 = 0.0175
TOWER_DRAG_COEFFICIENT = 1.2

TUNNEL_WIDTH_M = 2.7
TUNNEL_HEIGHT_M = 1.8
TUNNEL_AREA_M2 = TUNNEL_WIDTH_M * TUNNEL_HEIGHT_M
BLOCKAGE_RATIO_ALPHA = ROTOR_AREA_M2 / TUNNEL_AREA_M2

CORRECTION_MAX_ITERATIONS = 50
CORRECTION_TOLERANCE = 1e-8

WELCH_NPERSEG = 512
WELCH_NOOVERLAP = WELCH_NPERSEG // 2


# ============================================================
# Basic helpers
# ============================================================

def unwrap_singleton(value):
    current = value
    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]
    return current


def load_model1(mat_path: Path):
    mat_data = loadmat(mat_path, squeeze_me=True, struct_as_record=False)
    data = unwrap_singleton(mat_data["Data"])
    models_data = unwrap_singleton(data.ModelsData)
    return unwrap_singleton(models_data.Model1)


def get_array(value) -> np.ndarray:
    arr = np.asarray(unwrap_singleton(value), dtype=float).ravel(order="C")
    if arr.size == 0:
        return np.array([np.nan], dtype=float)
    return arr


def nested(root, *names: str) -> np.ndarray:
    current = root
    for name in names:
        current = unwrap_singleton(current)
        if not hasattr(current, name):
            return np.array([np.nan], dtype=float)
        current = getattr(current, name)
    return get_array(current)


def block_average(values: np.ndarray, target_length: int) -> np.ndarray:
    values = np.asarray(values, dtype=float).ravel(order="C")
    if values.size % target_length != 0:
        raise ValueError("Block averaging requires an integer length ratio.")
    block_size = values.size // target_length
    with np.errstate(invalid="ignore"):
        return np.nanmean(values.reshape(target_length, block_size), axis=1)


def interpolate_to_target(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    values = np.asarray(values, dtype=float).ravel(order="C")
    if values.size == 0:
        return np.full(target_length, np.nan)
    if values.size == 1:
        return np.full(target_length, values[0])
    if values.size == target_length:
        return values.astype(float, copy=False)

    finite = np.isfinite(values)
    if finite.sum() == 0:
        return np.full(target_length, np.nan)
    if finite.sum() == 1:
        return np.full(target_length, values[finite][0])

    t_src = np.linspace(0.0, duration_s, values.size, endpoint=False)
    t_tar = np.linspace(0.0, duration_s, target_length, endpoint=False)
    return np.interp(t_tar, t_src[finite], values[finite])


def resample_to_2000(values: np.ndarray, name: str) -> np.ndarray:
    """Yaw-case-local processing: 2000 stay unchanged, 20000 -> block average, 1000 -> interpolation."""
    values = np.asarray(values, dtype=float).ravel(order="C")
    if values.size == TARGET_LENGTH:
        return values.astype(float, copy=False)
    if values.size > TARGET_LENGTH and values.size % TARGET_LENGTH == 0:
        return block_average(values, TARGET_LENGTH)
    return interpolate_to_target(values, TARGET_LENGTH)


def finite_values(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    return values[np.isfinite(values)]


def stat(values: np.ndarray, fn) -> float:
    clean = finite_values(values)
    if clean.size == 0:
        return math.nan
    return float(fn(clean))


def fill_nonfinite(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float).copy()
    finite = np.isfinite(values)
    if finite.sum() == 0:
        return np.full_like(values, np.nan)
    if finite.sum() == 1:
        return np.full_like(values, values[finite][0])
    if not finite.all():
        idx = np.arange(values.size)
        values[~finite] = np.interp(idx[~finite], idx[finite], values[finite])
    return values


def welch_dc_estimate(values: np.ndarray) -> float:
    """Robust Welch-based scalar replacing the former raw FFT-DC value."""
    values = fill_nonfinite(values)
    if values.size == 0 or not np.isfinite(values).any():
        return math.nan

    nperseg = min(WELCH_NPERSEG, values.size)
    noverlap = min(WELCH_NOOVERLAP, max(0, nperseg - 1))

    try:
        _, pxx = welch(
            values,
            fs=FS_HZ,
            window="boxcar",
            nperseg=nperseg,
            noverlap=noverlap,
            detrend=False,
            scaling="spectrum",
            average="median",
        )
    except TypeError:
        _, pxx = welch(
            values,
            fs=FS_HZ,
            window="boxcar",
            nperseg=nperseg,
            noverlap=noverlap,
            detrend=False,
            scaling="spectrum",
        )

    if pxx.size == 0 or not np.isfinite(pxx[0]):
        return math.nan

    magnitude = float(np.sqrt(max(pxx[0], 0.0)))
    sign_ref = stat(values, np.median)
    if not np.isfinite(sign_ref) or sign_ref == 0.0:
        sign_ref = stat(values, np.mean)
    sign = 1.0 if sign_ref >= 0.0 else -1.0
    return sign * magnitude


def welch_spectrum(values: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    values = fill_nonfinite(values)
    if values.size == 0 or not np.isfinite(values).any():
        return np.array([]), np.array([])

    nperseg = min(WELCH_NPERSEG, values.size)
    noverlap = min(WELCH_NOOVERLAP, max(0, nperseg - 1))

    try:
        f, pxx = welch(
            values,
            fs=FS_HZ,
            window="boxcar",
            nperseg=nperseg,
            noverlap=noverlap,
            detrend=False,
            scaling="spectrum",
            average="median",
        )
    except TypeError:
        f, pxx = welch(
            values,
            fs=FS_HZ,
            window="boxcar",
            nperseg=nperseg,
            noverlap=noverlap,
            detrend=False,
            scaling="spectrum",
        )
    return f, pxx


# ============================================================
# Glauert and drag correction
# ============================================================

def axial_induction_from_ct(ct: np.ndarray) -> np.ndarray:
    ct_limited = np.clip(np.asarray(ct, dtype=float), 0.0, 0.999999)
    return 0.5 * (1.0 - np.sqrt(1.0 - ct_limited))


def glauert_velocity_from_zero_yaw_thrust(
    u_tunnel: np.ndarray,
    air_density: np.ndarray,
    rotor_thrust_n: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Glauert correction using tunnel velocity, not corrected velocity."""
    denom = 0.5 * air_density * ROTOR_AREA_M2 * u_tunnel**2
    ct_glauert = np.divide(
        rotor_thrust_n,
        denom,
        out=np.full(TARGET_LENGTH, np.nan),
        where=np.isfinite(rotor_thrust_n) & np.isfinite(denom) & (denom > 0.0),
    )

    factor = np.full(TARGET_LENGTH, np.nan)
    valid = (
        np.isfinite(u_tunnel)
        & np.isfinite(ct_glauert)
        & (u_tunnel > 0.0)
        & (ct_glauert < 1.0)
        & np.isfinite(1.0 - ct_glauert)
        & ((1.0 - ct_glauert) != 0.0)
    )
    factor[valid] = (
        1.0
        + (BLOCKAGE_RATIO_ALPHA / 4.0)
        * ct_glauert[valid]
        / (1.0 - ct_glauert[valid])
    )
    u_glauert = u_tunnel * factor
    return u_glauert, ct_glauert, factor


def drag_corrected_zero_yaw(
    tower_fa_moment: np.ndarray,
    pitot_velocity: np.ndarray,
    air_density: np.ndarray,
) -> dict[str, np.ndarray]:
    """Self-consistent drag and Glauert calculation for 0-degree yaw references."""
    measured_positive_moment = -tower_fa_moment
    thrust = measured_positive_moment / THRUST_LEVER_ARM_M

    u_glauert = pitot_velocity.copy()
    ct_glauert = np.full(TARGET_LENGTH, np.nan)
    factor = np.full(TARGET_LENGTH, np.nan)

    nacelle_drag = np.zeros(TARGET_LENGTH)
    nacelle_moment = np.zeros(TARGET_LENGTH)
    tower_drag = np.zeros(TARGET_LENGTH)
    tower_moment = np.zeros(TARGET_LENGTH)
    tower_velocity = np.zeros(TARGET_LENGTH)
    axial_induction = np.zeros(TARGET_LENGTH)

    for _ in range(CORRECTION_MAX_ITERATIONS):
        old = thrust.copy()
        u_glauert, ct_glauert, factor = glauert_velocity_from_zero_yaw_thrust(
            u_tunnel=pitot_velocity,
            air_density=air_density,
            rotor_thrust_n=thrust,
        )

        ct_denom = 0.5 * air_density * ROTOR_AREA_M2 * u_glauert**2
        ct_corr = np.divide(
            thrust,
            ct_denom,
            out=np.full(TARGET_LENGTH, np.nan),
            where=np.isfinite(thrust) & np.isfinite(ct_denom) & (ct_denom > 0.0),
        )
        axial_induction = axial_induction_from_ct(ct_corr)

        nacelle_drag = 0.5 * air_density * u_glauert**2 * NACELLE_HUB_SCD_M2
        nacelle_moment = nacelle_drag * NACELLE_HUB_MOMENT_ARM_M

        tower_velocity = np.maximum(u_glauert * (1.0 - 2.0 * axial_induction), 0.0)
        tower_drag = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * TOWER_HEIGHT_M
        )
        tower_moment = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * (TOWER_HEIGHT_M**2 / 2.0)
        )

        rotor_moment = measured_positive_moment - nacelle_moment - tower_moment
        thrust = np.divide(
            rotor_moment,
            THRUST_LEVER_ARM_M,
            out=np.full(TARGET_LENGTH, np.nan),
            where=np.isfinite(rotor_moment),
        )
        thrust = np.where(np.isfinite(thrust), np.maximum(thrust, 0.0), np.nan)

        diff = np.abs(thrust - old)
        diff = diff[np.isfinite(diff)]
        if diff.size == 0 or np.max(diff) < CORRECTION_TOLERANCE:
            break

    return {
        "rotor_thrust": thrust,
        "u_glauert": u_glauert,
        "ct_glauert": ct_glauert,
        "blockage_factor": factor,
        "nacelle_drag": nacelle_drag,
        "nacelle_moment": nacelle_moment,
        "tower_drag": tower_drag,
        "tower_moment": tower_moment,
        "tower_velocity": tower_velocity,
        "axial_induction": axial_induction,
    }


def drag_corrected_with_fixed_u(
    tower_fa_moment: np.ndarray,
    air_density: np.ndarray,
    u_glauert_reference: np.ndarray,
) -> dict[str, np.ndarray]:
    """Yawed cases: drag correction uses the matching 0-degree Glauert velocity."""
    measured_positive_moment = -tower_fa_moment
    thrust = measured_positive_moment / THRUST_LEVER_ARM_M
    u_glauert = np.asarray(u_glauert_reference, dtype=float)

    nacelle_drag = np.zeros(TARGET_LENGTH)
    nacelle_moment = np.zeros(TARGET_LENGTH)
    tower_drag = np.zeros(TARGET_LENGTH)
    tower_moment = np.zeros(TARGET_LENGTH)
    tower_velocity = np.zeros(TARGET_LENGTH)
    axial_induction = np.zeros(TARGET_LENGTH)

    for _ in range(CORRECTION_MAX_ITERATIONS):
        old = thrust.copy()

        ct_denom = 0.5 * air_density * ROTOR_AREA_M2 * u_glauert**2
        ct_corr = np.divide(
            thrust,
            ct_denom,
            out=np.full(TARGET_LENGTH, np.nan),
            where=np.isfinite(thrust) & np.isfinite(ct_denom) & (ct_denom > 0.0),
        )
        axial_induction = axial_induction_from_ct(ct_corr)

        nacelle_drag = 0.5 * air_density * u_glauert**2 * NACELLE_HUB_SCD_M2
        nacelle_moment = nacelle_drag * NACELLE_HUB_MOMENT_ARM_M

        tower_velocity = np.maximum(u_glauert * (1.0 - 2.0 * axial_induction), 0.0)
        tower_drag = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * TOWER_HEIGHT_M
        )
        tower_moment = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * (TOWER_HEIGHT_M**2 / 2.0)
        )

        rotor_moment = measured_positive_moment - nacelle_moment - tower_moment
        thrust = np.divide(
            rotor_moment,
            THRUST_LEVER_ARM_M,
            out=np.full(TARGET_LENGTH, np.nan),
            where=np.isfinite(rotor_moment),
        )
        thrust = np.where(np.isfinite(thrust), np.maximum(thrust, 0.0), np.nan)

        diff = np.abs(thrust - old)
        diff = diff[np.isfinite(diff)]
        if diff.size == 0 or np.max(diff) < CORRECTION_TOLERANCE:
            break

    return {
        "rotor_thrust": thrust,
        "u_glauert": u_glauert,
        "nacelle_drag": nacelle_drag,
        "nacelle_moment": nacelle_moment,
        "tower_drag": tower_drag,
        "tower_moment": tower_moment,
        "tower_velocity": tower_velocity,
        "axial_induction": axial_induction,
    }


# ============================================================
# Excel and grouping
# ============================================================

def load_test_matrix() -> list[dict]:
    wb = openpyxl.load_workbook(TEST_MATRIX_PATH, data_only=True)
    ws = wb["Task1 Performance G2"]

    headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    rows = []

    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value is None:
            continue

        record = {}
        for col, header in enumerate(headers, start=1):
            record[header] = ws.cell(row=row, column=col).value
        rows.append(record)

    return rows


def yaw_suffix(yaw_value: float) -> str:
    yaw_int = int(round(float(yaw_value)))
    if yaw_int > 0:
        return f"plus_{yaw_int}"
    if yaw_int < 0:
        return f"minus_{abs(yaw_int)}"
    return "0"


def group_records_by_yaw(records: list[dict]) -> dict[str, dict]:
    grouped: dict[str, dict] = {}

    for record in records:
        yaw = record["Yaw [deg]"]
        if yaw is None:
            continue

        suffix = yaw_suffix(yaw)
        key = f"yaw_{suffix}"

        grouped.setdefault(
            key,
            {
                "yaw_target": yaw,
                "suffix": suffix,
                "records": [],
            },
        )
        grouped[key]["records"].append(record)

    for group in grouped.values():
        group["records"].sort(key=lambda r: (r["Test number"], r["Top"]))

    return grouped


def operating_key(record: dict) -> tuple:
    columns = [
        "Requested wind speed [m/s]",
        "Blade pitch [deg]",
        "Rotor speed [rpm]",
        "TSR [-]",
    ]

    key = []
    for column in columns:
        value = record.get(column)
        if isinstance(value, (int, float)):
            key.append(round(float(value), 8))
        else:
            key.append(value)
    return tuple(key)


def mat_path_from_record(record: dict) -> Path:
    return DATA_DIR / f"Results_File_{int(record['File number'])}_TOP_{int(record['Top'])}.mat"


# ============================================================
# Loading and derived quantities
# ============================================================

def load_measured(mat_path: Path) -> dict[str, np.ndarray]:
    model1 = load_model1(mat_path)

    raw = {
        "rotor_speed_rpm": nested(model1, "RotorSpeed"),
        "pitot_velocity": nested(model1, "PitotVelocity"),
        "air_density": nested(model1, "AirDensity"),
        "hub_torque_raw": nested(model1, "Hub", "Torque"),
        "tower_fa_moment": nested(model1, "Tower", "FA"),
        "yaw_actual": nested(model1, "Yaw"),
    }

    measured = {name: resample_to_2000(values, name) for name, values in raw.items()}

    measured["omega"] = measured["rotor_speed_rpm"] * 2.0 * math.pi / 60.0
    measured["tip_speed"] = measured["omega"] * ROTOR_RADIUS_M
    measured["torque"] = measured["hub_torque_raw"] * TORQUE_SCALE_TO_NM

    return measured


def build_zero_yaw_references(grouped: dict[str, dict]) -> dict[tuple, dict]:
    zero_group = grouped.get("yaw_0")
    if zero_group is None:
        raise RuntimeError("No 0-degree yaw group found. It is required for the Glauert reference.")

    references = {}

    for record in zero_group["records"]:
        mat_path = mat_path_from_record(record)
        if not mat_path.exists():
            print(f"Missing zero-yaw reference file: {mat_path.name}")
            continue

        measured = load_measured(mat_path)
        correction = drag_corrected_zero_yaw(
            tower_fa_moment=measured["tower_fa_moment"],
            pitot_velocity=measured["pitot_velocity"],
            air_density=measured["air_density"],
        )

        references[operating_key(record)] = {
            "record": record,
            "measured": measured,
            "correction": correction,
            "u_glauert": correction["u_glauert"],
            "ct_glauert": correction["ct_glauert"],
            "blockage_factor": correction["blockage_factor"],
        }

    if not references:
        raise RuntimeError("No zero-yaw reference could be created.")

    return references


def nearest_zero_reference(record: dict, references: dict[tuple, dict]) -> dict:
    key = operating_key(record)
    if key in references:
        return references[key]

    target_tsr = record.get("TSR [-]")
    target_rpm = record.get("Rotor speed [rpm]")

    best_ref = None
    best_distance = math.inf

    for reference in references.values():
        zero_record = reference["record"]
        distance = 0.0

        if isinstance(target_tsr, (int, float)) and isinstance(zero_record.get("TSR [-]"), (int, float)):
            distance += abs(float(target_tsr) - float(zero_record["TSR [-]"]))

        if isinstance(target_rpm, (int, float)) and isinstance(zero_record.get("Rotor speed [rpm]"), (int, float)):
            distance += abs(float(target_rpm) - float(zero_record["Rotor speed [rpm]"])) / 1000.0

        if distance < best_distance:
            best_distance = distance
            best_ref = reference

    if best_ref is None:
        raise RuntimeError(f"No zero-yaw reference found for test {record.get('Test number')}.")

    print(
        f"WARNING: no exact zero-yaw reference for test {record.get('Test number')}; "
        "nearest operating point was used."
    )

    return best_ref


def compute_derived(
    measured: dict[str, np.ndarray],
    record: dict,
    zero_ref: dict,
) -> dict[str, np.ndarray]:
    gamma_deg = float(record["Yaw [deg]"])
    cos_gamma = math.cos(math.radians(gamma_deg))
    if cos_gamma <= 0.0:
        raise ValueError(f"Invalid cos(gamma) for yaw={gamma_deg} deg.")

    u_glauert = zero_ref["u_glauert"].copy()
    is_zero = abs(gamma_deg) < 1e-9 and operating_key(record) == operating_key(zero_ref["record"])

    if is_zero:
        correction = zero_ref["correction"]
    else:
        correction = drag_corrected_with_fixed_u(
            tower_fa_moment=measured["tower_fa_moment"],
            air_density=measured["air_density"],
            u_glauert_reference=u_glauert,
        )

    thrust = correction["rotor_thrust"]
    power = measured["torque"] * measured["omega"]

    # For yawed cases, cP and cT use the orthogonal inflow.
    u_orthogonal = u_glauert * cos_gamma

    # TSR always uses the Glauert-corrected velocity, not the yaw-corrected velocity.
    tsr = np.divide(
        measured["tip_speed"],
        u_glauert,
        out=np.full(TARGET_LENGTH, np.nan),
        where=np.isfinite(u_glauert) & (u_glauert != 0.0),
    )

    cp_denom = 0.5 * measured["air_density"] * ROTOR_AREA_M2 * u_orthogonal**3
    cp = np.divide(
        power,
        cp_denom,
        out=np.full(TARGET_LENGTH, np.nan),
        where=np.isfinite(cp_denom) & (cp_denom != 0.0),
    )

    ct_denom = 0.5 * measured["air_density"] * ROTOR_AREA_M2 * u_orthogonal**2
    ct = np.divide(
        thrust,
        ct_denom,
        out=np.full(TARGET_LENGTH, np.nan),
        where=np.isfinite(ct_denom) & (ct_denom != 0.0),
    )

    return {
        "u_glauert": u_glauert,
        "u_orthogonal": u_orthogonal,
        "tsr": tsr,
        "cp": cp,
        "ct": ct,
        "thrust": thrust,
        "power": power,
        "torque": measured["torque"],
        "omega": measured["omega"],
        "ct_glauert_reference": zero_ref["ct_glauert"],
        "blockage_factor_reference": zero_ref["blockage_factor"],
    }


def build_yaw_cases(grouped: dict[str, dict], zero_refs: dict[tuple, dict]) -> dict[str, dict]:
    yaw_cases = {}

    for yaw_key, group in grouped.items():
        records = []
        measured_list = []
        derived_list = []

        for record in group["records"]:
            mat_path = mat_path_from_record(record)
            if not mat_path.exists():
                print(f"Missing MAT file, skipped: {mat_path.name}")
                continue

            measured = load_measured(mat_path)
            zero_ref = nearest_zero_reference(record, zero_refs)
            derived = compute_derived(measured, record, zero_ref)

            record_copy = dict(record)
            record_copy["mat_file"] = mat_path.name

            records.append(record_copy)
            measured_list.append(measured)
            derived_list.append(derived)

        yaw_cases[yaw_key] = {
            "yaw_target": group["yaw_target"],
            "suffix": group["suffix"],
            "records": records,
            "measured": measured_list,
            "derived": derived_list,
        }

    return yaw_cases


# ============================================================
# Rows and output
# ============================================================

def add_pair(row: dict, prefix: str, values: np.ndarray) -> None:
    row[f"{prefix}_median"] = stat(values, np.median)
    row[f"{prefix}_welch"] = welch_dc_estimate(values)


def compute_rows(yaw_cases: dict[str, dict]) -> list[dict]:
    rows = []

    for yaw_key, case in yaw_cases.items():
        for record, measured, derived in zip(case["records"], case["measured"], case["derived"]):
            row = {
                "yaw_group": yaw_key,
                "test": record["Test number"],
                "top": record["Top"],
                "file_number": record["File number"],
                "mat_file": record["mat_file"],
                "yaw_target": record["Yaw [deg]"],
                "yaw_actual_median": stat(measured["yaw_actual"], np.median),
                "tsr_expected": record.get("TSR [-]"),
                "cp_expected": record.get("Expected C_P [-]"),
                "ct_expected": record.get("Expected C_T (thrust) [-]"),
                "cos_gamma": math.cos(math.radians(float(record["Yaw [deg]"]))),
            }

            for prefix, values in [
                ("tsr", derived["tsr"]),
                ("cp", derived["cp"]),
                ("ct", derived["ct"]),
                ("u_inf_glauert_corrected", derived["u_glauert"]),
                ("u_orthogonal_for_performance", derived["u_orthogonal"]),
                ("omega", derived["omega"]),
                ("thrust", derived["thrust"]),
                ("power", derived["power"]),
                ("torque", derived["torque"]),
                ("ct_glauert_reference", derived["ct_glauert_reference"]),
                ("blockage_factor_reference", derived["blockage_factor_reference"]),
            ]:
                add_pair(row, prefix, values)

            f_cp, pxx_cp = welch_spectrum(derived["cp"])
            f_ct, pxx_ct = welch_spectrum(derived["ct"])
            row["cp_welch_frequency_hz"] = f_cp
            row["cp_welch_spectrum"] = pxx_cp
            row["ct_welch_frequency_hz"] = f_ct
            row["ct_welch_spectrum"] = pxx_ct

            if pxx_cp.size > 1:
                idx = int(np.argmax(pxx_cp[1:])) + 1
                row["cp_welch_dominant_frequency_hz"] = float(f_cp[idx])
                row["cp_welch_dominant_spectrum"] = float(pxx_cp[idx])
            else:
                row["cp_welch_dominant_frequency_hz"] = math.nan
                row["cp_welch_dominant_spectrum"] = math.nan

            if pxx_ct.size > 1:
                idx = int(np.argmax(pxx_ct[1:])) + 1
                row["ct_welch_dominant_frequency_hz"] = float(f_ct[idx])
                row["ct_welch_dominant_spectrum"] = float(pxx_ct[idx])
            else:
                row["ct_welch_dominant_frequency_hz"] = math.nan
                row["ct_welch_dominant_spectrum"] = math.nan

            rows.append(row)

    return rows


def save_summary_csv(rows: list[dict]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / "cp_ct_welch_summary_grouped_by_yaw_zero_glauert.csv"

    columns = [
        "yaw_group", "test", "top", "file_number", "mat_file",
        "yaw_target", "yaw_actual_median", "cos_gamma",
        "tsr_expected", "cp_expected", "ct_expected",
        "tsr_median", "tsr_welch",
        "cp_median", "cp_welch",
        "ct_median", "ct_welch",
        "u_inf_glauert_corrected_median", "u_inf_glauert_corrected_welch",
        "u_orthogonal_for_performance_median", "u_orthogonal_for_performance_welch",
        "omega_median", "omega_welch",
        "thrust_median", "thrust_welch",
        "power_median", "power_welch",
        "torque_median", "torque_welch",
        "ct_glauert_reference_median", "ct_glauert_reference_welch",
        "blockage_factor_reference_median", "blockage_factor_reference_welch",
        "cp_welch_dominant_frequency_hz", "cp_welch_dominant_spectrum",
        "ct_welch_dominant_frequency_hz", "ct_welch_dominant_spectrum",
    ]

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, delimiter=";")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c, math.nan) for c in columns})

    return path


def save_arrays_npz(yaw_cases: dict[str, dict]) -> Path:
    path = OUTPUT_DIR / "yaw_grouped_time_series_arrays_welch_zero_glauert.npz"
    arrays = {}

    for yaw_key, case in yaw_cases.items():
        suffix = case["suffix"]

        for i, (record, measured, derived) in enumerate(zip(case["records"], case["measured"], case["derived"])):
            # Store each test number explicitly.
            test = int(record["Test number"])
            for name, values in measured.items():
                arrays[f"{name}_{suffix}_test_{test}"] = values
            for name, values in derived.items():
                arrays[f"{name}_{suffix}_test_{test}"] = values

    if arrays:
        np.savez_compressed(path, **arrays)

    return path


# ============================================================
# Plotting
# ============================================================

def annotate(x, y, tests, dx=4, dy=4):
    for xi, yi, test in zip(x, y, tests):
        if np.isfinite(xi) and np.isfinite(yi):
            plt.annotate(
                str(int(test)),
                (xi, yi),
                textcoords="offset points",
                xytext=(dx, dy),
                fontsize=8,
            )


def plot_performance(rows: list[dict], yaw_group: str, yaw_target: float, coeff: str):
    selected = [r for r in rows if r["yaw_group"] == yaw_group]
    selected.sort(key=lambda r: r["tsr_median"])

    tests = [r["test"] for r in selected]
    x_med = [r["tsr_median"] for r in selected]
    y_med = [r[f"{coeff}_median"] for r in selected]
    x_welch = [r["tsr_welch"] for r in selected]
    y_welch = [r[f"{coeff}_welch"] for r in selected]

    label = "cP" if coeff == "cp" else "cT"

    plt.figure(figsize=(11, 6.5))
    plt.plot(x_med, y_med, marker="o", label=f"{label} median")
    plt.plot(x_welch, y_welch, marker="s", linestyle="--", label=f"{label} Welch estimate")

    annotate(x_med, y_med, tests, 4, 5)
    annotate(x_welch, y_welch, tests, 4, -12)

    plt.xlabel("TSR using Glauert-corrected velocity [-]")
    plt.ylabel(f"{label} [-]")
    plt.title(f"{label} over TSR at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    suffix = yaw_suffix(yaw_target)
    plt.savefig(OUTPUT_DIR / f"{label}_median_welch_yaw_{suffix}_test_numbers.png", dpi=200)
    plt.close()


def plot_variable_by_test(rows, yaw_group, yaw_target, prefix, ylabel, title):
    selected = [r for r in rows if r["yaw_group"] == yaw_group]
    selected.sort(key=lambda r: r["test"])

    x = [r["test"] for r in selected]
    y_med = [r[f"{prefix}_median"] for r in selected]
    y_welch = [r[f"{prefix}_welch"] for r in selected]

    plt.figure(figsize=(11, 6.5))
    plt.plot(x, y_med, marker="o", label="Median")
    plt.plot(x, y_welch, marker="s", linestyle="--", label="Welch estimate")
    plt.xlabel("Test number from Excel [-]")
    plt.ylabel(ylabel)
    plt.title(f"{title} by test number at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    suffix = yaw_suffix(yaw_target)
    name = prefix.replace("u_inf_glauert_corrected", "u-glauert").replace("_", "-")
    plt.savefig(OUTPUT_DIR / f"{name}_median_welch_by_test_yaw_{suffix}.png", dpi=200)
    plt.close()


def plot_mean_welch_spectrum(rows, yaw_group, yaw_target, coeff):
    selected = [r for r in rows if r["yaw_group"] == yaw_group]
    if not selected:
        return

    fk = f"{coeff}_welch_frequency_hz"
    sk = f"{coeff}_welch_spectrum"
    lengths = [len(r[fk]) for r in selected if len(r[fk]) > 0]

    if not lengths:
        return

    common = min(lengths)
    f = selected[0][fk][:common]
    spectra = np.vstack([r[sk][:common] for r in selected])
    mean_spectrum = np.nanmean(spectra, axis=0)

    label = "cP" if coeff == "cp" else "cT"

    plt.figure(figsize=(11, 6.5))
    plt.plot(f, mean_spectrum, label="Mean Welch spectrum")
    plt.xlabel("Frequency [Hz]")
    plt.ylabel(f"{label} Welch spectrum")
    plt.title(f"{label} Welch spectrum at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    suffix = yaw_suffix(yaw_target)
    plt.savefig(OUTPUT_DIR / f"{label}_welch_spectrum_yaw_{suffix}.png", dpi=200)
    plt.close()


def create_plots(rows: list[dict], yaw_cases: dict[str, dict]):
    for yaw_group, case in yaw_cases.items():
        if not case["records"]:
            continue

        yaw_target = case["yaw_target"]

        plot_performance(rows, yaw_group, yaw_target, "cp")
        plot_performance(rows, yaw_group, yaw_target, "ct")

        variable_plots = [
            ("u_inf_glauert_corrected", "U'_infinity Glauert-corrected [m/s]", "Glauert-corrected velocity"),
            ("omega", "Omega [rad/s]", "Rotor angular velocity"),
            ("thrust", "Drag-corrected thrust [N]", "Rotor thrust"),
            ("power", "Power [W]", "Aerodynamic power"),
            ("torque", "Torque [Nm]", "Aerodynamic torque"),
        ]

        for prefix, ylabel, title in variable_plots:
            plot_variable_by_test(rows, yaw_group, yaw_target, prefix, ylabel, title)

        plot_mean_welch_spectrum(rows, yaw_group, yaw_target, "cp")
        plot_mean_welch_spectrum(rows, yaw_group, yaw_target, "ct")


# ============================================================
# Main
# ============================================================

def main():
    if not DATA_DIR.exists():
        raise FileNotFoundError(f"Data folder not found: {DATA_DIR}")

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(f"Excel file not found: {TEST_MATRIX_PATH}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    records = load_test_matrix()
    grouped = group_records_by_yaw(records)

    # Glauert correction is calculated once from the 0-degree yaw cases.
    zero_refs = build_zero_yaw_references(grouped)

    # Yaw cases remain separated. Each case uses the matching zero-yaw Glauert reference.
    yaw_cases = build_yaw_cases(grouped, zero_refs)

    rows = compute_rows(yaw_cases)
    if not rows:
        raise RuntimeError("No tests could be processed.")

    summary_csv = save_summary_csv(rows)
    arrays_npz = save_arrays_npz(yaw_cases)
    create_plots(rows, yaw_cases)

    print("Finished.")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")
    print(f"Summary CSV: {summary_csv.resolve()}")
    print(f"Grouped arrays NPZ: {arrays_npz.resolve()}")

    print("\nYaw groups processed separately:")
    for yaw_key, case in yaw_cases.items():
        print(f" - {yaw_key}: {len(case['records'])} tests, yaw target = {case['yaw_target']}°")

    print("\nImportant:")
    print("Excel is the reference for test number, Top, File number and yaw grouping.")
    print("Glauert correction is calculated only from the matching 0-degree yaw case.")
    print("C_T for Glauert uses tunnel/Pitot velocity, not Glauert-corrected velocity.")
    print("Yawed cP and cT use U'_infinity*cos(gamma).")
    print("TSR uses U'_infinity without yaw cosine correction.")
    print("Raw FFT was replaced by scipy.signal.welch.")
    print("cP/cT plots are annotated with Excel test numbers.")
    print(f"Welch nperseg={WELCH_NPERSEG}, noverlap={WELCH_NOOVERLAP}.")
    print(f"Blockage ratio alpha={BLOCKAGE_RATIO_ALPHA:.5f}.")


if __name__ == "__main__":
    main()
