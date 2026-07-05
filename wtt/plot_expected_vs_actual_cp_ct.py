from __future__ import annotations

from pathlib import Path
import math
import csv

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat


# ============================================================
# Paths and constants
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
TEST_MATRIX_PATH = DATA_DIR / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "plots_expected_vs_actual_cp_ct"

MEASUREMENT_DURATION_S = 8.0
TARGET_LENGTH = 2000

ROTOR_RADIUS_M = 0.5436
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2

# Hub.Torque appears to be in mNm.
TORQUE_SCALE_TO_NM = 1e-3

# Same simple thrust conversion as in the blockage-only script.
THRUST_LEVER_ARM_M = 0.72

# Wind-tunnel dimensions
TUNNEL_WIDTH_M = 2.7
TUNNEL_HEIGHT_M = 1.8
TUNNEL_AREA_M2 = TUNNEL_WIDTH_M * TUNNEL_HEIGHT_M
BLOCKAGE_RATIO_ALPHA = ROTOR_AREA_M2 / TUNNEL_AREA_M2


# ============================================================
# Helpers
# ============================================================

def unwrap_singleton(value):
    current = value
    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]
    return current


def load_model1(mat_path: Path):
    mat_data = loadmat(
        mat_path,
        squeeze_me=True,
        struct_as_record=False,
    )
    data = unwrap_singleton(mat_data["Data"])
    models_data = unwrap_singleton(data.ModelsData)
    model1 = unwrap_singleton(models_data.Model1)
    return model1


def get_array(value) -> np.ndarray:
    array = np.asarray(unwrap_singleton(value), dtype=float).ravel(order="C")
    if array.size == 0:
        return np.array([np.nan], dtype=float)
    return array


def block_average_to_target_length(
    values: np.ndarray,
    target_length: int,
) -> np.ndarray:
    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size % target_length != 0:
        raise ValueError(
            "Block averaging requires the input length to be an integer "
            "multiple of target_length."
        )

    block_size = values.size // target_length
    blocks = values.reshape(target_length, block_size)

    with np.errstate(invalid="ignore"):
        return np.nanmean(blocks, axis=1)


def interpolate_to_target_length(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    values = np.asarray(values, dtype=float).ravel(order="C")

    if target_length <= 0:
        raise ValueError("target_length must be greater than zero.")

    if values.size == 0:
        return np.full(target_length, np.nan, dtype=float)

    if values.size == 1:
        return np.full(target_length, values[0], dtype=float)

    if values.size == target_length:
        return values.astype(float, copy=False)

    finite_mask = np.isfinite(values)

    if finite_mask.sum() == 0:
        return np.full(target_length, np.nan, dtype=float)

    if finite_mask.sum() == 1:
        return np.full(target_length, values[finite_mask][0], dtype=float)

    source_time = np.linspace(
        0.0,
        duration_s,
        values.size,
        endpoint=False,
    )
    target_time = np.linspace(
        0.0,
        duration_s,
        target_length,
        endpoint=False,
    )

    return np.interp(
        target_time,
        source_time[finite_mask],
        values[finite_mask],
    )


def resample_to_2000(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size == TARGET_LENGTH:
        return values.astype(float, copy=False)

    if values.size > TARGET_LENGTH and values.size % TARGET_LENGTH == 0:
        return block_average_to_target_length(values, TARGET_LENGTH)

    return interpolate_to_target_length(values, TARGET_LENGTH)


def finite_stat(values: np.ndarray, fn) -> float:
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]

    if values.size == 0:
        return math.nan

    return float(fn(values))


def glauert_blockage_corrected_velocity(
    pitot_velocity: np.ndarray,
    air_density: np.ndarray,
    thrust_n: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Applies the Glauert blockage correction to the time series.

    U'_inf / U_inf = 1 + (alpha / 4) * C_T / (1 - C_T)
    C_T = T / (0.5 * rho * A_R * U_inf^2)
    """

    pitot_velocity = np.asarray(pitot_velocity, dtype=float)
    air_density = np.asarray(air_density, dtype=float)
    thrust_n = np.asarray(thrust_n, dtype=float)

    ct_denominator_uncorrected = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * pitot_velocity**2
    )

    ct_for_blockage = np.divide(
        thrust_n,
        ct_denominator_uncorrected,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(thrust_n)
            & np.isfinite(ct_denominator_uncorrected)
            & (ct_denominator_uncorrected > 0.0)
        ),
    )

    blockage_factor = np.full(TARGET_LENGTH, np.nan, dtype=float)

    valid = (
        np.isfinite(pitot_velocity)
        & np.isfinite(ct_for_blockage)
        & (pitot_velocity > 0.0)
        & (ct_for_blockage >= 0.0)
        & (ct_for_blockage < 1.0)
        & np.isfinite(1.0 - ct_for_blockage)
        & ((1.0 - ct_for_blockage) != 0.0)
    )

    blockage_factor[valid] = (
        1.0
        + (BLOCKAGE_RATIO_ALPHA / 4.0)
        * ct_for_blockage[valid]
        / (1.0 - ct_for_blockage[valid])
    )

    u_blockage_corrected = pitot_velocity * blockage_factor

    return u_blockage_corrected, ct_for_blockage, blockage_factor


# ============================================================
# Actual-value calculation
# ============================================================

def compute_actual_test_statistics(mat_path: Path) -> dict:
    model1 = load_model1(mat_path)

    rotor_speed_rpm = get_array(model1.RotorSpeed)
    pitot_velocity = get_array(model1.PitotVelocity)
    air_density = get_array(model1.AirDensity)
    hub_torque_raw = get_array(model1.Hub.Torque)
    tower_fa_moment = get_array(model1.Tower.FA)
    yaw_actual = get_array(model1.Yaw)

    rotor_speed_common = resample_to_2000(rotor_speed_rpm)
    pitot_velocity_common = resample_to_2000(pitot_velocity)
    air_density_common = resample_to_2000(air_density)
    torque_common_raw = resample_to_2000(hub_torque_raw)
    tower_fa_common = resample_to_2000(tower_fa_moment)
    yaw_common = resample_to_2000(yaw_actual)

    omega = rotor_speed_common * 2.0 * math.pi / 60.0
    torque_nm = torque_common_raw * TORQUE_SCALE_TO_NM
    power_w = torque_nm * omega

    # Same assumption as in the blockage-only script:
    # negative Tower.FA corresponds to positive thrust direction.
    thrust_n = -tower_fa_common / THRUST_LEVER_ARM_M

    u_blockage_corrected, ct_for_blockage, blockage_factor = (
        glauert_blockage_corrected_velocity(
            pitot_velocity=pitot_velocity_common,
            air_density=air_density_common,
            thrust_n=thrust_n,
        )
    )

    tsr_series = np.divide(
        omega * ROTOR_RADIUS_M,
        u_blockage_corrected,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(u_blockage_corrected)
            & (u_blockage_corrected != 0.0)
        ),
    )

    cp_denominator = (
        0.5
        * air_density_common
        * ROTOR_AREA_M2
        * u_blockage_corrected**3
    )

    cp_series = np.divide(
        power_w,
        cp_denominator,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(cp_denominator)
            & (cp_denominator != 0.0)
        ),
    )

    ct_denominator_corrected = (
        0.5
        * air_density_common
        * ROTOR_AREA_M2
        * u_blockage_corrected**2
    )

    ct_series = np.divide(
        thrust_n,
        ct_denominator_corrected,
        out=np.full(TARGET_LENGTH, np.nan, dtype=float),
        where=(
            np.isfinite(ct_denominator_corrected)
            & (ct_denominator_corrected != 0.0)
        ),
    )

    return {
        "yaw_actual_median": finite_stat(yaw_common, np.median),
        "pitot_velocity_median": finite_stat(pitot_velocity_common, np.median),
        "u_blockage_corrected_median": finite_stat(u_blockage_corrected, np.median),
        "tsr_actual_median": finite_stat(tsr_series, np.median),
        "cp_actual_median": finite_stat(cp_series, np.median),
        "ct_actual_median": finite_stat(ct_series, np.median),
        "ct_for_blockage_median": finite_stat(ct_for_blockage, np.median),
        "blockage_factor_median": finite_stat(blockage_factor, np.median),
    }


# ============================================================
# Read expected values from Excel and combine
# ============================================================

def load_test_matrix(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Task1 Performance G2"]

    headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    records = []

    for row_idx in range(2, ws.max_row + 1):
        test_number = ws.cell(row=row_idx, column=1).value

        if test_number is None:
            continue

        record = {}
        for col_idx, header in enumerate(headers, start=1):
            record[header] = ws.cell(row=row_idx, column=col_idx).value

        records.append(record)

    return records


def build_comparison_rows() -> list[dict]:
    records = load_test_matrix(TEST_MATRIX_PATH)
    rows = []

    for record in records:
        top = record["Top"]
        file_number = record["File number"]

        if top is None or file_number is None:
            continue

        mat_name = f"Results_File_{int(file_number)}_TOP_{int(top)}.mat"
        mat_path = DATA_DIR / mat_name

        if not mat_path.exists():
            print(f"Missing MAT file, skipped: {mat_name}")
            continue

        actual = compute_actual_test_statistics(mat_path)

        rows.append(
            {
                "test": record["Test number"],
                "yaw_target": record["Yaw [deg]"],
                "yaw_actual_median": actual["yaw_actual_median"],

                "tsr_expected": record["TSR [-]"],
                "cp_expected": record["Expected C_P [-]"],
                "ct_expected": record["Expected C_T (thrust) [-]"],

                "tsr_actual_median": actual["tsr_actual_median"],
                "cp_actual_median": actual["cp_actual_median"],
                "ct_actual_median": actual["ct_actual_median"],

                "pitot_velocity_median": actual["pitot_velocity_median"],
                "u_blockage_corrected_median": actual["u_blockage_corrected_median"],
                "ct_for_blockage_median": actual["ct_for_blockage_median"],
                "blockage_factor_median": actual["blockage_factor_median"],

                "top": top,
                "file_number": file_number,
                "mat_file": mat_name,
            }
        )

    return rows


# ============================================================
# Output
# ============================================================

def save_summary_csv(rows: list[dict]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / "expected_vs_actual_cp_ct_summary.csv"

    headers = [
        "test",
        "yaw_target",
        "yaw_actual_median",
        "tsr_expected",
        "cp_expected",
        "ct_expected",
        "tsr_actual_median",
        "cp_actual_median",
        "ct_actual_median",
        "pitot_velocity_median",
        "u_blockage_corrected_median",
        "ct_for_blockage_median",
        "blockage_factor_median",
        "top",
        "file_number",
        "mat_file",
    ]

    with output_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)

    return output_path


def plot_expected_vs_actual(
    rows: list[dict],
    yaw_value: float,
    coefficient_prefix: str,
) -> None:
    yaw_rows = [r for r in rows if r["yaw_target"] == yaw_value]

    expected_pairs = [
        (r["tsr_expected"], r[f"{coefficient_prefix}_expected"])
        for r in yaw_rows
        if r["tsr_expected"] is not None and r[f"{coefficient_prefix}_expected"] is not None
    ]
    actual_pairs = [
        (r["tsr_actual_median"], r[f"{coefficient_prefix}_actual_median"])
        for r in yaw_rows
        if np.isfinite(r["tsr_actual_median"]) and np.isfinite(r[f"{coefficient_prefix}_actual_median"])
    ]

    expected_pairs.sort(key=lambda x: x[0])
    actual_pairs.sort(key=lambda x: x[0])

    x_expected = [pair[0] for pair in expected_pairs]
    y_expected = [pair[1] for pair in expected_pairs]

    x_actual = [pair[0] for pair in actual_pairs]
    y_actual = [pair[1] for pair in actual_pairs]

    label = "cP" if coefficient_prefix == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(
        x_expected,
        y_expected,
        marker="o",
        linestyle="--",
        label=f"Expected {label}",
    )
    plt.plot(
        x_actual,
        y_actual,
        marker="o",
        label=f"Actual {label} (measured, median)",
    )

    plt.xlabel("Tip-speed ratio TSR [-]")
    plt.ylabel(f"{label} [-]")
    plt.title(f"Expected vs actual {label} at yaw = {yaw_value}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    out_name = (
        f"{label}_expected_vs_actual_yaw_{int(yaw_value):+d}.png"
        .replace("+", "plus")
        .replace("-", "minus")
    )
    plt.savefig(OUTPUT_DIR / out_name, dpi=200)
    plt.close()


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(f"Excel file not found: {TEST_MATRIX_PATH}")

    if not DATA_DIR.exists():
        raise FileNotFoundError(f"Data folder not found: {DATA_DIR}")

    rows = build_comparison_rows()

    if not rows:
        raise RuntimeError("No comparison rows could be generated.")

    summary_path = save_summary_csv(rows)

    yaw_values = sorted(
        {
            r["yaw_target"]
            for r in rows
            if r["yaw_target"] is not None
        }
    )

    for yaw in yaw_values:
        plot_expected_vs_actual(rows, yaw, "cp")
        plot_expected_vs_actual(rows, yaw, "ct")

    print("Finished.")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")
    print(f"Summary CSV: {summary_path.resolve()}")
    print("Created comparison plots for expected vs actual cP and cT.")
    print("Expected values come from the Excel columns:")
    print(" - TSR [-]")
    print(" - Expected C_P [-]")
    print(" - Expected C_T (thrust) [-]")
    print("Actual values are computed from the MAT files using:")
    print(" - 2000-sample processing")
    print(" - block averaging for Hub.Torque when needed")
    print(" - interpolation for AirDensity when needed")
    print(" - Glauert blockage-corrected velocity u_blockage_corrected")
    print(" - median actual TSR, cP and cT")


if __name__ == "__main__":
    main()
