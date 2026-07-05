from __future__ import annotations

from pathlib import Path
import math

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat


# =========================
# User settings / constants
# =========================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct"
TEST_MATRIX_PATH = (
    PROJECT_ROOT
    / "data"
    / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"
)

# From the repository / README:
ROTOR_RADIUS_M = 0.5436
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2

# Hub.Torque in the MAT files appears to be in mNm for the uploaded examples.
# Therefore convert to Nm:
TORQUE_SCALE_TO_NM = 1e-3

# IMPORTANT:
# To compute cT from Tower.FA you need a lever arm / calibration distance.
# Replace this with the correct distance from your setup if necessary.
# Example assumption:
THRUST_LEVER_ARM_M = 0.80


# =========================
# Helper functions
# =========================

def unwrap_singleton(value):
    current = value
    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]
    return current


def load_model1(mat_path: Path):
    mat_data = loadmat(mat_path, squeeze_me=True, struct_as_record=False)
    data = unwrap_singleton(mat_data["Data"])
    models_data = unwrap_singleton(data.ModelsData)
    model1 = unwrap_singleton(models_data.Model1)
    return model1


def get_array(value) -> np.ndarray:
    array = np.asarray(unwrap_singleton(value), dtype=float).ravel(order="C")
    if array.size == 0:
        return np.array([np.nan], dtype=float)
    return array


def finite_only(*arrays):
    """Return arrays filtered to indices where all arrays are finite."""
    stacked = [np.asarray(a, dtype=float).ravel(order="C") for a in arrays]
    min_len = min(len(a) for a in stacked)
    stacked = [a[:min_len] for a in stacked]
    mask = np.ones(min_len, dtype=bool)
    for a in stacked:
        mask &= np.isfinite(a)
    return [a[mask] for a in stacked]


def safe_stat(values: np.ndarray, fn):
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]
    if values.size == 0:
        return math.nan
    return float(fn(values))


def compute_test_statistics(mat_path: Path) -> dict:
    """Compute measured TSR, cP and cT statistics from one MAT file."""
    model1 = load_model1(mat_path)

    rotor_speed_rpm = get_array(model1.RotorSpeed)
    pitot_velocity = get_array(model1.PitotVelocity)
    air_density = get_array(model1.AirDensity)
    hub_torque_raw = get_array(model1.Hub.Torque)
    tower_fa_moment = get_array(model1.Tower.FA)
    yaw_actual = get_array(model1.Yaw)

    # Angular speed in rad/s
    omega = rotor_speed_rpm * 2.0 * math.pi / 60.0

    # Torque in Nm
    torque_nm = hub_torque_raw * TORQUE_SCALE_TO_NM

    # Measured power
    power_w = torque_nm * omega

    # Align the arrays used for TSR / cP
    omega_cp, u_cp, rho_cp, power_cp = finite_only(omega, pitot_velocity, air_density, power_w)

    tsr_series = omega_cp * ROTOR_RADIUS_M / u_cp
    cp_series = 2.0 * power_cp / (rho_cp * ROTOR_AREA_M2 * u_cp**3)

    # Convert Tower.FA moment to thrust force
    # Sign convention: tower fore-aft is often negative under thrust loading.
    # The minus sign makes thrust positive in the usual operating direction.
    thrust_n = -tower_fa_moment / THRUST_LEVER_ARM_M

    thrust_ct, u_ct, rho_ct = finite_only(thrust_n, pitot_velocity, air_density)
    ct_series = thrust_ct / (0.5 * rho_ct * ROTOR_AREA_M2 * u_ct**2)

    return {
        "yaw_actual_median": safe_stat(yaw_actual, np.median),
        "tsr_median": safe_stat(tsr_series, np.median),
        "tsr_std": safe_stat(tsr_series, np.std),
        "cp_median": safe_stat(cp_series, np.median),
        "cp_std": safe_stat(cp_series, np.std),
        "cp_min": safe_stat(cp_series, np.min),
        "cp_max": safe_stat(cp_series, np.max),
        "ct_median": safe_stat(ct_series, np.median),
        "ct_std": safe_stat(ct_series, np.std),
        "ct_min": safe_stat(ct_series, np.min),
        "ct_max": safe_stat(ct_series, np.max),
    }


def load_test_matrix(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Task1 Performance G2"]

    headers = [ws.cell(row=1, column=i).value for i in range(1, 31)]

    records = []
    for row_idx in range(2, ws.max_row + 1):
        test_number = ws.cell(row=row_idx, column=1).value
        if test_number is None:
            continue

        record = {}
        for col_idx, header in enumerate(headers, start=1):
            record[header] = ws.cell(row=row_idx, column=col_idx).value
        records.append(record)

    return records


def build_dataframe_like_records() -> list[dict]:
    records = load_test_matrix(TEST_MATRIX_PATH)
    rows = []

    for record in records:
        top = record["Top"]
        file_number = record["File number"]

        if top is None or file_number is None:
            continue

        mat_name = f"Results_File_{int(file_number)}_TOP_{int(top)}.mat"
        mat_path = DATA_DIR / mat_name

        if not mat_path.exists():
            print(f"Datei fehlt, übersprungen: {mat_name}")
            continue

        stats = compute_test_statistics(mat_path)

        rows.append(
            {
                "test": record["Test number"],
                "yaw_target": record["Yaw [deg]"],
                "yaw_actual_median": stats["yaw_actual_median"],
                "tsr_expected": record["TSR [-]"],
                "tsr_median": stats["tsr_median"],
                "tsr_std": stats["tsr_std"],
                "cp_median": stats["cp_median"],
                "cp_plus_std": (
                    stats["cp_median"] + stats["cp_std"]
                    if np.isfinite(stats["cp_median"]) and np.isfinite(stats["cp_std"])
                    else math.nan
                ),
                "cp_minus_std": (
                    stats["cp_median"] - stats["cp_std"]
                    if np.isfinite(stats["cp_median"]) and np.isfinite(stats["cp_std"])
                    else math.nan
                ),
                "cp_min": stats["cp_min"],
                "cp_max": stats["cp_max"],
                "ct_median": stats["ct_median"],
                "ct_plus_std": (
                    stats["ct_median"] + stats["ct_std"]
                    if np.isfinite(stats["ct_median"]) and np.isfinite(stats["ct_std"])
                    else math.nan
                ),
                "ct_minus_std": (
                    stats["ct_median"] - stats["ct_std"]
                    if np.isfinite(stats["ct_median"]) and np.isfinite(stats["ct_std"])
                    else math.nan
                ),
                "ct_min": stats["ct_min"],
                "ct_max": stats["ct_max"],
                "mat_file": mat_name,
            }
        )

    return rows


def plot_for_yaw(rows: list[dict], yaw_value: float, coefficient_prefix: str):
    """
    coefficient_prefix: 'cp' or 'ct'
    Produces one plot with:
    median, median+std, median-std, min, max
    """
    yaw_rows = [r for r in rows if r["yaw_target"] == yaw_value]
    yaw_rows.sort(key=lambda r: (r["tsr_median"] if r["tsr_median"] is not None else math.inf))

    x = [r["tsr_median"] for r in yaw_rows]
    y_median = [r[f"{coefficient_prefix}_median"] for r in yaw_rows]
    y_plus_std = [r[f"{coefficient_prefix}_plus_std"] for r in yaw_rows]
    y_minus_std = [r[f"{coefficient_prefix}_minus_std"] for r in yaw_rows]
    y_min = [r[f"{coefficient_prefix}_min"] for r in yaw_rows]
    y_max = [r[f"{coefficient_prefix}_max"] for r in yaw_rows]

    coeff_label = "cP" if coefficient_prefix == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(x, y_median, marker="o", label="Median")
    plt.plot(x, y_plus_std, marker="o", linestyle="--", label="Median + Standardabweichung")
    plt.plot(x, y_minus_std, marker="o", linestyle="--", label="Median - Standardabweichung")
    plt.plot(x, y_min, marker="o", linestyle=":", label="Minimum")
    plt.plot(x, y_max, marker="o", linestyle=":", label="Maximum")

    plt.xlabel("TSR (gemessener Median)")
    plt.ylabel(coeff_label)
    plt.title(f"{coeff_label} über TSR bei Yaw = {yaw_value}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    out_name = f"{coeff_label}_yaw_{int(yaw_value):+d}.png".replace("+", "plus").replace("-", "minus")
    plt.savefig(OUTPUT_DIR / out_name, dpi=200)
    plt.close()


def save_summary_csv(rows: list[dict]):
    import csv

    out_csv = OUTPUT_DIR / "cp_ct_summary_by_test.csv"
    fieldnames = [
        "test", "yaw_target", "yaw_actual_median", "tsr_expected", "tsr_median", "tsr_std",
        "cp_median", "cp_plus_std", "cp_minus_std", "cp_min", "cp_max",
        "ct_median", "ct_plus_std", "ct_minus_std", "ct_min", "ct_max",
        "mat_file",
    ]

    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(f"Excel-Datei nicht gefunden: {TEST_MATRIX_PATH}")

    if not DATA_DIR.exists():
        raise FileNotFoundError(f"data-Ordner nicht gefunden: {DATA_DIR}")

    rows = build_dataframe_like_records()
    if not rows:
        raise RuntimeError("Es konnten keine Testdaten verarbeitet werden.")

    save_summary_csv(rows)

    yaw_values = sorted(set(r["yaw_target"] for r in rows if r["yaw_target"] is not None))
    for yaw in yaw_values:
        plot_for_yaw(rows, yaw, "cp")
        plot_for_yaw(rows, yaw, "ct")

    print("Fertig.")
    print(f"Ausgabeordner: {OUTPUT_DIR.resolve()}")
    print("Erzeugte Dateien:")
    for path in sorted(OUTPUT_DIR.iterdir()):
        print(" -", path.name)

    print("\nWICHTIG:")
    print("cP wurde aus Hub.Torque berechnet.")
    print("cT wurde aus Tower.FA über einen Hebelarm berechnet.")
    print(f"Aktueller Hebelarm in diesem Skript: {THRUST_LEVER_ARM_M} m")
    print("Falls euer Prüfstand einen anderen Kalibrierwert hat, bitte THRUST_LEVER_ARM_M anpassen.")


if __name__ == "__main__":
    main()
