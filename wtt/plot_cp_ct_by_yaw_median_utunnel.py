from __future__ import annotations

import csv
import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat

from blockage import blockage_ratio, equivalent_free_air_speed


# ============================================================
# Paths and constants
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
TEST_MATRIX_PATH = DATA_DIR / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct_median_utunnel"

MEASUREMENT_DURATION_S = 8.0

ROTOR_RADIUS_M = 0.5436
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2

THRUST_LEVER_ARM_M = 0.72
TORQUE_SCALE_TO_NM = 1e-3

TUNNEL_WIDTH_M = 2.7
TUNNEL_HEIGHT_M = 1.8
TUNNEL_TEST_SECTION_LENGTH_M = 4.5
TUNNEL_CROSS_SECTION_M2 = TUNNEL_WIDTH_M * TUNNEL_HEIGHT_M

BLOCKAGE_RATIO = blockage_ratio(
    rotor_area=ROTOR_AREA_M2,
    tunnel_area=TUNNEL_CROSS_SECTION_M2,
)

# Drag corrections from the task slides.
NACELLE_HUB_SCD_M2 = 0.0175
TOWER_DIAMETER_M = 0.047

# Configurable assumption for the tower as circular cylinder.
TOWER_DRAG_COEFFICIENT = 1.2

# Effective tower height above the tower-base load measurement location.
EFFECTIVE_TOWER_HEIGHT_M = THRUST_LEVER_ARM_M

CORRECTION_MAX_ITERATIONS = 50
CORRECTION_TOLERANCE = 1e-8


# ============================================================
# Data helpers
# ============================================================

def unwrap_singleton(value):
    current = value
    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]
    return current


def get_array(value) -> np.ndarray:
    return np.asarray(
        unwrap_singleton(value),
        dtype=float,
    ).ravel(order="C")


def load_model1(mat_path: Path):
    mat_data = loadmat(
        mat_path,
        squeeze_me=True,
        struct_as_record=False,
    )
    return mat_data["Data"].ModelsData.Model1


def resample_to_common_time(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    values = np.asarray(values, dtype=float).ravel(order="C")

    if target_length <= 0:
        raise ValueError("target_length must be greater than zero.")

    if values.size == 0:
        return np.full(target_length, np.nan)

    if values.size == 1:
        return np.full(target_length, values[0])

    if values.size == target_length:
        return values

    finite = np.isfinite(values)

    if finite.sum() == 0:
        return np.full(target_length, np.nan)

    if finite.sum() == 1:
        return np.full(target_length, values[finite][0])

    source_time = np.linspace(
        0.0,
        duration_s,
        values.size,
        endpoint=False,
    )
    target_time = np.linspace(
        0.0,
        duration_s,
        target_length,
        endpoint=False,
    )

    return np.interp(
        target_time,
        source_time[finite],
        values[finite],
    )


def finite_stat(values: np.ndarray, function) -> float:
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]
    return float(function(values)) if values.size else math.nan


def finite_median(values: np.ndarray) -> float:
    return finite_stat(values, np.median)


def finite_std(values: np.ndarray) -> float:
    return finite_stat(values, np.std)


def finite_min(values: np.ndarray) -> float:
    return finite_stat(values, np.min)


def finite_max(values: np.ndarray) -> float:
    return finite_stat(values, np.max)


# ============================================================
# Drag and blockage correction
# ============================================================

def axial_induction_from_ct(ct: np.ndarray) -> np.ndarray:
    """Classical momentum-theory relation for 0 <= CT < 1."""

    ct = np.asarray(ct, dtype=float)
    ct_limited = np.clip(ct, 0.0, 0.999999)
    return 0.5 * (1.0 - np.sqrt(1.0 - ct_limited))


def correct_rotor_thrust_and_velocity_from_utunnel(
    tower_fa_moment: np.ndarray,
    requested_tunnel_speed: float,
    air_density: np.ndarray,
) -> dict[str, np.ndarray]:
    """Uses Excel requested tunnel speed instead of measured PitotVelocity.

    Sign convention:
    Tower.FA is negative in the measured loading direction. Therefore
    -Tower.FA is treated as the positive measured fore-aft bending moment.

    Moment balance:
    -M_FA = T_rotor * L + D_nacelle * L + M_tower_drag

    The blockage correction uses U_tunnel = requested wind speed from Excel.
    PitotVelocity is not used in the calculation of U'_infinity here.
    """

    measured_positive_moment = -np.asarray(tower_fa_moment, dtype=float)
    air_density = np.asarray(air_density, dtype=float)

    u_tunnel = np.full_like(
        measured_positive_moment,
        float(requested_tunnel_speed),
        dtype=float,
    )

    rotor_thrust = measured_positive_moment / THRUST_LEVER_ARM_M

    free_air_velocity = u_tunnel.copy()
    nacelle_drag = np.zeros_like(rotor_thrust)
    tower_drag_force = np.zeros_like(rotor_thrust)
    tower_drag_moment = np.zeros_like(rotor_thrust)

    for _ in range(CORRECTION_MAX_ITERATIONS):
        previous_thrust = rotor_thrust.copy()

        ct_tunnel_denominator = (
            0.5
            * air_density
            * ROTOR_AREA_M2
            * u_tunnel**2
        )

        ct_for_blockage = np.divide(
            rotor_thrust,
            ct_tunnel_denominator,
            out=np.full_like(rotor_thrust, np.nan),
            where=(
                np.isfinite(ct_tunnel_denominator)
                & (ct_tunnel_denominator > 0.0)
            ),
        )

        valid = (
            np.isfinite(u_tunnel)
            & np.isfinite(ct_for_blockage)
            & (u_tunnel > 0.0)
            & (ct_for_blockage >= 0.0)
            & (ct_for_blockage < 1.0)
        )

        free_air_velocity = np.full_like(u_tunnel, np.nan)
        free_air_velocity[valid] = equivalent_free_air_speed(
            commanded_tunnel_speed_value=u_tunnel[valid],
            thrust_coefficient=ct_for_blockage[valid],
            blockage_ratio_value=BLOCKAGE_RATIO,
        )

        ct_free_denominator = (
            0.5
            * air_density
            * ROTOR_AREA_M2
            * free_air_velocity**2
        )

        ct_free = np.divide(
            rotor_thrust,
            ct_free_denominator,
            out=np.full_like(rotor_thrust, np.nan),
            where=(
                np.isfinite(ct_free_denominator)
                & (ct_free_denominator > 0.0)
            ),
        )

        induction = axial_induction_from_ct(ct_free)

        nacelle_drag = (
            0.5
            * air_density
            * free_air_velocity**2
            * NACELLE_HUB_SCD_M2
        )
        nacelle_drag_moment = nacelle_drag * THRUST_LEVER_ARM_M

        tower_velocity = free_air_velocity * (1.0 - 2.0 * induction)
        tower_velocity = np.maximum(tower_velocity, 0.0)

        distributed_tower_drag = (
            0.5
            * air_density
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * TOWER_DIAMETER_M
        )  # N/m

        tower_drag_force = (
            distributed_tower_drag
            * EFFECTIVE_TOWER_HEIGHT_M
        )

        tower_drag_moment = (
            distributed_tower_drag
            * EFFECTIVE_TOWER_HEIGHT_M**2
            / 2.0
        )

        rotor_moment = (
            measured_positive_moment
            - nacelle_drag_moment
            - tower_drag_moment
        )

        rotor_thrust = rotor_moment / THRUST_LEVER_ARM_M
        rotor_thrust = np.where(
            np.isfinite(rotor_thrust),
            np.maximum(rotor_thrust, 0.0),
            np.nan,
        )

        if np.all(~np.isfinite(rotor_thrust - previous_thrust)):
            break

        difference = np.nanmax(
            np.abs(rotor_thrust - previous_thrust)
        )

        if np.isfinite(difference) and difference < CORRECTION_TOLERANCE:
            break

    return {
        "u_tunnel": u_tunnel,
        "rotor_thrust": rotor_thrust,
        "free_air_velocity": free_air_velocity,
        "nacelle_drag": nacelle_drag,
        "tower_drag_force": tower_drag_force,
        "tower_drag_moment": tower_drag_moment,
    }


# ============================================================
# Test calculation
# ============================================================

def compute_test(
    mat_path: Path,
    requested_tunnel_speed: float,
) -> dict:
    model1 = load_model1(mat_path)

    rotor_speed_rpm = get_array(model1.RotorSpeed)
    pitot_velocity = get_array(model1.PitotVelocity)
    air_density = get_array(model1.AirDensity)
    torque_raw = get_array(model1.Hub.Torque)
    tower_fa = get_array(model1.Tower.FA)
    yaw = get_array(model1.Yaw)

    common_length = max(
        rotor_speed_rpm.size,
        pitot_velocity.size,
        air_density.size,
        torque_raw.size,
        tower_fa.size,
        yaw.size,
    )

    rotor_speed_rpm = resample_to_common_time(
        rotor_speed_rpm,
        common_length,
    )
    pitot_velocity = resample_to_common_time(
        pitot_velocity,
        common_length,
    )
    air_density = resample_to_common_time(
        air_density,
        common_length,
    )
    torque_raw = resample_to_common_time(
        torque_raw,
        common_length,
    )
    tower_fa = resample_to_common_time(
        tower_fa,
        common_length,
    )
    yaw = resample_to_common_time(
        yaw,
        common_length,
    )

    omega = rotor_speed_rpm * 2.0 * math.pi / 60.0
    tip_speed = omega * ROTOR_RADIUS_M

    torque_nm = torque_raw * TORQUE_SCALE_TO_NM
    power_w = torque_nm * omega

    correction = correct_rotor_thrust_and_velocity_from_utunnel(
        tower_fa_moment=tower_fa,
        requested_tunnel_speed=requested_tunnel_speed,
        air_density=air_density,
    )

    free_air_velocity = correction["free_air_velocity"]
    rotor_thrust = correction["rotor_thrust"]

    cp_denominator = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * free_air_velocity**3
    )

    cp = np.divide(
        power_w,
        cp_denominator,
        out=np.full(common_length, np.nan),
        where=np.isfinite(cp_denominator) & (cp_denominator > 0.0),
    )

    ct_denominator = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * free_air_velocity**2
    )

    ct = np.divide(
        rotor_thrust,
        ct_denominator,
        out=np.full(common_length, np.nan),
        where=np.isfinite(ct_denominator) & (ct_denominator > 0.0),
    )

    # Representative TSR for the test point.
    # Median numerator and denominator are formed separately.
    tip_speed_median = finite_median(tip_speed)
    free_air_velocity_median = finite_median(free_air_velocity)

    if (
        np.isfinite(tip_speed_median)
        and np.isfinite(free_air_velocity_median)
        and free_air_velocity_median != 0.0
    ):
        tsr_median = tip_speed_median / free_air_velocity_median
    else:
        tsr_median = math.nan

    return {
        "yaw_median": finite_median(yaw),
        "rotor_speed_rpm_median": finite_median(rotor_speed_rpm),
        "tip_speed_median": tip_speed_median,
        "requested_tunnel_speed": float(requested_tunnel_speed),
        "pitot_velocity_median": finite_median(pitot_velocity),
        "free_air_velocity_median": free_air_velocity_median,
        "tsr_median": tsr_median,

        "cp_median": finite_median(cp),
        "cp_std": finite_std(cp),
        "cp_min": finite_min(cp),
        "cp_max": finite_max(cp),

        "ct_median": finite_median(ct),
        "ct_std": finite_std(ct),
        "ct_min": finite_min(ct),
        "ct_max": finite_max(ct),

        "rotor_thrust_median": finite_median(rotor_thrust),
        "nacelle_drag_median": finite_median(correction["nacelle_drag"]),
        "tower_drag_force_median": finite_median(
            correction["tower_drag_force"]
        ),
        "tower_drag_moment_median": finite_median(
            correction["tower_drag_moment"]
        ),
    }


# ============================================================
# Test matrix and processing
# ============================================================

def load_test_matrix() -> list[dict]:
    workbook = openpyxl.load_workbook(
        TEST_MATRIX_PATH,
        data_only=True,
    )
    sheet = workbook["Task1 Performance G2"]

    records = []

    for row_index in range(2, sheet.max_row + 1):
        test_number = sheet.cell(row_index, 1).value

        if test_number is None:
            continue

        records.append(
            {
                "test": test_number,
                "requested_tunnel_speed": sheet.cell(row_index, 2).value,
                "excel_equivalent_free_air_speed": sheet.cell(row_index, 3).value,
                "yaw_target": sheet.cell(row_index, 6).value,
                "tsr_expected": sheet.cell(row_index, 8).value,
                "top": sheet.cell(row_index, 29).value,
                "file_number": sheet.cell(row_index, 30).value,
            }
        )

    return records


def process_all_tests() -> list[dict]:
    results = []

    for record in load_test_matrix():
        if record["top"] is None or record["file_number"] is None:
            continue

        if record["requested_tunnel_speed"] is None:
            continue

        mat_name = (
            f"Results_File_{int(record['file_number'])}"
            f"_TOP_{int(record['top'])}.mat"
        )
        mat_path = DATA_DIR / mat_name

        if not mat_path.exists():
            print(f"Missing file, skipped: {mat_name}")
            continue

        stats = compute_test(
            mat_path=mat_path,
            requested_tunnel_speed=float(record["requested_tunnel_speed"]),
        )

        results.append(
            {
                **record,
                **stats,
                "cp_plus_std": (
                    stats["cp_median"] + stats["cp_std"]
                    if np.isfinite(stats["cp_median"])
                    and np.isfinite(stats["cp_std"])
                    else math.nan
                ),
                "cp_minus_std": (
                    stats["cp_median"] - stats["cp_std"]
                    if np.isfinite(stats["cp_median"])
                    and np.isfinite(stats["cp_std"])
                    else math.nan
                ),
                "ct_plus_std": (
                    stats["ct_median"] + stats["ct_std"]
                    if np.isfinite(stats["ct_median"])
                    and np.isfinite(stats["ct_std"])
                    else math.nan
                ),
                "ct_minus_std": (
                    stats["ct_median"] - stats["ct_std"]
                    if np.isfinite(stats["ct_median"])
                    and np.isfinite(stats["ct_std"])
                    else math.nan
                ),
                "mat_file": mat_name,
            }
        )

    return results


# ============================================================
# Output
# ============================================================

def save_summary_csv(results: list[dict]) -> None:
    output_path = OUTPUT_DIR / "cp_ct_median_utunnel_summary.csv"

    columns = [
        "test",
        "yaw_target",
        "yaw_median",
        "requested_tunnel_speed",
        "pitot_velocity_median",
        "excel_equivalent_free_air_speed",
        "free_air_velocity_median",
        "rotor_speed_rpm_median",
        "tip_speed_median",
        "tsr_expected",
        "tsr_median",
        "cp_median",
        "cp_std",
        "cp_min",
        "cp_max",
        "ct_median",
        "ct_std",
        "ct_min",
        "ct_max",
        "rotor_thrust_median",
        "nacelle_drag_median",
        "tower_drag_force_median",
        "tower_drag_moment_median",
        "mat_file",
    ]

    with output_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=columns,
            delimiter=";",
        )
        writer.writeheader()

        for result in results:
            writer.writerow(
                {column: result[column] for column in columns}
            )


def plot_performance(
    results: list[dict],
    yaw_target: float,
    coefficient: str,
) -> None:
    selected = [
        result
        for result in results
        if result["yaw_target"] == yaw_target
    ]
    selected.sort(key=lambda result: result["tsr_median"])

    x = [result["tsr_median"] for result in selected]
    median = [result[f"{coefficient}_median"] for result in selected]
    plus_std = [result[f"{coefficient}_plus_std"] for result in selected]
    minus_std = [result[f"{coefficient}_minus_std"] for result in selected]
    minimum = [result[f"{coefficient}_min"] for result in selected]
    maximum = [result[f"{coefficient}_max"] for result in selected]

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(x, median, marker="o", label="Median")
    plt.plot(
        x,
        plus_std,
        marker="o",
        linestyle="--",
        label="Median + standard deviation",
    )
    plt.plot(
        x,
        minus_std,
        marker="o",
        linestyle="--",
        label="Median - standard deviation",
    )
    plt.plot(
        x,
        minimum,
        marker="o",
        linestyle=":",
        label="Minimum",
    )
    plt.plot(
        x,
        maximum,
        marker="o",
        linestyle=":",
        label="Maximum",
    )

    plt.xlabel("Tip-speed ratio, TSR [-]")
    plt.ylabel(f"{label} [-]")
    plt.title(
        f"{label} versus TSR at yaw = {yaw_target}° "
        "(using requested tunnel speed)"
    )
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    yaw_name = str(int(yaw_target)).replace("-", "minus")

    plt.savefig(
        OUTPUT_DIR / f"{label}_median_utunnel_yaw_{yaw_name}.png",
        dpi=200,
    )
    plt.close()


def main() -> None:
    if not DATA_DIR.exists():
        raise FileNotFoundError(f"Data folder not found: {DATA_DIR}")

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(
            f"Excel file not found: {TEST_MATRIX_PATH}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    results = process_all_tests()

    if not results:
        raise RuntimeError("No tests could be processed.")

    save_summary_csv(results)

    yaw_values = sorted(
        {
            result["yaw_target"]
            for result in results
            if result["yaw_target"] is not None
        }
    )

    for yaw_target in yaw_values:
        plot_performance(results, yaw_target, "cp")
        plot_performance(results, yaw_target, "ct")

    print("Finished.")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")
    print(f"Blockage ratio alpha: {BLOCKAGE_RATIO:.5f}")
    print(
        "This script uses the Excel requested tunnel speed, not PitotVelocity, "
        "as U_tunnel for the blockage correction."
    )
    print(
        "Final TSR, cP and cT use the resulting equivalent free-air velocity."
    )


if __name__ == "__main__":
    main()
