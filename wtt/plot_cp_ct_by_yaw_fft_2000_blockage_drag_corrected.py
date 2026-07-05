from __future__ import annotations

import csv
import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat


# ============================================================
# Paths and constants
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
TEST_MATRIX_PATH = DATA_DIR / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct_fft_2000_blockage_drag_corrected"

MEASUREMENT_DURATION_S = 8.0
TARGET_LENGTH = 2000

ROTOR_RADIUS_M = 0.5436
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2

TORQUE_SCALE_TO_NM = 1e-3

# Geometry for the thrust / moment correction.
# The user-provided hub height is used as the thrust/nacelle moment arm.
HUB_HEIGHT_M = 0.80
THRUST_LEVER_ARM_M = HUB_HEIGHT_M
NACELLE_HUB_MOMENT_ARM_M = HUB_HEIGHT_M

# Nacelle-hub drag from the slides:
# D_nacelle = 0.5 * rho * U^2 * SCD
NACELLE_HUB_SCD_M2 = 0.0175

# Tower drag from the slides:
# D_T = 0.5 * rho * D * integral_0^H V(h)^2 * C_D(h,Re) dh
TOWER_HEIGHT_M = HUB_HEIGHT_M
TOWER_DIAMETER_M = 0.047

# Required modelling assumption for C_D(h,Re).
# A typical first estimate for a circular cylinder is used.
TOWER_DRAG_COEFFICIENT = 1.2

# Wind-tunnel dimensions for blockage ratio alpha = A_rotor / A_tunnel.
TUNNEL_WIDTH_M = 2.7
TUNNEL_HEIGHT_M = 1.8
TUNNEL_CROSS_SECTION_M2 = TUNNEL_WIDTH_M * TUNNEL_HEIGHT_M
BLOCKAGE_RATIO_ALPHA = ROTOR_AREA_M2 / TUNNEL_CROSS_SECTION_M2

CORRECTION_MAX_ITERATIONS = 50
CORRECTION_TOLERANCE = 1e-8


# ============================================================
# Data helpers
# ============================================================

def unwrap_singleton(value):
    current = value

    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]

    return current


def get_array(value) -> np.ndarray:
    return np.asarray(
        unwrap_singleton(value),
        dtype=float,
    ).ravel(order="C")


def load_model1(mat_path: Path):
    mat_data = loadmat(
        mat_path,
        squeeze_me=True,
        struct_as_record=False,
    )
    return mat_data["Data"].ModelsData.Model1


def block_average_to_target_length(
    values: np.ndarray,
    target_length: int,
) -> np.ndarray:
    """Downsamples by block averaging.

    Example:
    Hub.Torque has 20000 samples. For target_length=2000, this function
    forms 2000 blocks with 10 samples each and returns one mean per block.
    """

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size % target_length != 0:
        raise ValueError(
            "Block averaging requires the signal length to be an integer "
            "multiple of target_length."
        )

    block_size = values.size // target_length
    blocks = values.reshape(target_length, block_size)

    with np.errstate(invalid="ignore"):
        return np.nanmean(blocks, axis=1)


def interpolate_to_target_length(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    """Linearly interpolates a signal over the common measurement period."""

    values = np.asarray(values, dtype=float).ravel(order="C")

    if target_length <= 0:
        raise ValueError("target_length must be greater than zero.")

    if values.size == 0:
        return np.full(target_length, np.nan)

    if values.size == 1:
        return np.full(target_length, values[0])

    if values.size == target_length:
        return values.astype(float, copy=False)

    finite = np.isfinite(values)

    if finite.sum() == 0:
        return np.full(target_length, np.nan)

    if finite.sum() == 1:
        return np.full(target_length, values[finite][0])

    source_time = np.linspace(
        0.0,
        duration_s,
        values.size,
        endpoint=False,
    )
    target_time = np.linspace(
        0.0,
        duration_s,
        target_length,
        endpoint=False,
    )

    return np.interp(
        target_time,
        source_time[finite],
        values[finite],
    )


def resample_to_2000(values: np.ndarray, signal_name: str) -> np.ndarray:
    """Resamples all signals to 2000 samples.

    - 2000-sample signals remain unchanged.
    - Longer integer-ratio signals are block-averaged.
      This is intended for Hub.Torque: 20000 -> 2000.
    - Shorter signals are interpolated.
      This is intended for AirDensity: 1000 -> 2000.
    """

    values = np.asarray(values, dtype=float).ravel(order="C")

    if values.size == TARGET_LENGTH:
        return values.astype(float, copy=False)

    if values.size > TARGET_LENGTH and values.size % TARGET_LENGTH == 0:
        return block_average_to_target_length(values, TARGET_LENGTH)

    return interpolate_to_target_length(values, TARGET_LENGTH)


def fill_nonfinite(values: np.ndarray) -> np.ndarray:
    """Fills invalid values linearly before the Fourier transform."""

    values = np.asarray(values, dtype=float).copy()
    finite = np.isfinite(values)

    if finite.sum() == 0:
        return np.full_like(values, np.nan)

    if finite.sum() == 1:
        return np.full_like(values, values[finite][0])

    if not finite.all():
        indices = np.arange(values.size)
        values[~finite] = np.interp(
            indices[~finite],
            indices[finite],
            values[finite],
        )

    return values


def safe_stat(values: np.ndarray, fn) -> float:
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]

    if values.size == 0:
        return math.nan

    return float(fn(values))


# ============================================================
# Fourier analysis
# ============================================================

def fourier_analysis(
    values: np.ndarray,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> dict:
    """Returns the one-sided FFT spectrum and the DC component.

    The normalized FFT component at 0 Hz equals the arithmetic mean of the
    original time series.
    """

    values = fill_nonfinite(values)

    if values.size == 0 or not np.isfinite(values).any():
        return {
            "dc": math.nan,
            "frequency_hz": np.array([]),
            "amplitude": np.array([]),
            "dominant_frequency_hz": math.nan,
            "dominant_amplitude": math.nan,
        }

    n = values.size
    dt = duration_s / n

    fft_values = np.fft.rfft(values)
    frequency_hz = np.fft.rfftfreq(n, d=dt)
    amplitude = np.abs(fft_values) / n

    if amplitude.size > 1:
        if n % 2 == 0:
            amplitude[1:-1] *= 2.0
        else:
            amplitude[1:] *= 2.0

    dc = float(np.real(fft_values[0]) / n)

    if amplitude.size > 1:
        dominant_index = int(np.argmax(amplitude[1:])) + 1
        dominant_frequency_hz = float(frequency_hz[dominant_index])
        dominant_amplitude = float(amplitude[dominant_index])
    else:
        dominant_frequency_hz = math.nan
        dominant_amplitude = math.nan

    return {
        "dc": dc,
        "frequency_hz": frequency_hz,
        "amplitude": amplitude,
        "dominant_frequency_hz": dominant_frequency_hz,
        "dominant_amplitude": dominant_amplitude,
    }


# ============================================================
# Blockage, nacelle-hub drag and tower drag
# ============================================================

def axial_induction_from_ct(ct: np.ndarray) -> np.ndarray:
    """Momentum-theory estimate for axial induction.

    CT = 4a(1-a) gives a = 0.5 * (1 - sqrt(1 - CT)).
    Values are clipped to the classical range.
    """

    ct = np.asarray(ct, dtype=float)
    ct_limited = np.clip(ct, 0.0, 0.999999)

    return 0.5 * (1.0 - np.sqrt(1.0 - ct_limited))


def glauert_velocity_from_thrust(
    u_inf: np.ndarray,
    air_density: np.ndarray,
    rotor_thrust_n: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Applies the Glauert blockage correction to the time series.

    Formula:
        U'_inf / U_inf = 1 + (alpha / 4) * C_T / (1 - C_T)

    with:
        C_T = T / (0.5 * rho * A_R * U_inf^2)

    C_T is the rotor thrust coefficient, not the torque coefficient.
    """

    u_inf = np.asarray(u_inf, dtype=float)
    air_density = np.asarray(air_density, dtype=float)
    rotor_thrust_n = np.asarray(rotor_thrust_n, dtype=float)

    ct_denominator_uncorrected = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * u_inf**2
    )

    ct_for_blockage = np.divide(
        rotor_thrust_n,
        ct_denominator_uncorrected,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(rotor_thrust_n)
            & np.isfinite(ct_denominator_uncorrected)
            & (ct_denominator_uncorrected > 0.0)
        ),
    )

    blockage_factor = np.full(TARGET_LENGTH, np.nan)

    valid = (
        np.isfinite(u_inf)
        & np.isfinite(ct_for_blockage)
        & (u_inf > 0.0)
        & (ct_for_blockage < 1.0)
        & np.isfinite(1.0 - ct_for_blockage)
        & ((1.0 - ct_for_blockage) != 0.0)
    )

    blockage_factor[valid] = (
        1.0
        + (BLOCKAGE_RATIO_ALPHA / 4.0)
        * ct_for_blockage[valid]
        / (1.0 - ct_for_blockage[valid])
    )

    u_blockage_corrected = u_inf * blockage_factor

    return u_blockage_corrected, ct_for_blockage, blockage_factor


def drag_corrected_rotor_thrust_and_velocity(
    tower_fa_moment: np.ndarray,
    pitot_velocity: np.ndarray,
    air_density: np.ndarray,
) -> dict[str, np.ndarray]:
    """Removes nacelle-hub and tower-drag moments from Tower.FA.

    Moment balance around the tower-base load measurement location:

        -M_FA = T_rotor * L
                + D_nacelle * z_hub
                + M_tower_drag

    Therefore:

        T_rotor = (-M_FA - D_nacelle*z_hub - M_tower_drag) / L

    Nacelle-hub drag:
        D_nacelle = 0.5 * rho * U^2 * SCD

    Tower drag force from the slides:
        D_T = 0.5 * rho * D * integral_0^H V(h)^2 * C_D(h,Re) dh

    Tower-base moment from tower drag:
        M_T = 0.5 * rho * D * integral_0^H h * V(h)^2 * C_D(h,Re) dh

    Assumption used here:
    - constant C_D = TOWER_DRAG_COEFFICIENT
    - uniform tower velocity V(h) = U'_inf * (1 - 2a)
    - a is estimated from rotor C_T
    """

    tower_fa_moment = np.asarray(tower_fa_moment, dtype=float)
    pitot_velocity = np.asarray(pitot_velocity, dtype=float)
    air_density = np.asarray(air_density, dtype=float)

    measured_positive_moment = -tower_fa_moment

    # Initial estimate: all measured moment is due to rotor thrust.
    rotor_thrust_n = measured_positive_moment / THRUST_LEVER_ARM_M

    u_blockage_corrected = pitot_velocity.copy()
    ct_for_blockage = np.full(TARGET_LENGTH, np.nan)
    blockage_factor = np.full(TARGET_LENGTH, np.nan)

    nacelle_drag_n = np.zeros(TARGET_LENGTH)
    nacelle_drag_moment_nm = np.zeros(TARGET_LENGTH)

    tower_drag_n = np.zeros(TARGET_LENGTH)
    tower_drag_moment_nm = np.zeros(TARGET_LENGTH)
    tower_velocity = np.zeros(TARGET_LENGTH)
    axial_induction = np.zeros(TARGET_LENGTH)

    for _ in range(CORRECTION_MAX_ITERATIONS):
        previous_rotor_thrust_n = rotor_thrust_n.copy()

        (
            u_blockage_corrected,
            ct_for_blockage,
            blockage_factor,
        ) = glauert_velocity_from_thrust(
            u_inf=pitot_velocity,
            air_density=air_density,
            rotor_thrust_n=rotor_thrust_n,
        )

        ct_corrected_denominator = (
            0.5
            * air_density
            * ROTOR_AREA_M2
            * u_blockage_corrected**2
        )

        ct_corrected = np.divide(
            rotor_thrust_n,
            ct_corrected_denominator,
            out=np.full(TARGET_LENGTH, np.nan),
            where=(
                np.isfinite(rotor_thrust_n)
                & np.isfinite(ct_corrected_denominator)
                & (ct_corrected_denominator > 0.0)
            ),
        )

        axial_induction = axial_induction_from_ct(ct_corrected)

        # Nacelle-hub drag at hub height, using the blockage-corrected velocity.
        nacelle_drag_n = (
            0.5
            * air_density
            * u_blockage_corrected**2
            * NACELLE_HUB_SCD_M2
        )

        nacelle_drag_moment_nm = (
            nacelle_drag_n
            * NACELLE_HUB_MOMENT_ARM_M
        )

        # Tower drag downstream of the rotor.
        tower_velocity = u_blockage_corrected * (1.0 - 2.0 * axial_induction)
        tower_velocity = np.maximum(tower_velocity, 0.0)

        # Integral approximation for tower drag force:
        # D_T = 0.5*rho*D*V^2*C_D*H
        tower_drag_n = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * TOWER_HEIGHT_M
        )

        # Integral approximation for tower-base moment:
        # M_T = 0.5*rho*D*V^2*C_D*H^2/2
        tower_drag_moment_nm = (
            0.5
            * air_density
            * TOWER_DIAMETER_M
            * tower_velocity**2
            * TOWER_DRAG_COEFFICIENT
            * (TOWER_HEIGHT_M**2 / 2.0)
        )

        rotor_moment_nm = (
            measured_positive_moment
            - nacelle_drag_moment_nm
            - tower_drag_moment_nm
        )

        rotor_thrust_n = np.divide(
            rotor_moment_nm,
            THRUST_LEVER_ARM_M,
            out=np.full(TARGET_LENGTH, np.nan),
            where=np.isfinite(rotor_moment_nm),
        )

        # Negative corrected rotor thrust is non-physical here.
        rotor_thrust_n = np.where(
            np.isfinite(rotor_thrust_n),
            np.maximum(rotor_thrust_n, 0.0),
            np.nan,
        )

        difference = np.abs(rotor_thrust_n - previous_rotor_thrust_n)
        finite_difference = difference[np.isfinite(difference)]

        if finite_difference.size == 0:
            break

        if np.max(finite_difference) < CORRECTION_TOLERANCE:
            break

    return {
        "measured_positive_moment": measured_positive_moment,
        "rotor_thrust_n": rotor_thrust_n,
        "u_blockage_corrected": u_blockage_corrected,
        "ct_for_blockage": ct_for_blockage,
        "blockage_factor": blockage_factor,

        "nacelle_drag_n": nacelle_drag_n,
        "nacelle_drag_moment_nm": nacelle_drag_moment_nm,

        "tower_drag_n": tower_drag_n,
        "tower_drag_moment_nm": tower_drag_moment_nm,
        "tower_velocity": tower_velocity,
        "axial_induction": axial_induction,
    }


# ============================================================
# cP, cT, TSR and FFT
# ============================================================

def compute_test(mat_path: Path) -> dict:
    model1 = load_model1(mat_path)

    rotor_speed_rpm = get_array(model1.RotorSpeed)
    pitot_velocity = get_array(model1.PitotVelocity)
    air_density = get_array(model1.AirDensity)
    torque_raw = get_array(model1.Hub.Torque)
    tower_fa = get_array(model1.Tower.FA)
    yaw = get_array(model1.Yaw)

    # Fixed common target length: 2000 samples.
    # Hub.Torque 20000 -> 2000 by block averaging.
    # AirDensity 1000 -> 2000 by interpolation.
    rotor_speed_rpm = resample_to_2000(rotor_speed_rpm, "RotorSpeed")
    pitot_velocity = resample_to_2000(pitot_velocity, "PitotVelocity")
    air_density = resample_to_2000(air_density, "AirDensity")
    torque_raw = resample_to_2000(torque_raw, "Hub.Torque")
    tower_fa = resample_to_2000(tower_fa, "Tower.FA")
    yaw = resample_to_2000(yaw, "Yaw")

    omega = rotor_speed_rpm * 2.0 * math.pi / 60.0
    tip_speed = omega * ROTOR_RADIUS_M
    torque_nm = torque_raw * TORQUE_SCALE_TO_NM
    power_w = torque_nm * omega

    correction = drag_corrected_rotor_thrust_and_velocity(
        tower_fa_moment=tower_fa,
        pitot_velocity=pitot_velocity,
        air_density=air_density,
    )

    thrust_raw_from_tower_fa = -tower_fa / THRUST_LEVER_ARM_M
    thrust_n = correction["rotor_thrust_n"]
    u_blockage_corrected = correction["u_blockage_corrected"]

    # From this point onward, only u_blockage_corrected is used for TSR, cP and cT.
    # The thrust used for cT and the Glauert correction is rotor thrust after
    # nacelle-hub and tower-drag moment removal.
    tsr = np.divide(
        tip_speed,
        u_blockage_corrected,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(u_blockage_corrected)
            & (u_blockage_corrected != 0.0)
        ),
    )

    cp_denominator = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * u_blockage_corrected**3
    )

    cp = np.divide(
        power_w,
        cp_denominator,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(cp_denominator)
            & (cp_denominator != 0.0)
        ),
    )

    ct_denominator = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * u_blockage_corrected**2
    )

    ct = np.divide(
        thrust_n,
        ct_denominator,
        out=np.full(TARGET_LENGTH, np.nan),
        where=(
            np.isfinite(ct_denominator)
            & (ct_denominator != 0.0)
        ),
    )

    return {
        "yaw_fft": fourier_analysis(yaw),
        "tsr_fft": fourier_analysis(tsr),
        "tip_speed_fft": fourier_analysis(tip_speed),
        "cp_fft": fourier_analysis(cp),
        "ct_fft": fourier_analysis(ct),

        "pitot_fft": fourier_analysis(pitot_velocity),
        "u_blockage_corrected_fft": fourier_analysis(u_blockage_corrected),
        "blockage_factor_fft": fourier_analysis(correction["blockage_factor"]),
        "ct_for_blockage_fft": fourier_analysis(correction["ct_for_blockage"]),

        "thrust_raw_fft": fourier_analysis(thrust_raw_from_tower_fa),
        "rotor_thrust_fft": fourier_analysis(thrust_n),

        "nacelle_drag_fft": fourier_analysis(correction["nacelle_drag_n"]),
        "nacelle_drag_moment_fft": fourier_analysis(
            correction["nacelle_drag_moment_nm"]
        ),

        "tower_drag_fft": fourier_analysis(correction["tower_drag_n"]),
        "tower_drag_moment_fft": fourier_analysis(
            correction["tower_drag_moment_nm"]
        ),
        "tower_velocity_fft": fourier_analysis(correction["tower_velocity"]),
        "axial_induction_fft": fourier_analysis(correction["axial_induction"]),
    }


# ============================================================
# Test matrix
# ============================================================

def load_test_matrix() -> list[dict]:
    workbook = openpyxl.load_workbook(
        TEST_MATRIX_PATH,
        data_only=True,
    )
    sheet = workbook["Task1 Performance G2"]

    rows = []

    for row_index in range(2, sheet.max_row + 1):
        test_number = sheet.cell(row_index, 1).value

        if test_number is None:
            continue

        rows.append(
            {
                "test": test_number,
                "yaw_target": sheet.cell(row_index, 6).value,
                "tsr_expected": sheet.cell(row_index, 8).value,
                "top": sheet.cell(row_index, 29).value,
                "file_number": sheet.cell(row_index, 30).value,
            }
        )

    return rows


def process_all_tests() -> list[dict]:
    results = []

    for test in load_test_matrix():
        if test["top"] is None or test["file_number"] is None:
            continue

        mat_name = (
            f"Results_File_{int(test['file_number'])}"
            f"_TOP_{int(test['top'])}.mat"
        )
        mat_path = DATA_DIR / mat_name

        if not mat_path.exists():
            print(f"Missing file, skipped: {mat_name}")
            continue

        fft_result = compute_test(mat_path)

        results.append(
            {
                **test,
                "mat_file": mat_name,

                "yaw_fft_dc": fft_result["yaw_fft"]["dc"],
                "tsr_fft_dc": fft_result["tsr_fft"]["dc"],
                "tip_speed_fft_dc": fft_result["tip_speed_fft"]["dc"],
                "cp_fft_dc": fft_result["cp_fft"]["dc"],
                "ct_fft_dc": fft_result["ct_fft"]["dc"],

                "pitot_fft_dc": fft_result["pitot_fft"]["dc"],
                "u_blockage_corrected_fft_dc":
                    fft_result["u_blockage_corrected_fft"]["dc"],
                "blockage_factor_fft_dc":
                    fft_result["blockage_factor_fft"]["dc"],
                "ct_for_blockage_fft_dc":
                    fft_result["ct_for_blockage_fft"]["dc"],

                "thrust_raw_fft_dc": fft_result["thrust_raw_fft"]["dc"],
                "rotor_thrust_fft_dc":
                    fft_result["rotor_thrust_fft"]["dc"],

                "nacelle_drag_fft_dc":
                    fft_result["nacelle_drag_fft"]["dc"],
                "nacelle_drag_moment_fft_dc":
                    fft_result["nacelle_drag_moment_fft"]["dc"],

                "tower_drag_fft_dc":
                    fft_result["tower_drag_fft"]["dc"],
                "tower_drag_moment_fft_dc":
                    fft_result["tower_drag_moment_fft"]["dc"],
                "tower_velocity_fft_dc":
                    fft_result["tower_velocity_fft"]["dc"],
                "axial_induction_fft_dc":
                    fft_result["axial_induction_fft"]["dc"],

                "cp_dominant_frequency_hz":
                    fft_result["cp_fft"]["dominant_frequency_hz"],
                "cp_dominant_amplitude":
                    fft_result["cp_fft"]["dominant_amplitude"],
                "ct_dominant_frequency_hz":
                    fft_result["ct_fft"]["dominant_frequency_hz"],
                "ct_dominant_amplitude":
                    fft_result["ct_fft"]["dominant_amplitude"],

                "cp_frequency_hz": fft_result["cp_fft"]["frequency_hz"],
                "cp_amplitude": fft_result["cp_fft"]["amplitude"],
                "ct_frequency_hz": fft_result["ct_fft"]["frequency_hz"],
                "ct_amplitude": fft_result["ct_fft"]["amplitude"],
            }
        )

    return results


# ============================================================
# Output
# ============================================================

def save_summary_csv(results: list[dict]) -> None:
    output_path = OUTPUT_DIR / "cp_ct_fft_summary_by_test.csv"

    columns = [
        "test",
        "yaw_target",
        "yaw_fft_dc",

        "tsr_expected",
        "tsr_fft_dc",
        "tip_speed_fft_dc",

        "cp_fft_dc",
        "ct_fft_dc",

        "pitot_fft_dc",
        "u_blockage_corrected_fft_dc",
        "blockage_factor_fft_dc",
        "ct_for_blockage_fft_dc",

        "thrust_raw_fft_dc",
        "rotor_thrust_fft_dc",

        "nacelle_drag_fft_dc",
        "nacelle_drag_moment_fft_dc",

        "tower_drag_fft_dc",
        "tower_drag_moment_fft_dc",
        "tower_velocity_fft_dc",
        "axial_induction_fft_dc",

        "cp_dominant_frequency_hz",
        "cp_dominant_amplitude",
        "ct_dominant_frequency_hz",
        "ct_dominant_amplitude",

        "mat_file",
    ]

    with output_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=columns,
            delimiter=";",
        )
        writer.writeheader()

        for result in results:
            writer.writerow(
                {column: result[column] for column in columns}
            )


def plot_performance(
    results: list[dict],
    yaw_target: float,
    coefficient: str,
) -> None:
    selected = [
        result
        for result in results
        if result["yaw_target"] == yaw_target
    ]
    selected.sort(key=lambda result: result["tsr_fft_dc"])

    x = [result["tsr_fft_dc"] for result in selected]
    y = [result[f"{coefficient}_fft_dc"] for result in selected]

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(
        x,
        y,
        marker="o",
        label="FFT DC component",
    )
    plt.xlabel("TSR from FFT DC component")
    plt.ylabel(f"{label} [-]")
    plt.title(f"{label} over TSR at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    yaw_name = str(int(yaw_target)).replace("-", "minus")
    plt.savefig(
        OUTPUT_DIR / f"{label}_fft_yaw_{yaw_name}.png",
        dpi=200,
    )
    plt.close()


def plot_mean_spectrum(
    results: list[dict],
    yaw_target: float,
    coefficient: str,
) -> None:
    selected = [
        result
        for result in results
        if result["yaw_target"] == yaw_target
    ]

    if not selected:
        return

    frequency_key = f"{coefficient}_frequency_hz"
    amplitude_key = f"{coefficient}_amplitude"

    lengths = [
        len(result[frequency_key])
        for result in selected
        if len(result[frequency_key]) > 0
    ]

    if not lengths:
        return

    common_length = min(lengths)
    frequency = selected[0][frequency_key][:common_length]
    amplitudes = np.vstack(
        [
            result[amplitude_key][:common_length]
            for result in selected
        ]
    )
    mean_amplitude = np.nanmean(amplitudes, axis=0)

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(
        frequency,
        mean_amplitude,
        label="Mean one-sided FFT spectrum",
    )
    plt.xlabel("Frequency [Hz]")
    plt.ylabel(f"{label} amplitude")
    plt.title(f"{label} Fourier spectrum at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    yaw_name = str(int(yaw_target)).replace("-", "minus")
    plt.savefig(
        OUTPUT_DIR / f"{label}_spectrum_yaw_{yaw_name}.png",
        dpi=200,
    )
    plt.close()


def main() -> None:
    if not DATA_DIR.exists():
        raise FileNotFoundError(f"Data folder not found: {DATA_DIR}")

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(
            f"Excel file not found: {TEST_MATRIX_PATH}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    results = process_all_tests()

    if not results:
        raise RuntimeError("No tests could be processed.")

    save_summary_csv(results)

    yaw_values = sorted(
        {
            result["yaw_target"]
            for result in results
            if result["yaw_target"] is not None
        }
    )

    for yaw_target in yaw_values:
        plot_performance(results, yaw_target, "cp")
        plot_performance(results, yaw_target, "ct")
        plot_mean_spectrum(results, yaw_target, "cp")
        plot_mean_spectrum(results, yaw_target, "ct")

    print("Finished.")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")

    print("\nIMPORTANT:")
    print("All signals are processed on a 2000-sample time axis.")
    print("Hub.Torque is block-averaged from 20000 to 2000 samples.")
    print("AirDensity is interpolated from 1000 to 2000 samples.")
    print("Nacelle-hub and tower-drag moments are removed from Tower.FA.")
    print("Glauert blockage correction is applied to the time series.")
    print("After u_blockage_corrected is computed, TSR, cP and cT use only this velocity.")
    print(f"Hub height / thrust arm: {HUB_HEIGHT_M:.3f} m")
    print(f"Nacelle-hub SCD: {NACELLE_HUB_SCD_M2:.5f} m²")
    print(f"Tower diameter: {TOWER_DIAMETER_M:.3f} m")
    print(f"Tower drag coefficient assumption: {TOWER_DRAG_COEFFICIENT:.2f}")
    print(f"Blockage ratio alpha: {BLOCKAGE_RATIO_ALPHA:.5f}")
    print(
        "The cP/cT performance curves use the normalized FFT value at 0 Hz. "
        "This DC component equals the arithmetic mean, not the median."
    )


if __name__ == "__main__":
    main()
