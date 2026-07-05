from __future__ import annotations

import csv
import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.io import loadmat

from blockage import blockage_ratio, equivalent_free_air_speed


# ============================================================
# Paths and constants
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
TEST_MATRIX_PATH = DATA_DIR / "Task1_Performance_TestMatrix_Group2_measurements.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "plots_cp_ct_fft"

MEASUREMENT_DURATION_S = 8.0

ROTOR_RADIUS_M = 0.5436
ROTOR_AREA_M2 = math.pi * ROTOR_RADIUS_M**2

THRUST_LEVER_ARM_M = 0.72
TORQUE_SCALE_TO_NM = 1e-3

TUNNEL_WIDTH_M = 2.7
TUNNEL_HEIGHT_M = 1.8
TUNNEL_CROSS_SECTION_M2 = TUNNEL_WIDTH_M * TUNNEL_HEIGHT_M

BLOCKAGE_RATIO = blockage_ratio(
    rotor_area=ROTOR_AREA_M2,
    tunnel_area=TUNNEL_CROSS_SECTION_M2,
)


# ============================================================
# Data helpers
# ============================================================

def unwrap_singleton(value):
    current = value
    while isinstance(current, np.ndarray) and current.size == 1:
        current = current.flat[0]
    return current


def get_array(value) -> np.ndarray:
    return np.asarray(
        unwrap_singleton(value),
        dtype=float,
    ).ravel(order="C")


def load_model1(mat_path: Path):
    mat_data = loadmat(
        mat_path,
        squeeze_me=True,
        struct_as_record=False,
    )
    return mat_data["Data"].ModelsData.Model1


def resample_to_common_time(
    values: np.ndarray,
    target_length: int,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> np.ndarray:
    """Linear interpolation over the common eight-second measurement period."""

    values = np.asarray(values, dtype=float).ravel(order="C")

    if target_length <= 0:
        raise ValueError("target_length must be greater than zero.")

    if values.size == 0:
        return np.full(target_length, np.nan)

    if values.size == 1:
        return np.full(target_length, values[0])

    if values.size == target_length:
        return values

    finite = np.isfinite(values)

    if finite.sum() == 0:
        return np.full(target_length, np.nan)

    if finite.sum() == 1:
        return np.full(target_length, values[finite][0])

    source_time = np.linspace(
        0.0,
        duration_s,
        values.size,
        endpoint=False,
    )
    target_time = np.linspace(
        0.0,
        duration_s,
        target_length,
        endpoint=False,
    )

    return np.interp(
        target_time,
        source_time[finite],
        values[finite],
    )


def fill_nonfinite(values: np.ndarray) -> np.ndarray:
    """Fills internal invalid values linearly before the Fourier transform."""

    values = np.asarray(values, dtype=float).copy()
    finite = np.isfinite(values)

    if finite.sum() == 0:
        return np.full_like(values, np.nan)

    if finite.sum() == 1:
        return np.full_like(values, values[finite][0])

    if not finite.all():
        indices = np.arange(values.size)
        values[~finite] = np.interp(
            indices[~finite],
            indices[finite],
            values[finite],
        )

    return values


# ============================================================
# Fourier analysis
# ============================================================

def fourier_analysis(
    values: np.ndarray,
    duration_s: float = MEASUREMENT_DURATION_S,
) -> dict:
    """Returns the one-sided FFT spectrum and its zero-frequency component.

    The FFT value at 0 Hz is the DC component. After normalization it equals
    the arithmetic mean of the original signal. This replaces the previously
    used median as the representative value for each test.
    """

    values = fill_nonfinite(values)

    if values.size == 0 or not np.isfinite(values).any():
        return {
            "dc": math.nan,
            "frequency_hz": np.array([]),
            "amplitude": np.array([]),
            "dominant_frequency_hz": math.nan,
            "dominant_amplitude": math.nan,
        }

    n = values.size
    dt = duration_s / n

    fft_values = np.fft.rfft(values)
    frequency_hz = np.fft.rfftfreq(n, d=dt)
    amplitude = np.abs(fft_values) / n

    if amplitude.size > 1:
        if n % 2 == 0:
            amplitude[1:-1] *= 2.0
        else:
            amplitude[1:] *= 2.0

    dc = float(np.real(fft_values[0]) / n)

    if amplitude.size > 1:
        dominant_index = int(np.argmax(amplitude[1:])) + 1
        dominant_frequency_hz = float(frequency_hz[dominant_index])
        dominant_amplitude = float(amplitude[dominant_index])
    else:
        dominant_frequency_hz = math.nan
        dominant_amplitude = math.nan

    return {
        "dc": dc,
        "frequency_hz": frequency_hz,
        "amplitude": amplitude,
        "dominant_frequency_hz": dominant_frequency_hz,
        "dominant_amplitude": dominant_amplitude,
    }


# ============================================================
# cP, cT and TSR
# ============================================================

def compute_test(mat_path: Path) -> dict:
    model1 = load_model1(mat_path)

    rotor_speed_rpm = get_array(model1.RotorSpeed)
    pitot_velocity = get_array(model1.PitotVelocity)
    air_density = get_array(model1.AirDensity)
    torque_raw = get_array(model1.Hub.Torque)
    tower_fa = get_array(model1.Tower.FA)
    yaw = get_array(model1.Yaw)

    common_length = max(
        rotor_speed_rpm.size,
        pitot_velocity.size,
        air_density.size,
        torque_raw.size,
        tower_fa.size,
        yaw.size,
    )

    rotor_speed_rpm = resample_to_common_time(
        rotor_speed_rpm,
        common_length,
    )
    pitot_velocity = resample_to_common_time(
        pitot_velocity,
        common_length,
    )
    air_density = resample_to_common_time(
        air_density,
        common_length,
    )
    torque_raw = resample_to_common_time(
        torque_raw,
        common_length,
    )
    tower_fa = resample_to_common_time(
        tower_fa,
        common_length,
    )
    yaw = resample_to_common_time(
        yaw,
        common_length,
    )

    omega = rotor_speed_rpm * 2.0 * math.pi / 60.0
    torque_nm = torque_raw * TORQUE_SCALE_TO_NM
    power_w = torque_nm * omega
    thrust_n = -tower_fa / THRUST_LEVER_ARM_M

    # Initial cT based on measured Pitot velocity.
    ct_uncorrected_denominator = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * pitot_velocity**2
    )

    ct_uncorrected = np.divide(
        thrust_n,
        ct_uncorrected_denominator,
        out=np.full(common_length, np.nan),
        where=(
            np.isfinite(ct_uncorrected_denominator)
            & (ct_uncorrected_denominator != 0.0)
        ),
    )

    # Equivalent free-air velocity using blockage.py.
    valid = (
        np.isfinite(pitot_velocity)
        & np.isfinite(ct_uncorrected)
        & (pitot_velocity > 0.0)
        & (ct_uncorrected < 1.0)
    )

    free_air_velocity = np.full(common_length, np.nan)

    free_air_velocity[valid] = equivalent_free_air_speed(
        commanded_tunnel_speed_value=pitot_velocity[valid],
        thrust_coefficient=ct_uncorrected[valid],
        blockage_ratio_value=BLOCKAGE_RATIO,
    )

    tsr = np.divide(
        omega * ROTOR_RADIUS_M,
        free_air_velocity,
        out=np.full(common_length, np.nan),
        where=np.isfinite(free_air_velocity) & (free_air_velocity != 0.0),
    )

    cp_denominator = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * free_air_velocity**3
    )
    cp = np.divide(
        power_w,
        cp_denominator,
        out=np.full(common_length, np.nan),
        where=np.isfinite(cp_denominator) & (cp_denominator != 0.0),
    )

    ct_denominator = (
        0.5
        * air_density
        * ROTOR_AREA_M2
        * free_air_velocity**2
    )
    ct = np.divide(
        thrust_n,
        ct_denominator,
        out=np.full(common_length, np.nan),
        where=np.isfinite(ct_denominator) & (ct_denominator != 0.0),
    )

    return {
        "yaw_fft": fourier_analysis(yaw),
        "tsr_fft": fourier_analysis(tsr),
        "cp_fft": fourier_analysis(cp),
        "ct_fft": fourier_analysis(ct),
        "pitot_fft": fourier_analysis(pitot_velocity),
        "free_air_fft": fourier_analysis(free_air_velocity),
    }


# ============================================================
# Test matrix
# ============================================================

def load_test_matrix() -> list[dict]:
    workbook = openpyxl.load_workbook(
        TEST_MATRIX_PATH,
        data_only=True,
    )
    sheet = workbook["Task1 Performance G2"]

    rows = []

    for row_index in range(2, sheet.max_row + 1):
        test_number = sheet.cell(row_index, 1).value

        if test_number is None:
            continue

        rows.append(
            {
                "test": test_number,
                "yaw_target": sheet.cell(row_index, 6).value,
                "tsr_expected": sheet.cell(row_index, 8).value,
                "top": sheet.cell(row_index, 29).value,
                "file_number": sheet.cell(row_index, 30).value,
            }
        )

    return rows


def process_all_tests() -> list[dict]:
    results = []

    for test in load_test_matrix():
        if test["top"] is None or test["file_number"] is None:
            continue

        mat_name = (
            f"Results_File_{int(test['file_number'])}"
            f"_TOP_{int(test['top'])}.mat"
        )
        mat_path = DATA_DIR / mat_name

        if not mat_path.exists():
            print(f"Missing file, skipped: {mat_name}")
            continue

        fft_result = compute_test(mat_path)

        results.append(
            {
                **test,
                "mat_file": mat_name,
                "yaw_fft_dc": fft_result["yaw_fft"]["dc"],
                "tsr_fft_dc": fft_result["tsr_fft"]["dc"],
                "cp_fft_dc": fft_result["cp_fft"]["dc"],
                "ct_fft_dc": fft_result["ct_fft"]["dc"],
                "pitot_fft_dc": fft_result["pitot_fft"]["dc"],
                "free_air_fft_dc": fft_result["free_air_fft"]["dc"],
                "cp_dominant_frequency_hz":
                    fft_result["cp_fft"]["dominant_frequency_hz"],
                "cp_dominant_amplitude":
                    fft_result["cp_fft"]["dominant_amplitude"],
                "ct_dominant_frequency_hz":
                    fft_result["ct_fft"]["dominant_frequency_hz"],
                "ct_dominant_amplitude":
                    fft_result["ct_fft"]["dominant_amplitude"],
                "cp_frequency_hz": fft_result["cp_fft"]["frequency_hz"],
                "cp_amplitude": fft_result["cp_fft"]["amplitude"],
                "ct_frequency_hz": fft_result["ct_fft"]["frequency_hz"],
                "ct_amplitude": fft_result["ct_fft"]["amplitude"],
            }
        )

    return results


# ============================================================
# Output
# ============================================================

def save_summary_csv(results: list[dict]) -> None:
    output_path = OUTPUT_DIR / "cp_ct_fft_summary_by_test.csv"

    columns = [
        "test",
        "yaw_target",
        "yaw_fft_dc",
        "tsr_expected",
        "tsr_fft_dc",
        "cp_fft_dc",
        "ct_fft_dc",
        "pitot_fft_dc",
        "free_air_fft_dc",
        "cp_dominant_frequency_hz",
        "cp_dominant_amplitude",
        "ct_dominant_frequency_hz",
        "ct_dominant_amplitude",
        "mat_file",
    ]

    with output_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=columns,
            delimiter=";",
        )
        writer.writeheader()

        for result in results:
            writer.writerow(
                {column: result[column] for column in columns}
            )


def plot_performance(
    results: list[dict],
    yaw_target: float,
    coefficient: str,
) -> None:
    selected = [
        result
        for result in results
        if result["yaw_target"] == yaw_target
    ]
    selected.sort(key=lambda result: result["tsr_fft_dc"])

    x = [result["tsr_fft_dc"] for result in selected]
    y = [result[f"{coefficient}_fft_dc"] for result in selected]

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(
        x,
        y,
        marker="o",
        label="FFT DC component",
    )
    plt.xlabel("TSR from FFT DC component")
    plt.ylabel(label)
    plt.title(f"{label} over TSR at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    yaw_name = str(int(yaw_target)).replace("-", "minus")
    plt.savefig(
        OUTPUT_DIR / f"{label}_fft_yaw_{yaw_name}.png",
        dpi=200,
    )
    plt.close()


def plot_mean_spectrum(
    results: list[dict],
    yaw_target: float,
    coefficient: str,
) -> None:
    selected = [
        result
        for result in results
        if result["yaw_target"] == yaw_target
    ]

    if not selected:
        return

    frequency_key = f"{coefficient}_frequency_hz"
    amplitude_key = f"{coefficient}_amplitude"

    lengths = [
        len(result[frequency_key])
        for result in selected
        if len(result[frequency_key]) > 0
    ]

    if not lengths:
        return

    common_length = min(lengths)
    frequency = selected[0][frequency_key][:common_length]
    amplitudes = np.vstack(
        [
            result[amplitude_key][:common_length]
            for result in selected
        ]
    )
    mean_amplitude = np.nanmean(amplitudes, axis=0)

    label = "cP" if coefficient == "cp" else "cT"

    plt.figure(figsize=(10, 6))
    plt.plot(
        frequency,
        mean_amplitude,
        label="Mean one-sided FFT spectrum",
    )
    plt.xlabel("Frequency [Hz]")
    plt.ylabel(f"{label} amplitude")
    plt.title(f"{label} Fourier spectrum at yaw = {yaw_target}°")
    plt.grid(True)
    plt.legend()
    plt.tight_layout()

    yaw_name = str(int(yaw_target)).replace("-", "minus")
    plt.savefig(
        OUTPUT_DIR / f"{label}_spectrum_yaw_{yaw_name}.png",
        dpi=200,
    )
    plt.close()


def main() -> None:
    if not DATA_DIR.exists():
        raise FileNotFoundError(f"Data folder not found: {DATA_DIR}")

    if not TEST_MATRIX_PATH.exists():
        raise FileNotFoundError(
            f"Excel file not found: {TEST_MATRIX_PATH}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    results = process_all_tests()

    if not results:
        raise RuntimeError("No tests could be processed.")

    save_summary_csv(results)

    yaw_values = sorted(
        {
            result["yaw_target"]
            for result in results
            if result["yaw_target"] is not None
        }
    )

    for yaw_target in yaw_values:
        plot_performance(results, yaw_target, "cp")
        plot_performance(results, yaw_target, "ct")
        plot_mean_spectrum(results, yaw_target, "cp")
        plot_mean_spectrum(results, yaw_target, "ct")

    print("Finished.")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")
    print(f"Blockage ratio: {BLOCKAGE_RATIO:.5f}")
    print(
        "The cP/cT performance curves use the normalized FFT value at 0 Hz. "
        "This DC component equals the arithmetic mean, not the median."
    )


if __name__ == "__main__":
    main()
